"""Appendix 1 checks that compare a generated deck with its template.

Palette, text contrast, layout guides, edge margins, picture proportions, brand elements,
slide fill and chart axis labels. They read only coordinates, sizes, colour codes and
layout references in the files, so the same deck always gives the same issues.
"""

from __future__ import annotations

import colorsys
import re
from pathlib import Path
from typing import Any, NamedTuple

from lxml import etree
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

from .audit import _issue
from .models import PreparedTemplate, SlideContent

P15 = "{http://schemas.microsoft.com/office/powerpoint/2012/main}"
EMU_PER_CM = 360000
GUIDE_TOLERANCE = 54864        # 0.06 in: an edge this close sits on the guide
MARGIN_TOLERANCE = 91440       # 0.1 in beyond the layout's content area
BRAND_TOLERANCE = 19050        # 1.5 pt
COLOR_TOLERANCE = 30           # RGB distance still read as the same palette colour
NEUTRAL_SPREAD = 24            # black, white and greys are not brand colours
MIN_CONTRAST, LARGE_TEXT_CONTRAST = 4.5, 3.0
MIN_FILL, MAX_FILL = 0.25, 0.75
DISTORTION = 0.03
CONTENT_PLACEHOLDERS = {
    PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY,
    PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.CHART, PP_PLACEHOLDER.TABLE, PP_PLACEHOLDER.PICTURE,
    PP_PLACEHOLDER.VERTICAL_TITLE, PP_PLACEHOLDER.VERTICAL_BODY, PP_PLACEHOLDER.VERTICAL_OBJECT,
}
FOOTER_PLACEHOLDERS = {PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE}
AXIS_FREE_CHARTS = {XL_CHART_TYPE.PIE, XL_CHART_TYPE.PIE_EXPLODED, XL_CHART_TYPE.DOUGHNUT,
                    XL_CHART_TYPE.DOUGHNUT_EXPLODED, XL_CHART_TYPE.THREE_D_PIE, XL_CHART_TYPE.THREE_D_PIE_EXPLODED}

Rgb = tuple[int, int, int]


class Block(NamedTuple):
    """A generated block the grid checks judge: one text box, chart, table or a whole diagram."""
    shape_id: int
    name: str
    box: tuple[int, int, int, int]


# --- Colours -----------------------------------------------------------------------------

class _Theme:
    """Colour scheme, colour map and fill styles of one slide master."""

    def __init__(self, master):
        theme = etree.fromstring(master.part.part_related_by(RT.THEME).blob)
        self.scheme: dict[str, Rgb] = {}
        scheme = theme.find(".//" + qn("a:clrScheme"))
        for item in scheme if scheme is not None else []:
            value = _base_color(item[0]) if len(item) else None
            if value is not None:
                self.scheme[etree.QName(item).localname] = value
        mapping = master._element.find(qn("p:clrMap"))
        self.map = dict(mapping.attrib) if mapping is not None else {}
        fills, backgrounds = theme.find(".//" + qn("a:fillStyleLst")), theme.find(".//" + qn("a:bgFillStyleLst"))
        self.fills = list(fills) if fills is not None else []
        self.background_fills = list(backgrounds) if backgrounds is not None else []

    def color(self, element) -> Rgb | None:
        """a:srgbClr, a:schemeClr, a:sysClr or a:prstClr with its luminance modifiers."""
        name = etree.QName(element).localname
        if name == "schemeClr":
            key = element.get("val", "")
            base = self.scheme.get(self.map.get(key, key))
        else:
            base = _base_color(element)
        if base is None:
            return None
        red, green, blue = (channel / 255 for channel in base)
        hue, light, saturation = colorsys.rgb_to_hls(red, green, blue)
        for modifier in element:
            kind, value = etree.QName(modifier).localname, int(modifier.get("val", "100000")) / 100000
            if kind == "lumMod":
                light *= value
            elif kind == "lumOff":
                light += value
            elif kind == "tint":
                light = light + (1 - light) * (1 - value)
            elif kind == "shade":
                light *= value
        red, green, blue = colorsys.hls_to_rgb(hue, max(0.0, min(1.0, light)), saturation)
        return round(red * 255), round(green * 255), round(blue * 255)

    def fill(self, parent) -> tuple[str | None, Rgb | None]:
        """('solid', colour), ('none', None), ('complex', None) or (None, None) when not set here."""
        if parent is None:
            return None, None
        for child in parent:
            kind = etree.QName(child).localname
            if kind == "solidFill":
                return ("solid", self.color(child[0])) if len(child) else ("complex", None)
            if kind == "noFill":
                return "none", None
            if kind in {"gradFill", "blipFill", "pattFill", "grpFill"}:
                return "complex", None
        return None, None

    def styled_fill(self, reference, styles) -> tuple[str, Rgb | None]:
        """A fill that a theme style supplies (p:style/a:fillRef or p:bgRef)."""
        index = int(reference.get("idx", "0"))
        index = index - 1000 if index > 1000 else index
        if index <= 0 or index > len(styles):
            return "none", None
        if etree.QName(styles[index - 1]).localname != "solidFill":
            return "complex", None
        return "solid", self.color(reference[0]) if len(reference) else None


def _parse_hex(value: str | None) -> Rgb | None:
    value = (value or "").lstrip("#")
    if not re.fullmatch(r"[0-9A-Fa-f]{6}", value):
        return None
    return int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16)


def _base_color(element) -> Rgb | None:
    name = etree.QName(element).localname
    value = None
    if name == "srgbClr":
        value = element.get("val")
    elif name == "sysClr":
        value = element.get("lastClr") or {"windowText": "000000", "window": "FFFFFF"}.get(element.get("val", ""))
    elif name == "prstClr":
        value = {"black": "000000", "white": "FFFFFF"}.get(element.get("val", ""))
    return _parse_hex(value)


def _hex(color: Rgb) -> str:
    return "#%02X%02X%02X" % color


def _neutral(color: Rgb) -> bool:
    return max(color) - min(color) <= NEUTRAL_SPREAD


def _distance(first: Rgb, second: Rgb) -> float:
    return sum((a - b) ** 2 for a, b in zip(first, second)) ** 0.5


def _luminance(color: Rgb) -> float:
    def linear(channel: int) -> float:
        value = channel / 255
        return value / 12.92 if value <= 0.04045 else ((value + 0.055) / 1.055) ** 2.4
    return sum(weight * linear(channel) for weight, channel in zip((0.2126, 0.7152, 0.0722), color))


def contrast(first: Rgb, second: Rgb) -> float:
    light, dark = sorted((_luminance(first), _luminance(second)), reverse=True)
    return (light + 0.05) / (dark + 0.05)


def template_palette(deck) -> set[Rgb]:
    """Theme colours plus every explicit colour the template's masters, layouts and slides use."""
    colors: set[Rgb] = set()
    parts = []
    for master in deck.slide_masters:
        colors.update(_Theme(master).scheme.values())
        parts.append(master.part)
        parts += [layout.part for layout in master.slide_layouts]
    parts += [slide.part for slide in deck.slides]
    for part in parts:
        for element in part._element.iter(qn("a:srgbClr")):
            value = _base_color(element)
            if value is not None:
                colors.add(value)
    return colors


# --- Geometry ----------------------------------------------------------------------------

def _box(shape) -> tuple[int, int, int, int] | None:
    if shape.left is None or shape.top is None or shape.width is None or shape.height is None:
        return None
    return int(shape.left), int(shape.top), int(shape.left + shape.width), int(shape.top + shape.height)


def _covered(cover, target) -> float:
    """Share of target's area under cover."""
    first, second = _box(cover), _box(target)
    if first is None or second is None:
        return 0.0
    width = max(0, min(first[2], second[2]) - max(first[0], second[0]))
    height = max(0, min(first[3], second[3]) - max(first[1], second[1]))
    return width * height / max(1, (second[2] - second[0]) * (second[3] - second[1]))


def _cm(value: float) -> str:
    return f"{abs(value) / EMU_PER_CM:.1f}".replace(".", ",") + " см"


def _guides(deck, layout) -> list[int]:
    """Vertical drawing guides (p15:sldGuideLst) of the presentation, master and layout, in EMU."""
    positions = []
    for owner in (deck.part._element, layout.slide_master._element, layout._element):
        for guide in owner.iter(P15 + "guide"):
            if guide.get("orient", "vert") == "vert":
                positions.append(round(int(guide.get("pos", "0")) * 1587.5))
    return positions


def content_area(template) -> tuple[int, int, int, int]:
    """Where the template itself puts text and pictures; generated content beyond it enters the edge margins."""
    width, height = int(template.slide_width), int(template.slide_height)
    boxes = []
    for slide in template.slides:
        for shape in slide.shapes:
            box = _box(shape)
            if box is None or box[2] - box[0] >= width * 0.9 or box[3] - box[1] >= height * 0.9:
                continue
            if (shape.has_text_frame and shape.text.strip()) or shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                boxes.append(box)
    for master in template.slide_masters:
        for layout in master.slide_layouts:
            boxes += [_box(item) for item in layout.placeholders if item.placeholder_format.type in CONTENT_PLACEHOLDERS]
    boxes = [box for box in boxes if box is not None]
    if not boxes:
        return round(width * 0.04), round(height * 0.04), round(width * 0.96), round(height * 0.96)
    return (max(0, min(box[0] for box in boxes)), max(0, min(box[1] for box in boxes)),
            min(width, max(box[2] for box in boxes)), min(height, max(box[3] for box in boxes)))


# --- The audit ---------------------------------------------------------------------------

def design_issues(
    pptx_path: Path,
    template_path: Path,
    slides: list[SlideContent],
    prepared: PreparedTemplate | None = None,
    source_slides: list[int] | None = None,
    extra_colors: dict[str, str] | None = None,
) -> list[dict[str, Any]]:
    """Palette, contrast, guides, margins, pictures, brand, fill and chart axes of a generated deck.

    source_slides holds the 1-based template slide each generated slide was built from; without it
    the brand check is skipped and every distorted picture is reported.
    """
    deck = Presentation(str(pptx_path))
    template = Presentation(str(template_path))
    width, height = int(deck.slide_width), int(deck.slide_height)
    palette = template_palette(template)
    area = content_area(template)
    # Variant B's dark palette is designed for the deck and approved with it.
    palette.update(color for color in map(_parse_hex, map(str, (extra_colors or {}).values())) if color)
    generated_lines = {" ".join(text.casefold().split()) for item in slides for text in [item.title, *item.bullets]}
    compositions = {item.source_slide_index: item for item in (prepared.compositions if prepared else [])}
    issues: list[dict[str, Any]] = []
    for index, slide in enumerate(deck.slides):
        spec = slides[index] if index < len(slides) else None
        source = None
        if source_slides and index < len(source_slides) and 1 <= source_slides[index] <= len(template.slides):
            source = template.slides[source_slides[index] - 1]
        theme = _Theme(slide.slide_layout.slide_master)

        def generated(shape) -> bool:
            if not shape.has_text_frame:
                return False
            lines = [" ".join(line.casefold().split()) for line in re.split(r"[\n\v]", shape.text) if line.strip()]
            return bool(lines) and all(line in generated_lines for line in lines)

        blocks = _blocks(slide, generated)
        issues += _palette_issues(index, slide, theme, palette)
        issues += _contrast_issues(index, slide, theme, generated)
        issues += _guide_issues(index, slide, deck, blocks, width)
        issues += _margin_issues(index, blocks, area)
        issues += _picture_issues(index, slide, source)
        if source is not None:
            issues += _brand_issues(index, slide, source, compositions.get(source_slides[index] - 1))
        if spec is not None and any(item.strip() for item in spec.bullets):
            issues += _fill_issues(index, slide, theme, width, height)
        issues += _chart_axis_issues(index, slide)
    return issues


def _blocks(slide, generated) -> list[Block]:
    blocks, nodes = [], []
    for shape in slide.shapes:
        box = _box(shape)
        if box is None:
            continue
        if shape.name.startswith("Infographic node"):
            nodes.append((shape, box))
        elif generated(shape) or getattr(shape, "has_chart", False) or getattr(shape, "has_table", False):
            blocks.append(Block(shape.shape_id, shape.name, box))
    if nodes:
        # A diagram's nodes follow their own rhythm; the diagram as a whole sits on the grid.
        boxes = [box for _, box in nodes]
        blocks.append(Block(nodes[0][0].shape_id, "Инфографика", (min(b[0] for b in boxes), min(b[1] for b in boxes),
                                                                   max(b[2] for b in boxes), max(b[3] for b in boxes))))
    return blocks


def _palette_issues(index, slide, theme: _Theme, palette: set[Rgb]) -> list[dict[str, Any]]:
    elements = list(slide._element.iter(qn("a:srgbClr")))
    for shape in slide.shapes:
        if shape.has_chart:
            elements += list(shape.chart._chartSpace.iter(qn("a:srgbClr")))
    foreign: list[Rgb] = []
    for element in elements:
        color = theme.color(element)
        if color is None or _neutral(color) or color in foreign:
            continue
        if not any(_distance(color, known) <= COLOR_TOLERANCE for known in palette):
            foreign.append(color)
    if not foreign:
        return []
    listed = ", ".join(_hex(color) for color in foreign[:4]) + (" и др." if len(foreign) > 4 else "")
    return [_issue(index, None, "off_palette_color", "warning",
                   f"Цвета не из палитры шаблона: {listed}.", "manual")]


def _text_color(shape, paragraph, run, theme: _Theme) -> Rgb | None:
    from .fonts import _style_chain

    candidates = [run._r.rPr]
    ppr = paragraph._p.pPr
    candidates.append(ppr.find(qn("a:defRPr")) if ppr is not None else None)
    # List styles are numbered from one (a:lvl1pPr) while paragraph levels start at zero.
    chain = [item.find(qn("a:defRPr")) for item in _style_chain(shape, paragraph.level + 1)]
    style = shape._element.find(qn("p:style"))
    reference = style.find(qn("a:fontRef")) if style is not None else None
    # A shape's own list style wins over its theme style; masters and defaults come after both.
    candidates += chain[:1]
    for props in candidates:
        kind, color = theme.fill(props)
        if kind == "solid":
            return color
    if reference is not None and len(reference):
        return theme.color(reference[0])
    for props in chain[1:]:
        kind, color = theme.fill(props)
        if kind == "solid":
            return color
    return theme.scheme.get(theme.map.get("tx1", "dk1"))


def _background(slide, theme: _Theme) -> Rgb | None:
    for owner in (slide, slide.slide_layout, slide.slide_layout.slide_master):
        background = owner._element.cSld.find(qn("p:bg"))
        if background is None:
            continue
        properties = background.find(qn("p:bgPr"))
        if properties is not None:
            kind, color = theme.fill(properties)
            return color if kind == "solid" else None
        reference = background.find(qn("p:bgRef"))
        if reference is not None:
            kind, color = theme.styled_fill(reference, theme.background_fills)
            return color if kind == "solid" else None
    return theme.scheme.get(theme.map.get("bg1", "lt1"), (255, 255, 255))


def _shape_fill(shape, theme: _Theme) -> tuple[str, Rgb | None]:
    kind, color = theme.fill(shape._element.find(qn("p:spPr")))
    if kind is not None:
        return kind, color
    style = shape._element.find(qn("p:style"))
    reference = style.find(qn("a:fillRef")) if style is not None else None
    if reference is not None:
        return theme.styled_fill(reference, theme.fills)
    return "none", None


def _surface(slide, shape, theme: _Theme) -> Rgb | None:
    """The solid colour right behind a text box, or None when a picture or gradient makes it unknown."""
    surface = _background(slide, theme)
    layers = [item for owner in (slide.slide_layout.slide_master, slide.slide_layout)
              for item in owner.shapes if not item.is_placeholder]
    for item in slide.shapes:
        if item.shape_id == shape.shape_id:
            break
        layers.append(item)
    for cover in layers:
        share = _covered(cover, shape)
        if share < 0.3:
            continue
        if cover.shape_type in {MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.GROUP, MSO_SHAPE_TYPE.MEDIA} \
                or getattr(cover, "has_chart", False) or getattr(cover, "has_table", False):
            surface = None
            continue
        kind, color = _shape_fill(cover, theme)
        if kind == "solid" and share >= 0.9:
            surface = color
        elif kind in {"solid", "complex"}:
            surface = None
    kind, color = _shape_fill(shape, theme)
    if kind == "solid":
        return color
    return None if kind == "complex" else surface


def _contrast_issues(index, slide, theme: _Theme, generated) -> list[dict[str, Any]]:
    from .fonts import effective_font

    issues = []
    for shape in slide.shapes:
        if not generated(shape):
            continue
        surface = _surface(slide, shape, theme)
        if surface is None:
            continue
        worst = None
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if not run.text.strip():
                    continue
                color = _text_color(shape, paragraph, run, theme)
                if color is None:
                    continue
                _, size, bold = effective_font(shape, paragraph, run)
                large = size >= 18 or (bold and size >= 14)
                ratio = contrast(color, surface)
                needed = LARGE_TEXT_CONTRAST if large else MIN_CONTRAST
                if ratio < needed and (worst is None or ratio < worst[0]):
                    worst = ratio, needed, color
        if worst is not None:
            ratio, needed, color = worst
            # Hardly readable text blocks the variant; a near miss is left to the author.
            severity = "blocking" if ratio < 2.0 or (needed == MIN_CONTRAST and ratio < 3.0) else "warning"
            issues.append(_issue(index, shape.shape_id, "low_contrast", severity,
                                 f"Контраст текста {ratio:.1f}:1 (нужно не ниже {needed:g}:1): "
                                 f"{_hex(color)} на {_hex(surface)}.", "manual"))
    return issues


def _guide_issues(index, slide, deck, blocks, width: int) -> list[dict[str, Any]]:
    layout = slide.slide_layout
    references = _guides(deck, layout)
    for item in layout.placeholders:
        box = _box(item)
        if box is not None:
            references += [box[0], box[2]]
    generated_ids = {block.shape_id for block in blocks}
    for shape in slide.shapes:
        box = _box(shape)
        # Template objects kept on the slide are part of its grid.
        if (box is not None and shape.shape_id not in generated_ids and not shape.name.startswith("Infographic")
                and box[2] - box[0] < width * 0.9):
            references += [box[0], box[2]]
    slide_area = width * int(slide.part.package.presentation_part.presentation.slide_height)
    containers = [box for shape in slide.shapes if shape.shape_id not in generated_ids
                  and not shape.name.startswith("Infographic") and (box := _box(shape)) is not None
                  and (box[2] - box[0]) * (box[3] - box[1]) < slide_area * 0.9]
    issues = []
    for block in blocks:
        left, top, right, bottom = block.box
        # Text inside a template card follows the card, which already sits on the template's grid.
        if any(box[0] <= left and box[1] <= top and right <= box[2] and bottom <= box[3] for box in containers):
            continue
        own = references + [edge for other in blocks if other.shape_id != block.shape_id
                            for edge in (other.box[0], other.box[2])]
        centered = abs((left + right) / 2 - width / 2) <= GUIDE_TOLERANCE
        if centered or any(abs(left - x) <= GUIDE_TOLERANCE or abs(right - x) <= GUIDE_TOLERANCE for x in own):
            continue
        nearest = min(own, key=lambda x: abs(left - x), default=None)
        detail = f" Ближайшая направляющая в {_cm(left - nearest)} от левого края блока." if nearest is not None else ""
        issues.append(_issue(index, block.shape_id, "misaligned_block", "warning",
                             f"Блок «{block.name}» не выровнен по направляющим макета.{detail}", "manual"))
    return issues


def _margin_issues(index, blocks, area: tuple[int, int, int, int]) -> list[dict[str, Any]]:
    left, top, right, bottom = area
    issues = []
    for block in blocks:
        box = block.box
        overshoot = {"левого": left - box[0], "верхнего": top - box[1],
                     "правого": box[2] - right, "нижнего": box[3] - bottom}
        side, amount = max(overshoot.items(), key=lambda item: item[1])
        if amount > MARGIN_TOLERANCE:
            issues.append(_issue(index, block.shape_id, "margin_intrusion", "warning",
                                 f"Блок «{block.name}» заходит в поле у {side} края на {_cm(amount)}.", "manual"))
    return issues


def _picture_issues(index, slide, source) -> list[dict[str, Any]]:
    originals = {item.shape_id: _box(item) for item in source.shapes} if source is not None else {}
    issues = []
    for shape in slide.shapes:
        if shape.shape_type != MSO_SHAPE_TYPE.PICTURE or not shape.width or not shape.height:
            continue
        # The template's own picture geometry is its design, even when it is stretched there.
        if originals.get(shape.shape_id) == _box(shape):
            continue
        try:
            pixels_wide, pixels_high = shape.image.size
        except (AttributeError, KeyError, ValueError, OSError):
            continue
        crop = 1 - shape.crop_left - shape.crop_right, 1 - shape.crop_top - shape.crop_bottom
        if pixels_wide <= 0 or pixels_high <= 0 or crop[0] <= 0 or crop[1] <= 0:
            continue
        native = pixels_wide * crop[0] / (pixels_high * crop[1])
        shown = shape.width / shape.height
        distortion = shown / native - 1
        if abs(distortion) > DISTORTION:
            issues.append(_issue(index, shape.shape_id, "image_distorted", "blocking",
                                 f"Изображение «{shape.name}» {'растянуто' if distortion > 0 else 'сжато'} "
                                 f"по ширине на {abs(distortion) * 100:.0f}%: пропорции нарушены.", "automatic"))
    return issues


def _brand_issues(index, slide, source, composition) -> list[dict[str, Any]]:
    fixed = {int(key) for key, role in (composition.object_roles if composition else {}).items()
             if role == "fixed_brand" and str(key).isdigit()}
    originals = {item.shape_id: item for item in source.shapes}
    layout_footers = {item.placeholder_format.type: item for item in slide.slide_layout.placeholders
                      if item.placeholder_format.type in FOOTER_PLACEHOLDERS}
    issues = []
    for shape in slide.shapes:
        footer = shape.is_placeholder and shape.placeholder_format.type in FOOTER_PLACEHOLDERS
        if shape.shape_id not in fixed and not footer:
            continue
        reference = originals.get(shape.shape_id)
        if reference is None and footer:
            reference = layout_footers.get(shape.placeholder_format.type)
        if reference is None or _box(reference) is None or _box(shape) is None:
            continue
        shift = max(abs(a - b) for a, b in zip(_box(shape), _box(reference)))
        if shift > BRAND_TOLERANCE:
            kind = "Колонтитул" if footer else "Фирменный элемент"
            issues.append(_issue(index, shape.shape_id, "brand_moved", "warning",
                                 f"{kind} «{shape.name}» сдвинут относительно шаблона на {_cm(shift)}.", "automatic"))
    return issues


def _visible(shape, theme: _Theme) -> bool:
    """A drawn shape: filled, or outlined by its own line or its theme style."""
    if _shape_fill(shape, theme)[0] in {"solid", "complex"}:
        return True
    line = shape._element.find(qn("p:spPr") + "/" + qn("a:ln"))
    if line is not None:
        return line.find(qn("a:noFill")) is None
    style = shape._element.find(qn("p:style"))
    reference = style.find(qn("a:lnRef")) if style is not None else None
    return reference is not None and int(reference.get("idx", "0")) > 0


def _fill_issues(index, slide, theme: _Theme, width: int, height: int) -> list[dict[str, Any]]:
    """Share of the slide covered by content: text as tall as its lines, cards, charts, tables and pictures."""
    from .fonts import text_block_height

    columns, rows = 64, 36
    cells = [[False] * columns for _ in range(rows)]
    slide_area = width * height
    # Artwork of the layout and master fills the slide as much as the slide's own pictures.
    artwork = [item for owner in (slide.slide_layout.slide_master, slide.slide_layout) for item in owner.shapes
               if not item.is_placeholder and item.shape_type == MSO_SHAPE_TYPE.PICTURE]
    for shape in [*artwork, *slide.shapes]:
        box = _box(shape)
        if box is None or (box[2] - box[0]) * (box[3] - box[1]) >= slide_area * 0.9:
            continue
        left, top, right, bottom = box
        if shape.has_text_frame and shape.text.strip():
            try:
                lines = text_block_height(shape) * 12700
            except (AttributeError, ValueError, TypeError):
                lines = bottom - top
            frame = shape.text_frame
            bottom = min(bottom, top + int(lines + (frame.margin_top or 0) + (frame.margin_bottom or 0)))
        elif not (shape.shape_type in {MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.GROUP, MSO_SHAPE_TYPE.CHART}
                  or getattr(shape, "has_table", False) or getattr(shape, "has_chart", False)
                  or (shape.shape_type in {MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM} and _visible(shape, theme))):
            continue
        for row in range(rows):
            y = (row + 0.5) * height / rows
            if not top <= y <= bottom:
                continue
            for column in range(columns):
                if left <= (column + 0.5) * width / columns <= right:
                    cells[row][column] = True
    share = sum(map(sum, cells)) / (columns * rows)
    if share < MIN_FILL:
        return [_issue(index, None, "slide_underfilled", "warning",
                       f"Контент занимает {share * 100:.0f}% слайда (норма от 25%).", "manual")]
    if share > MAX_FILL:
        return [_issue(index, None, "slide_overfilled", "warning",
                       f"Контент занимает {share * 100:.0f}% слайда (норма до 75%).", "manual")]
    return []


def _chart_axis_issues(index, slide) -> list[dict[str, Any]]:
    issues = []
    for shape in slide.shapes:
        if not getattr(shape, "has_chart", False):
            continue
        chart = shape.chart
        try:
            if chart.chart_type in AXIS_FREE_CHARTS:
                continue
            category, value = chart.category_axis, chart.value_axis
        except (ValueError, NotImplementedError, KeyError):
            continue
        problems = []

        def hidden(axis) -> bool:
            deleted = axis._element.find(qn("c:delete"))
            return (deleted is not None and deleted.get("val") in {"1", "true"}) \
                or axis.tick_label_position == XL_TICK_LABEL_POSITION.NONE

        if hidden(category):
            problems.append("подписи категорий скрыты")
        labels = any(plot.has_data_labels for plot in chart.plots)
        if hidden(value) and not labels:
            problems.append("значения не подписаны: ось скрыта и подписей данных нет")
        elif not hidden(value) and not value.has_title:
            problems.append("у оси значений нет подписи с единицами")
        if problems:
            issues.append(_issue(index, shape.shape_id, "chart_axis_unlabeled", "warning",
                                 "Диаграмма: " + "; ".join(problems) + ".", "manual"))
    return issues
