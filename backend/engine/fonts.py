"""Template typefaces: exact resolution, a local font store and PPTX embedding.

Generated text must look exactly like the template's text, on any computer.
The family and size a run really gets are resolved the way PowerPoint does
(run, paragraph, shape list style, layout and master placeholders, master text
styles, theme fonts). Fonts the template names are taken from the system or
the application's font store; an open font the store lacks is fetched once
from Google Fonts. Every generated PPTX embeds the fonts it uses, so PowerPoint
and LibreOffice draw them without installing anything.
"""

from __future__ import annotations

import io
import logging
import os
import re
import struct
import urllib.parse
import urllib.request
from functools import lru_cache
from pathlib import Path

from lxml import etree
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn

from .typography import _key, font_dir, font_files, refresh_fonts

logger = logging.getLogger(__name__)

# Shipped with Windows, macOS and Office: every viewer already has them.
CORE_FONTS = {_key(name) for name in (
    "Arial", "Arial Black", "Arial Narrow", "Calibri", "Calibri Light", "Cambria", "Candara", "Century Gothic",
    "Comic Sans MS", "Consolas", "Constantia", "Corbel", "Courier New", "Franklin Gothic Medium", "Garamond",
    "Georgia", "Impact", "Lucida Console", "Lucida Sans Unicode", "Palatino Linotype", "Segoe UI",
    "Segoe UI Light", "Segoe UI Semibold", "Tahoma", "Times New Roman", "Trebuchet MS", "Verdana", "Aptos",
    "Aptos Display", "Symbol", "Wingdings", "Webdings",
)}
TITLE_PLACEHOLDERS = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
BODY_PLACEHOLDERS = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.OBJECT,
                     PP_PLACEHOLDER.VERTICAL_BODY, PP_PLACEHOLDER.VERTICAL_OBJECT}
WEIGHTS = {"thin": 100, "extralight": 200, "ultralight": 200, "light": 300, "regular": 400, "medium": 500,
           "semibold": 600, "demibold": 600, "bold": 700, "extrabold": 800, "ultrabold": 800,
           "black": 900, "heavy": 900}
STYLE_SLOTS = ("regular", "bold", "italic", "boldItalic")


# --- What PowerPoint actually draws -------------------------------------------------------------

@lru_cache(maxsize=32)
def _theme_fonts_from_xml(theme_xml: bytes) -> tuple[str | None, str | None]:
    scheme = etree.fromstring(theme_xml).find(".//" + qn("a:fontScheme"))
    fonts = []
    for kind in ("a:majorFont", "a:minorFont"):
        latin = scheme.find(qn(kind) + "/" + qn("a:latin")) if scheme is not None else None
        fonts.append((latin.get("typeface") or None) if latin is not None else None)
    return fonts[0], fonts[1]


def theme_fonts(master) -> tuple[str | None, str | None]:
    """(heading, body) typefaces of a slide master's theme."""
    try:
        return _theme_fonts_from_xml(master.part.part_related_by(RT.THEME).blob)
    except (AttributeError, KeyError, ValueError, etree.XMLSyntaxError):
        return None, None


def _master(shape):
    try:
        return shape.part.slide.slide_layout.slide_master
    except AttributeError:
        return None


def _placeholder_type(shape):
    try:
        return shape.placeholder_format.type if shape.is_placeholder else None
    except (AttributeError, ValueError):
        return None


def _level_props(container, level: int):
    """lvlNpPr of an a:lstStyle-like element (lstStyle, titleStyle, bodyStyle...)."""
    if container is None:
        return None
    return container.find(qn(f"a:lvl{level}pPr"))


def _style_chain(shape, level: int) -> list:
    """Paragraph property elements a run inherits from, nearest first."""
    chain = []
    body = shape._element.find(qn("p:txBody"))
    if body is not None:
        chain.append(_level_props(body.find(qn("a:lstStyle")), level))
    if shape.is_placeholder:
        base = getattr(shape, "_base_placeholder", None)
        while base is not None:
            element = base._element.find(qn("p:txBody"))
            if element is not None:
                chain.append(_level_props(element.find(qn("a:lstStyle")), level))
            base = getattr(base, "_base_placeholder", None)
    master = _master(shape)
    styles = master._element.find(qn("p:txStyles")) if master is not None else None
    kind = _placeholder_type(shape)
    name = ("p:titleStyle" if kind in TITLE_PLACEHOLDERS else "p:bodyStyle" if kind in BODY_PLACEHOLDERS
            else "p:otherStyle")
    if styles is not None:
        chain.append(_level_props(styles.find(qn(name)), level))
    try:
        presentation = shape.part.package.presentation_part._element
        chain.append(_level_props(presentation.find(qn("p:defaultTextStyle")), level))
    except AttributeError:
        pass
    return [item for item in chain if item is not None]


def _latin(rpr) -> str | None:
    latin = rpr.find(qn("a:latin")) if rpr is not None else None
    return (latin.get("typeface") or None) if latin is not None else None


def effective_font(shape, paragraph=None, run=None) -> tuple[str | None, float, bool]:
    """(family, size in pt, bold) a run is drawn with, as PowerPoint resolves it."""
    level = (paragraph.level if paragraph is not None else 0) + 1
    rprs = []
    if run is not None and run._r.rPr is not None:
        rprs.append(run._r.rPr)
    if paragraph is not None and paragraph._p.pPr is not None:
        rprs.append(paragraph._p.pPr.find(qn("a:defRPr")))
    rprs += [props.find(qn("a:defRPr")) for props in _style_chain(shape, level)]
    family = size = bold = None
    for rpr in rprs:
        if rpr is None:
            continue
        family = family or _latin(rpr)
        if size is None and (rpr.get("sz") or "").isdigit():
            size = int(rpr.get("sz")) / 100
        if bold is None and rpr.get("b") in {"0", "1", "true", "false"}:
            bold = rpr.get("b") in {"1", "true"}
    master = _master(shape)
    major, minor = theme_fonts(master) if master is not None else (None, None)
    if family is None or family.startswith("+"):
        heading = (family or "").startswith("+mj") or (family is None and _placeholder_type(shape) in TITLE_PLACEHOLDERS)
        family = major if heading else minor
    return family, size or 18.0, bool(bold)


def paragraph_spacing(shape, paragraph) -> tuple[float, float, float]:
    """(line factor, space before+after in pt, left indent in pt) of a paragraph."""
    level = paragraph.level + 1
    chain = ([paragraph._p.pPr] if paragraph._p.pPr is not None else []) + _style_chain(shape, level)
    factor, extra, indent = 1.0, 0.0, 0.0
    found = set()
    for ppr in chain:
        line = ppr.find(qn("a:lnSpc") + "/" + qn("a:spcPct"))
        if "line" not in found and line is not None and (line.get("val") or "").isdigit():
            factor, found = int(line.get("val")) / 100000, found | {"line"}
        for tag in ("a:spcBef", "a:spcAft"):
            points = ppr.find(qn(tag) + "/" + qn("a:spcPts"))
            if tag not in found and points is not None and (points.get("val") or "").isdigit():
                extra, found = extra + int(points.get("val")) / 100, found | {tag}
        if "marL" not in found and (ppr.get("marL") or "").lstrip("-").isdigit():
            indent, found = max(0, int(ppr.get("marL"))) / 12700, found | {"marL"}
    return factor, extra, indent


def text_block_height(shape, lines: list[str] | None = None, width_points: float | None = None) -> float:
    """Height in pt the shape's paragraphs (or the given lines in its style) need.

    width_points measures the text for another width of the same region.
    """
    from .typography import text_height
    frame = shape.text_frame
    outer = width_points if width_points is not None else shape.width / 12700
    width = max(1.0, outer - (frame.margin_left + frame.margin_right) / 12700)
    paragraphs = [paragraph for paragraph in frame.paragraphs if paragraph.text.strip()] or frame.paragraphs[:1]
    if lines is not None:
        sample = paragraphs[0] if paragraphs else None
        pairs = [(line, sample) for line in lines]
    else:
        pairs = [(paragraph.text, paragraph) for paragraph in paragraphs]
    required = 0.0
    for text, paragraph in pairs:
        run = next((item for item in paragraph.runs if item.text.strip()), None) if paragraph is not None else None
        family, size, bold = effective_font(shape, paragraph, run)
        italic = bool(run is not None and run.font.italic)
        factor, extra, indent = paragraph_spacing(shape, paragraph) if paragraph is not None else (1.0, 0.0, 0.0)
        required += text_height([text], max(20.0, width - indent), size, family or "Arial", bold, italic, factor) + extra
    return required


# --- Font store --------------------------------------------------------------------------------

def _split_weight(family: str) -> tuple[str, int]:
    """'Montserrat SemiBold' -> ('Montserrat', 600); a plain family keeps weight 400."""
    words = family.split()
    if len(words) > 1 and _key(words[-1]) in WEIGHTS:
        return " ".join(words[:-1]), WEIGHTS[_key(words[-1])]
    return family, 400


def _download(url: str, limit: int = 15 * 1024 * 1024) -> bytes:
    # A plain client receives TrueType files (browsers get WOFF2).
    request = urllib.request.Request(url, headers={"User-Agent": "Wget/1.21"})
    with urllib.request.urlopen(request, timeout=20) as response:
        data = response.read(limit + 1)
    if len(data) > limit:
        raise ValueError("Font file is too large")
    return data


def fetch_google_font(family: str) -> list[Path]:
    """Store the regular, bold and italic faces of an open Google Fonts family."""
    base, weight = _split_weight(family)
    bold_weight = 700 if weight < 700 else 900
    folder = font_dir() / re.sub(r"[^\w.-]+", "_", family).strip("_")
    saved = []
    for slot, axes in (("regular", f"wght@{weight}"), ("bold", f"wght@{bold_weight}"),
                       ("italic", f"ital,wght@1,{weight}"), ("boldItalic", f"ital,wght@1,{bold_weight}")):
        url = ("https://fonts.googleapis.com/css2?family=" + urllib.parse.quote(base).replace("%20", "+")
               + ":" + axes)
        try:
            css = _download(url, 256 * 1024).decode("utf-8", "replace")
            source = re.search(r"url\((https://fonts\.gstatic\.com/[^)]+\.ttf)\)", css)
            if source is None:
                continue
            data = _download(source.group(1))
            if data[:4] not in {b"\x00\x01\x00\x00", b"true"}:
                continue
        except (OSError, ValueError) as error:
            # A missing style (no italic) or an unknown family is not an error.
            logger.info("Google Fonts: %s %s unavailable: %s", family, slot, error)
            continue
        folder.mkdir(parents=True, exist_ok=True)
        target = folder / f"{slot}.ttf"
        temporary = target.with_suffix(".tmp")
        temporary.write_bytes(data)
        temporary.replace(target)
        saved.append(target)
    if saved:
        refresh_fonts()
    return saved


def template_families(presentation) -> set[str]:
    """Latin typefaces the template's slides, layouts, masters and themes use."""
    families: set[str] = set()
    parts = [presentation.part, *(slide.part for slide in presentation.slides)]
    for master in presentation.slide_masters:
        parts.append(master.part)
        parts += [layout.part for layout in master.slide_layouts]
        families.update(name for name in theme_fonts(master) if name)
    for part in parts:
        for latin in part._element.iter(qn("a:latin")):
            typeface = latin.get("typeface") or ""
            if typeface and not typeface.startswith("+"):
                families.add(typeface)
    return families


def embedded_families(presentation) -> set[str]:
    font_list = presentation._element.find(qn("p:embeddedFontLst"))
    if font_list is None:
        return set()
    return {font.get("typeface") for font in font_list.iter(qn("p:font")) if font.get("typeface")}


_UNAVAILABLE: set[str] = set()


def ensure_fonts(families, *, download: bool | None = None) -> dict[str, list[str]]:
    """Make every family measurable and embeddable; fetch open fonts the store lacks."""
    if download is None:
        download = os.getenv("AYA_FONT_DOWNLOAD", "1") != "0"
    report: dict[str, list[str]] = {"available": [], "downloaded": [], "missing": []}
    for family in sorted({name for name in families if name}):
        if font_files(family):
            report["available"].append(family)
        elif download and family not in _UNAVAILABLE and fetch_google_font(family):
            report["downloaded"].append(family)
        else:
            if download:
                # Proprietary fonts (not on Google Fonts) are asked for once per process.
                _UNAVAILABLE.add(family)
            report["missing"].append(family)
    return report


# --- Embedding ---------------------------------------------------------------------------------

def _tables(data: bytes) -> dict[str, bytes]:
    if data[:4] not in {b"\x00\x01\x00\x00", b"true"}:
        raise ValueError("Only TrueType outlines can be embedded")
    count = struct.unpack_from(">H", data, 4)[0]
    tables = {}
    for index in range(count):
        tag, _, offset, length = struct.unpack_from(">4sIII", data, 12 + 16 * index)
        tables[tag.decode("latin-1")] = data[offset:offset + length]
    return tables


def _names(table: bytes) -> dict[int, str]:
    _, count, storage = struct.unpack_from(">HHH", table, 0)
    found: dict[int, tuple[int, str]] = {}
    for index in range(count):
        platform, encoding, language, name_id, length, offset = struct.unpack_from(">HHHHHH", table, 6 + 12 * index)
        raw = table[storage + offset:storage + offset + length]
        if platform == 3 and encoding in {0, 1}:
            text, rank = raw.decode("utf-16-be", "replace"), 0 if language == 0x409 else 1
        elif platform == 1 and encoding == 0:
            text, rank = raw.decode("latin-1"), 2
        else:
            continue
        if name_id not in found or rank < found[name_id][0]:
            found[name_id] = (rank, text)
    return {key: value for key, (_, value) in found.items()}


def eot_from_ttf(data: bytes) -> bytes:
    """Uncompressed Embedded OpenType 2.1: the format PowerPoint uses for .fntdata parts."""
    tables = _tables(data)
    os2, head, names = tables["OS/2"], tables["head"], _names(tables["name"])
    fs_type = struct.unpack_from(">H", os2, 8)[0]
    if fs_type & 0x000F == 0x0002 or fs_type & 0x0200:
        raise PermissionError("The font license forbids embedding")
    weight = struct.unpack_from(">H", os2, 4)[0]
    fs_selection = struct.unpack_from(">H", os2, 62)[0]
    unicode_ranges = struct.unpack_from(">IIII", os2, 42)
    code_pages = struct.unpack_from(">II", os2, 78) if len(os2) >= 86 else (0, 0)
    checksum = struct.unpack_from(">I", head, 8)[0]

    def name(value: str) -> bytes:
        encoded = value.encode("utf-16-le")
        return struct.pack("<HH", 0, len(encoded)) + encoded

    body = (os2[32:42] + struct.pack("<BBIHH", 1, fs_selection & 1, weight, fs_type, 0x504C)
            + struct.pack("<IIII", *unicode_ranges) + struct.pack("<II", *code_pages)
            + struct.pack("<I", checksum) + b"\0" * 16
            + name(names.get(1, "")) + name(names.get(2, "")) + name(names.get(5, "")) + name(names.get(4, ""))
            + struct.pack("<HH", 0, 0))
    header_size = 16 + len(body)
    return struct.pack("<IIII", header_size + len(data), len(data), 0x00020001, 0) + body + data


def _style_files(family: str) -> dict[str, str]:
    files: dict[str, str] = {}
    for path, style in font_files(family):
        if not path.lower().endswith((".ttf", ".otf")):
            continue
        words = set(style.replace("-", " ").split())
        slot = ("boldItalic" if "bold" in words and ({"italic", "oblique"} & words)
                else "bold" if "bold" in words else "italic" if {"italic", "oblique"} & words
                else "regular" if words <= {"regular", "normal", "book", "roman"} else None)
        stored = Path(path).stem if Path(path).parent.parent == font_dir() else None
        slot = stored if stored in STYLE_SLOTS else slot
        if slot and slot not in files:
            files[slot] = path
    return files


_LIST_SUCCESSORS = ("p:custShowLst", "p:photoAlbum", "p:custDataLst", "p:kinsoku", "p:defaultTextStyle",
                    "p:modifyVerifier", "p:extLst")


def embed_fonts(presentation, families) -> dict[str, list[str]]:
    """Embed the fonts a deck uses, so it renders the same on computers without them."""
    report: dict[str, list[str]] = {"embedded": [], "already_embedded": [], "core": [], "not_embeddable": []}
    present = embedded_families(presentation)
    root = presentation._element
    font_list = root.find(qn("p:embeddedFontLst"))
    package = presentation.part.package
    for family in sorted({name for name in families if name}):
        if family in present:
            report["already_embedded"].append(family)
            continue
        if _key(family) in CORE_FONTS:
            report["core"].append(family)
            continue
        files = _style_files(family)
        entry = etree.SubElement(etree.Element(qn("p:embeddedFontLst")), qn("p:embeddedFont"))
        font = etree.SubElement(entry, qn("p:font"))
        font.set("typeface", family)
        for slot in STYLE_SLOTS:
            if slot not in files:
                continue
            try:
                blob = eot_from_ttf(Path(files[slot]).read_bytes())
            except (OSError, KeyError, ValueError, PermissionError, struct.error) as error:
                logger.info("Font %s %s is not embedded: %s", family, slot, error)
                continue
            safe = re.sub(r"[^\w-]+", "", family.replace(" ", "")) or "font"
            partname = PackURI(f"/ppt/fonts/{safe}-{slot}.fntdata")
            part = Part(partname, "application/x-fontdata", package, blob)
            relation = etree.SubElement(entry, qn(f"p:{slot}"))
            relation.set(qn("r:id"), presentation.part.relate_to(part, RT.FONT))
        if len(entry) == 1:
            report["not_embeddable"].append(family)
            continue
        if font_list is None:
            font_list = etree.Element(qn("p:embeddedFontLst"))
            successor = next((root.find(qn(tag)) for tag in _LIST_SUCCESSORS if root.find(qn(tag)) is not None), None)
            if successor is not None:
                successor.addprevious(font_list)
            else:
                root.append(font_list)
        font_list.append(entry)
        report["embedded"].append(family)
    if report["embedded"]:
        root.set("embedTrueTypeFonts", "1")
    return report


def used_families(presentation) -> set[str]:
    """Families of every text run on the deck's slides, plus the theme body font charts use."""
    families: set[str] = set()
    for slide in presentation.slides:
        families.update(name for name in theme_fonts(slide.slide_layout.slide_master) if name)
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.text.strip():
                        family, _, _ = effective_font(shape, paragraph, run)
                        if family:
                            families.add(family)
    return families


def with_embedded_fonts(source: Path, target: Path) -> Path:
    """A copy of a PPTX that carries its fonts, for rendering reference previews."""
    from pptx import Presentation
    deck = Presentation(str(source))
    families = template_families(deck)
    ensure_fonts(families)
    embed_fonts(deck, families)
    buffer = io.BytesIO()
    deck.save(buffer)
    Path(target).write_bytes(buffer.getvalue())
    return Path(target)
