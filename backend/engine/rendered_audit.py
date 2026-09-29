"""Audit actual PDF glyph positions and typefaces against editable PPTX text frames."""
from collections import Counter
from pathlib import Path
import re
import unicodedata
import pdfplumber
from pptx import Presentation
from pptx.enum.dml import MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE
from .audit import _issue
from .fonts import effective_font
from .typography import _key, font_files


def _normalize(text):
    # Letters and digits only: bullets, arrows and other symbols may be drawn by a fallback
    # font and extracted as other characters, which must not hide the paragraph.
    return ''.join(c for c in unicodedata.normalize('NFKC', text).casefold() if c.isalnum())


def _same_family(expected: str, fontname: str) -> bool:
    """'Play' matches the PDF font 'ABCDEF+Play-Regular'."""
    wanted, drawn = _key(expected), _key(fontname.split('+', 1)[-1])
    return bool(wanted) and (drawn.startswith(wanted) or wanted.startswith(drawn))


def _font_substitution(index, shape, chars):
    """Font Inspector, rendered side: generated letters are drawn in the template's typeface."""
    paragraph = next((item for item in shape.text_frame.paragraphs if item.text.strip()), None)
    if paragraph is None:
        return []
    run = next((item for item in paragraph.runs if item.text.strip()), None)
    family = effective_font(shape, paragraph, run)[0]
    letters = [c['fontname'] for c in chars if c['text'].isalpha() and c.get('fontname')]
    foreign = Counter(name.split('+', 1)[-1] for name in letters if not _same_family(family or '', name))
    if not family or not letters or sum(foreign.values()) <= len(letters) * 0.1:
        return []
    # A font the platform has must render; a font nobody supplied is the author's to add.
    available = bool(font_files(family))
    return [_issue(index, shape.shape_id, 'font_substituted', 'blocking' if available else 'warning',
                   f'Шрифт «{family}» при отрисовке заменён на «{foreign.most_common(1)[0][0]}»: '
                   + ('он есть на сервере, но не встроен в презентацию.' if available else
                      'его нет ни в шаблоне, ни в системе, ни в Google Fonts. Встройте шрифт в шаблон '
                      'или положите файл в хранилище шрифтов (AYA_FONT_DIR).'), 'manual')]


def _line_boxes(chars) -> list[tuple[float, float, float, float]]:
    """Boxes of the drawn lines of one text object: letters that share a baseline band."""
    lines: list[list] = []
    for char in sorted(chars, key=lambda item: (item['top'], item['x0'])):
        line = lines[-1] if lines else None
        if line and char['top'] < line[0]['top'] + (line[0]['bottom'] - line[0]['top']) * 0.5:
            line.append(char)
        else:
            lines.append([char])
    return [(min(c['x0'] for c in line), min(c['top'] for c in line),
             max(c['x1'] for c in line), max(c['bottom'] for c in line)) for line in lines]


def text_collisions(first, second, tolerance: float = 1.5) -> list[tuple[float, float, float, float]]:
    """Line boxes where the letters of two objects really overlap.

    Frames, and even the outlines of two paragraphs, may cross in their empty parts
    (an indented or short last line); only drawn lines that intersect are a collision.
    """
    hits = []
    for a in first:
        for b in second:
            if min(a[2], b[2]) - max(a[0], b[0]) > tolerance and min(a[3], b[3]) - max(a[1], b[1]) > tolerance:
                hits.append((max(a[0], b[0]), max(a[1], b[1]), min(a[2], b[2]), min(a[3], b[3])))
    return hits


def _covers(shape) -> bool:
    """A picture, chart, table, group or filled figure hides what lies under it."""
    if shape.shape_type in {MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.GROUP, MSO_SHAPE_TYPE.MEDIA,
                            MSO_SHAPE_TYPE.CHART, MSO_SHAPE_TYPE.TABLE, MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT}:
        return True
    if getattr(shape, "has_chart", False) or getattr(shape, "has_table", False):
        return True
    try:
        return shape.fill.type in {MSO_FILL.SOLID, MSO_FILL.PICTURE, MSO_FILL.GRADIENT, MSO_FILL.PATTERNED}
    except (AttributeError, TypeError, ValueError, NotImplementedError):
        return False


def _field_checker(index, shape, drawn, order, clear, sx, sy):
    """Field Checker: the text stays in its field's free area and nothing is drawn over it."""
    issues = []
    if clear:
        area = (clear[0] * sx, clear[1] * sy, (clear[0] + clear[2]) * sx, (clear[1] + clear[3]) * sy)
        if any(line[0] < area[0] - 2 or line[1] < area[1] - 2 or line[2] > area[2] + 2 or line[3] > area[3] + 2
               for line in drawn):
            issues.append(_issue(index, shape.shape_id, 'field_text_outside_clear_area', 'blocking',
                                 'Текст выходит из свободной области поля на картинку, иконку или декор шаблона.', 'manual'))
    position = next((number for number, item in enumerate(order) if item.shape_id == shape.shape_id), len(order))
    for cover in order[position + 1:]:
        # Objects later in the z-order are drawn over the text.
        if (cover.has_text_frame and cover.text.strip()) or not _covers(cover) or cover.width is None:
            continue
        box = (cover.left * sx, cover.top * sy, (cover.left + cover.width) * sx, (cover.top + cover.height) * sy)
        if text_collisions(drawn, [box], 1.0):
            issues.append(_issue(index, shape.shape_id, 'field_text_covered', 'blocking',
                                 f'Текстовое поле закрыто сверху объектом {cover.shape_id} ({cover.name}).', 'manual'))
    return issues


def audit_pdf(pptx_path: Path, pdf_path: Path, slides=None, fields=None):
    """Letters of the exported PDF against the editable PPTX.

    Every text object is located paragraph by paragraph (bullet glyphs sit between them).
    Its drawn lines must stay in the frame, clear of pictures and of other objects' lines;
    slides limits the typeface check to text the platform wrote. fields maps a slide index to
    the Field Marker's free areas ({shape_id: [x, y, width, height]}) for the Field Checker.
    """
    deck = Presentation(str(pptx_path))
    generated = None if slides is None else {
        ' '.join(text.casefold().split()) for item in slides for text in [item.title, *item.bullets]}
    issues = []
    with pdfplumber.open(pdf_path) as pdf:
        if len(pdf.pages) != len(deck.slides):
            raise ValueError('PDF page count does not match the presentation')
        for index, (slide, page) in enumerate(zip(deck.slides, pdf.pages)):
            glyphs, normalized = [], []
            for char in page.chars:
                for value in _normalize(char['text']):
                    normalized.append(value)
                    glyphs.append(char)
            stream = ''.join(normalized)
            sx, sy = page.width / deck.slide_width, page.height / deck.slide_height
            found = []
            claimed = set()
            for shape in slide.shapes:
                if not shape.has_text_frame or not shape.text.strip() or shape.rotation:
                    continue
                expected = (shape.left*sx, shape.top*sy, (shape.left+shape.width)*sx, (shape.top+shape.height)*sy)
                chars, anchor, missing = [], (expected[0], expected[1]), False
                for paragraph in shape.text_frame.paragraphs:
                    text = _normalize(paragraph.text)
                    if not text:
                        continue
                    candidates = []
                    cursor = stream.find(text)
                    while cursor >= 0:
                        if cursor not in claimed:
                            candidates.append(cursor)
                        cursor = stream.find(text, cursor + 1)
                    start = min(candidates, key=lambda pos: abs(glyphs[pos]['x0']-anchor[0])
                                + abs(glyphs[pos]['top']-anchor[1]), default=-1)
                    if start < 0:
                        missing = True
                        break
                    claimed.add(start)
                    span = glyphs[start:start+len(text)]
                    chars += span
                    anchor = (expected[0], max(c['bottom'] for c in span))
                if missing or not chars:
                    issues.append(_issue(index, shape.shape_id, 'rendered_text_missing', 'blocking',
                                         'Текст PPTX не найден целиком в экспортированном PDF: ' + shape.text[:100], 'manual'))
                    continue
                box = (min(c['x0'] for c in chars), min(c['top'] for c in chars),
                       max(c['x1'] for c in chars), max(c['bottom'] for c in chars))
                if any((expected[0]-box[0] > 2.5, expected[1]-box[1] > 2.5,
                        box[2]-expected[2] > 2.5, box[3]-expected[3] > 2.5)):
                    issues.append(_issue(index, shape.shape_id, 'rendered_text_overflow', 'blocking',
                                         f'Глифы PDF {tuple(round(v,1) for v in box)} выходят за рамку {tuple(round(v,1) for v in expected)}.', 'manual'))
                lines = [' '.join(line.casefold().split()) for line in re.split(r'[\n\v]', shape.text) if line.strip()]
                if generated is None or (lines and all(line in generated for line in lines)):
                    issues += _font_substitution(index, shape, chars)
                drawn = _line_boxes(chars)
                found.append((shape, drawn))
                issues += _field_checker(index, shape, drawn, list(slide.shapes),
                                         (fields or {}).get(index, {}).get(shape.shape_id), sx, sy)
                for picture in slide.shapes:
                    if picture.shape_type not in {MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.GROUP}:
                        continue
                    if picture.width * picture.height >= deck.slide_width * deck.slide_height * .55:
                        continue
                    artwork = (picture.left*sx, picture.top*sy, (picture.left+picture.width)*sx, (picture.top+picture.height)*sy)
                    if text_collisions(drawn, [artwork], 2.5):
                        issues.append(_issue(index, shape.shape_id, 'rendered_text_image_overlap', 'blocking',
                                             f'Отрисованный текст пересекает изображение {picture.shape_id}.', 'manual'))
            for i, (first, a) in enumerate(found):
                for second, b in found[i+1:]:
                    hits = text_collisions(a, b)
                    if hits:
                        where = hits[0]
                        issues.append(_issue(index, first.shape_id, 'rendered_text_overlap', 'blocking',
                                             f'Буквы объектов {first.shape_id} и {second.shape_id} накладываются друг на друга в PDF '
                                             f'(строки в области {tuple(round(v, 1) for v in where)}).', 'manual'))
            page.close()
    return issues
