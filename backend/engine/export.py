"""PDF/PNG/standalone HTML export from an immutable PPTX version."""

from __future__ import annotations

import base64
import html
import os
import shutil
import subprocess
import tempfile
import time
import sys
from threading import Lock
from pathlib import Path

from pptx import Presentation
from pypdf import PdfReader


def _find_soffice() -> str:
    candidate = shutil.which("soffice.com") or shutil.which("soffice") or shutil.which("soffice.exe")
    if candidate:
        return candidate
    windows = Path("C:/Program Files/LibreOffice/program/soffice.com")
    if windows.exists():
        return str(windows)
    raise RuntimeError("LibreOffice is required for PDF export")


def _find_pdftoppm() -> str:
    candidate = shutil.which("pdftoppm") or shutil.which("pdftoppm.exe")
    if candidate:
        return candidate
    bundled = Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies/native/poppler/Library/bin/pdftoppm.exe"
    if bundled.exists():
        return str(bundled)
    raise RuntimeError("Poppler pdftoppm is required for PNG previews")


def _run_converter(command: list[str], cwd: Path, timeout: float):
    process = subprocess.Popen(command, cwd=cwd, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                               text=True, encoding="utf-8", errors="replace",
                               start_new_session=os.name != "nt")
    try:
        stdout, stderr = process.communicate(timeout=timeout)
    except subprocess.TimeoutExpired:
        # Kill the launcher and its children so they cannot retain profile locks.
        if os.name == "nt":
            subprocess.run(["taskkill", "/PID", str(process.pid), "/T", "/F"],
                           capture_output=True, timeout=10)
        else:
            import signal
            os.killpg(process.pid, signal.SIGKILL)
        process.communicate(timeout=10)
        raise
    return subprocess.CompletedProcess(command, process.returncode, stdout, stderr)


def _valid_pdf(path: Path) -> bool:
    try:
        return path.is_file() and path.stat().st_size > 0 and len(PdfReader(path).pages) > 0
    except Exception:
        return False


_PDF_EXPORT_LOCK = Lock()


def export_pdf(pptx_path: Path, output_dir: Path, *, timeout: float = 120) -> Path:
    # Concurrent Windows LibreOffice launches can deadlock during startup even
    # with separate profiles. Include queue wait in the caller's time budget.
    if os.name != "nt":
        return _export_pdf(pptx_path, output_dir, timeout=timeout)
    started = time.monotonic()
    if not _PDF_EXPORT_LOCK.acquire(timeout=max(0, timeout)):
        raise TimeoutError("Истекло время ожидания очереди экспорта PDF.")
    try:
        remaining = timeout - (time.monotonic() - started)
        if remaining <= 0:
            raise TimeoutError("Истекло время экспорта PDF.")
        return _export_pdf(pptx_path, output_dir, timeout=remaining)
    finally:
        _PDF_EXPORT_LOCK.release()


def _export_pdf(pptx_path: Path, output_dir: Path, *, timeout: float = 120) -> Path:
    pptx_path = Path(pptx_path).resolve(strict=True)
    output_dir = Path(output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    executable = _find_soffice()
    deadline = time.monotonic() + timeout
    diagnostics = "Conversion time budget exhausted"
    for attempt in range(2):
        with tempfile.TemporaryDirectory(prefix="aiya-lo-", ignore_cleanup_errors=True) as scratch:
            work = Path(scratch)
            profile_path = work / "profile"
            user_dir = profile_path / "user"
            user_dir.mkdir(parents=True)
            # Disable the Windows auto-updater in this isolated profile only.
            (user_dir / "registrymodifications.xcu").write_text(
                '<?xml version="1.0" encoding="UTF-8"?>'
                '<oor:items xmlns:oor="http://openoffice.org/2001/registry">'
                '<item oor:path="/org.openoffice.Office.Update/Update">'
                '<prop oor:name="Enabled" oor:op="fuse"><value>false</value></prop>'
                '</item></oor:items>', encoding="utf-8",
            )
            # Simple local names avoid source path/encoding issues. An old output
            # PDF must never be mistaken for the result of this conversion.
            shutil.copyfile(pptx_path, work / "input.pptx")
            remaining = deadline - time.monotonic()
            if remaining <= 0:
                break
            try:
                result = _run_converter([
                    executable, "-env:UserInstallation=" + profile_path.resolve().as_uri(),
                    "--headless", "--norestore", "--nodefault", "--nofirststartwizard",
                    "--convert-to", "pdf:impress_pdf_Export", "--outdir", str(work),
                    str(work / "input.pptx"),
                ], cwd=work, timeout=remaining)
            except subprocess.TimeoutExpired as exc:
                raise RuntimeError("Не удалось экспортировать PDF: LibreOffice превысил время ожидания. Повторите подготовку.") from exc
            candidate = work / "input.pdf"
            if result.returncode == 0 and _valid_pdf(candidate):
                pdf_path = output_dir / (pptx_path.stem + ".pdf")
                # Publish only a complete PDF, atomically on the output volume.
                with tempfile.NamedTemporaryFile(dir=output_dir, suffix=".pdf", delete=False) as staged:
                    staged_path = Path(staged.name)
                try:
                    shutil.copyfile(candidate, staged_path)
                    staged_path.replace(pdf_path)
                finally:
                    staged_path.unlink(missing_ok=True)
                return pdf_path
            diagnostics = (result.stderr or result.stdout or f"exit {result.returncode}; PDF not created")[-1000:]
    raise RuntimeError("LibreOffice PDF export failed: " + diagnostics)


def export_previews(pdf_path: Path, output_dir: Path, dpi: int = 100, *, timeout: float = 120) -> list[Path]:
    output_dir = Path(output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    pdf_path = Path(pdf_path).resolve(strict=True)
    page_count = len(PdfReader(pdf_path).pages)
    deadline = time.monotonic() + timeout
    errors = []
    # A crashed renderer must not publish stale or partial previews. Each
    # attempt has an isolated staging directory and a different rendering engine.
    for engine in ("poppler", "pdfium"):
        remaining = deadline - time.monotonic()
        if remaining <= 0:
            break
        with tempfile.TemporaryDirectory(prefix="preview-", dir=output_dir.parent) as scratch:
            staging = Path(scratch)
            try:
                if engine == "poppler":
                    command = [_find_pdftoppm(), "-png", "-r", str(dpi), str(pdf_path), str(staging / "slide")]
                else:
                    command = [sys.executable, "-m", "engine.preview", str(pdf_path), str(staging), str(dpi)]
                result = subprocess.run(command, capture_output=True, text=True,
                                        encoding="utf-8", errors="replace", timeout=min(remaining, 40) if engine == "poppler" else remaining)
                previews = sorted(staging.glob("slide-*.png"), key=lambda path: int(path.stem.rsplit("-", 1)[1]))
                if result.returncode != 0 or len(previews) != page_count:
                    errors.append(f"{engine}: exit {result.returncode}, {len(previews)}/{page_count} pages; {result.stderr[-300:]}")
                    continue
                from PIL import Image
                for path in previews:
                    with Image.open(path) as image:
                        image.verify()
                # Publish in canonical order only after validating every page.
                published = []
                for index, path in enumerate(previews, 1):
                    target = output_dir / f"slide-{index:02d}.png"
                    path.replace(target)
                    published.append(target)
                return published
            except (OSError, subprocess.TimeoutExpired, RuntimeError) as error:
                errors.append(f"{engine}: {type(error).__name__}: {error}")
    raise RuntimeError("PDF preview export failed: " + "; ".join(errors))


def export_html(pptx_path: Path, previews: list[Path], output_path: Path) -> Path:
    deck = Presentation(str(pptx_path))
    sections = []
    for index, slide in enumerate(deck.slides):
        texts = [shape.text.strip() for shape in slide.shapes if shape.has_text_frame and shape.text.strip()]
        for shape in slide.shapes:
            if shape.has_table:
                texts.extend(
                    " | ".join(shape.table.cell(row, column).text.strip() for column in range(len(shape.table.columns)))
                    for row in range(len(shape.table.rows))
                )
        preview = previews[index] if index < len(previews) else None
        encoded = base64.b64encode(preview.read_bytes()).decode("ascii") if preview else ""
        alt = html.escape(texts[0] if texts else f"Слайд {index + 1}")
        text_block = "".join(f"<li>{html.escape(text)}</li>" for text in texts)
        sections.append(
            f'<section class="slide" id="slide-{index+1}"><header>Слайд {index+1}</header>'
            + (f'<img src="data:image/png;base64,{encoded}" alt="{alt}">' if preview else "")
            + f'<details><summary>Текст слайда</summary><ul>{text_block}</ul></details></section>'
        )
    page = """<!doctype html><html lang="ru"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>Айя — презентация</title>
<style>
:root{color-scheme:light;font-family:system-ui,-apple-system,Segoe UI,sans-serif;background:#f3f5fa;color:#18243d}
*{box-sizing:border-box}body{max-width:1100px;margin:0 auto;padding:24px}
h1{font-size:24px}.slide{background:white;padding:16px;margin:24px 0;border-radius:16px;box-shadow:0 8px 30px #18315b16}
.slide header{font-size:14px;font-weight:700;color:#526079;margin:0 0 10px}.slide img{display:block;width:100%;height:auto;border:1px solid #e0e4ec}
details{margin-top:12px}li{margin:6px 0}
</style></head><body><h1>Айя · презентация</h1>""" + "\n".join(sections) + "</body></html>"
    output_path = Path(output_path)
    output_path.write_text(page, encoding="utf-8")
    return output_path
