from io import BytesIO

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from engine.audit import audit_pptx
from engine.ingest import inspect_template
from engine.models import Fact, SlideContent
from engine.renderer import render_variant


def test_template_palette_brand_and_image_clearance(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 240, 210)
    title = slide.shapes.add_textbox(Inches(.5), Inches(.3), Inches(8), Inches(.7))
    title.text = 'Old heading'
    body = slide.shapes.add_textbox(Inches(.5), Inches(1.3), Inches(8), Inches(3.4))
    body.text = 'Old caption'
    body.text_frame.paragraphs[0].runs[0].font.size = Pt(8)
    brand = slide.shapes.add_textbox(Inches(.5), Inches(6.4), Inches(3), Inches(.4))
    brand.text = 'Brand name'
    brand.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    stream = BytesIO()
    Image.new('RGB', (100, 150), 'red').save(stream, format='PNG')
    photo = slide.shapes.add_picture(stream, Inches(6), Inches(1.5), Inches(3), Inches(4))
    source = tmp_path / 'unknown.pptx'
    deck.save(source)
    template = inspect_template(source)
    template.compositions[0].object_roles.update({str(brand.shape_id): 'fixed_brand', str(photo.shape_id): 'reusable_asset'})
    points = ['The first supported explanation of the product.', 'The second supported explanation of the product.']
    content = [SlideContent('slide_1', 'New heading', points, ['f1'])]
    fact = Fact('f1', ' '.join(points), 'brief', '1', ' '.join(points))
    geometries = []
    for variant in range(3):
        target = tmp_path / f'{variant}.pptx'
        metrics = render_variant(source, template, content, variant, target)
        rendered = Presentation(target).slides[0]
        assert rendered.background.fill.fore_color.rgb == RGBColor(255, 240, 210)
        assert metrics['dark_palette'] is None
        kept = next(s for s in rendered.shapes if s.has_text_frame and s.text == 'Brand name')
        assert (kept.left, kept.top, kept.width, kept.height) == (brand.left, brand.top, brand.width, brand.height)
        bodies = [s for s in rendered.shapes if s.has_text_frame and any(p in s.text for p in points)]
        # Generated points keep the template's own body size, even a small caption size.
        assert all(r.font.size.pt == 8 for s in bodies for p in s.text_frame.paragraphs for r in p.runs)
        assert all(s.left + s.width < photo.left for s in bodies)
        geometries.append([(s.left, s.top, s.width, s.height) for s in bodies])
        assert not any(i['rule_id'] == 'text_image_overlap' for i in audit_pptx(target, content, [fact]))
    assert len({str(g) for g in geometries}) == 3


def test_deterministic_audit_detects_existing_failures(tmp_path):
    deck = Presentation()
    specs = []
    for i in range(2):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(.5), Inches(1), Inches(7), Inches(.3))
        box.text = 'TODO ' + 'Repeated text with a tiny caption font. ' * 5
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(8)
        stream = BytesIO()
        Image.new('RGB', (100, 100), 'red').save(stream, format='PNG')
        slide.shapes.add_picture(stream, Inches(4), Inches(.8), Inches(2), Inches(2))
        specs.append(SlideContent(f'slide_{i+1}', 'Heading', ['Repeated explanation'], ['f1']))
    target = tmp_path / 'bad.pptx'
    deck.save(target)
    issues = audit_pptx(target, specs, [Fact('f1', 'Fact', 'brief', '1', 'Fact')])
    assert {'text_image_overlap', 'small_body_text', 'placeholder_text', 'duplicate_content', 'possible_text_overflow'} <= {i['rule_id'] for i in issues}
    assert issues == audit_pptx(target, specs, [Fact('f1', 'Fact', 'brief', '1', 'Fact')])
