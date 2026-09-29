import io
import json
import time
from concurrent.futures import ThreadPoolExecutor
from threading import Lock

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from engine import pipeline, provider
from engine.models import Fact, SlideContent
from engine.model_registry import APPROVED_MODELS


@pytest.mark.parametrize("brief", [
    "Проект создаёт презентации из материалов.\nКоманда разрабатывает отдельный сервис.\nРезультат сохраняется в редактируемом файле.",
    "ИИ",
    "Проверка предела визуального ремонта",
])
def test_complete_agent_pipeline_on_an_unseen_template(tmp_path, monkeypatch, brief):
    """Real parsing/rendering + fake OpenRouter wire responses, no paid requests."""
    monkeypatch.setenv("OPENROUTER_API_KEY", "synthetic-key")
    monkeypatch.setenv("MODEL_TEMPLATE_CACHE_DIR", str(tmp_path / "cache"))
    monkeypatch.setenv("MODEL_ANALYSIS_RPM", "6000")
    source = tmp_path / "unseen-template.pptx"
    deck = Presentation()
    for label in ("Alpha", "Beta", "Gamma"):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(8), Inches(0.8)).text = label
        slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(8), Inches(4)).text = "Old content"
    deck.save(source)
    pdf_slides = {}
    def export_pdf(pptx, folder, **kwargs):
        folder.mkdir(parents=True, exist_ok=True)
        target = folder / (pptx.stem + ".pdf")
        target.write_bytes(b"synthetic pdf")
        pdf_slides[str(target)] = len(Presentation(pptx).slides)
        return target
    def export_previews(pdf, folder, **kwargs):
        folder.mkdir(parents=True, exist_ok=True)
        result = []
        for i in range(pdf_slides[str(pdf)]):
            target = folder / f"slide-{i+1}.png"
            Image.new("RGB", (100, 60), "white").save(target)
            result.append(target)
        return result
    pdf_audits = {}
    def audit_pdf(pptx, pdf, *args):
        # The first export of every deck shows letters overlapping on slide 3; the relaid one is clean.
        pdf_audits[str(pptx)] = pdf_audits.get(str(pptx), 0) + 1
        if pdf_audits[str(pptx)] > 1:
            return []
        return [{"issue_id": "issue_crowded", "slide_id": "slide_3", "object_id": None, "rule_id": "rendered_text_overlap",
                 "severity": "blocking", "evidence": "synthetic overlap", "repairability": "manual", "selected": False}]
    monkeypatch.setattr(pipeline, "audit_pdf", audit_pdf)
    monkeypatch.setattr(pipeline, "export_pdf", export_pdf)
    monkeypatch.setattr(pipeline, "export_previews", export_previews)
    calls = []
    critique_round = 0
    visual_issue_sent = False
    lock = Lock()
    def openrouter(request, timeout):
        nonlocal critique_round, visual_issue_sent
        assert request.full_url == "https://openrouter.ai/api/v1/chat/completions"
        payload = json.loads(request.data)
        model = payload["model"]
        assert model in APPROVED_MODELS
        messages = payload["messages"]
        assert len(messages) == 2
        system = messages[0]["content"]
        content = messages[1]["content"]
        user = content[0]["text"] if isinstance(content, list) else content
        if "Template Analyst" in system:
            role = "template_analyst"
            assert model == "qwen/qwen3-vl-32b-instruct"
            assert "reasoning" not in payload
            inventory = json.loads(user)
            answer = {"slides": [{"source_slide_index": row["source_slide_index"], "archetype": "content",
                                  "object_roles": {obj["object_id"]: "replaceable" for obj in row["objects"]}}
                                 for row in inventory]}
        elif "Deck Planner" in system:
            role = "deck_planner"
            facts = json.loads(user.split("Fact registry:\n")[1])
            answer = {"slides": [{"title": title, "fact_ids": [fact["fact_id"]]}
                                 for title, fact in zip(("Проект", "Команда", "Результат"), (facts[i % len(facts)] for i in range(3)))]}
        elif "Slide Worker" in system:
            role = "slide_worker"
            task = json.loads(user)
            answer = task["slide"]
            if task["sparse_source"]:
                answer["bullets"] = ["На слайде объясняется тема " + answer["title"].lower() + ".", "Этот раздел связывает тему с исходными сведениями о проекте."]
        elif "Slide Repair agent" in system:
            role = "slide_repair"
            repair_task = json.loads(user)
            if "visual_issues" in repair_task:
                assert brief in {"ИИ", "Проверка предела визуального ремонта"}
                assert isinstance(content, list)
                assert repair_task["visual_issues"][0]["rule_id"] == "visual_text_clipping"
                repair_task["slide"]["bullets"][0] += " Уточнение."
            answer = repair_task["slide"]
            answer["title"] = "Уточнённый проект"
        elif "Infographic Designer" in system:
            role = "infographic_designer"
            candidates = json.loads(user)["candidates"]
            assert candidates
            assert payload["response_format"]["type"] == "json_schema"
            answer = {"slides": [{"slide_id": candidates[0]["slide_id"], "layout": "modules"}]}
        elif "Dark Theme Designer" in system:
            role = "dark_theme_designer"
            assert model == "qwen/qwen3.8-27b"
            assert "Тёмная тема" in json.loads(user)["instruction"]
            assert payload["response_format"]["type"] == "json_schema"
            answer = {"background": "#111827", "surface": "#1F2937", "text": "#F8FAFC", "accent": "#60A5FA"}
        elif "Prompt Compliance Critic" in system:
            role = "prompt_critic"
            assert model == "google/gemma-4-31b-it"
            assert json.loads(user)["brief"] == brief
            answer = {"issues": []}
        elif "Deck Critic" in system:
            role = "deck_critic"
            assert model == "google/gemma-4-31b-it"
            critique_round += 1
            answer = {"issues": [{"slide_id": "slide_1", "rule_id": "unsupported_claim", "severity": "blocking",
                                   "evidence": "Уточните заголовок по источнику"}] if critique_round == 1 else []}
            if brief == "ИИ" and critique_round > 2:
                answer = {"issues": [{"slide_id": "slide_1", "rule_id": "unsupported_claim",
                                      "severity": "blocking", "evidence": "Visual rewrite changed the meaning"}]}
        else:
            role = "visual_critic"
            assert model == "google/gemma-4-31b-it"
            assert isinstance(content, list)
            assert payload["response_format"]["type"] == "json_schema"
            visual_task = json.loads(user)
            with lock:
                trigger = (brief in {"ИИ", "Проверка предела визуального ремонта"}
                           and visual_task["variant_theme"] == "source"
                           and (brief != "Проверка предела визуального ремонта"
                                or visual_task["expected_title"] == "Уточнённый проект")
                           and (brief == "Проверка предела визуального ремонта" or not visual_issue_sent))
                if trigger:
                    visual_issue_sent = True
            answer = {"issues": [{"rule_id": "visual_text_clipping", "severity": "blocking",
                                  "evidence": "Текст не помещается в карточку"}]} if trigger else {"issues": []}
        with lock:
            calls.append(role)
        return io.BytesIO(json.dumps({"choices": [{"message": {"content": json.dumps(answer)}}],
                                       "usage": {"prompt_tokens": 10, "completion_tokens": 20, "cost": 0},
                                       "provider": "synthetic", "id": "test-generation"}).encode())
    monkeypatch.setattr(provider, "_open", openrouter)
    prepared = pipeline.prepare(source, tmp_path / "prepared")
    variants = pipeline.generate(source, brief, 3, tmp_path / "out", prepared_path=prepared)
    assert len(variants) == 3
    for variant in variants:
        assert len(Presentation(variant.pptx_path).slides) == 3
        assert variant.html_path.is_file()
        assert len(variant.preview_paths) == 3
    report = json.loads((tmp_path / "out/run_report.json").read_text(encoding="utf-8"))
    assert set(calls) == set(provider.ROLE_DEFAULTS) - {"dark_theme_designer"}
    visual_attempts = 10 if brief == "Проверка предела визуального ремонта" else 1 if brief == "ИИ" else 0
    assert calls.count("visual_critic") == 9 + visual_attempts
    assert calls.count("slide_worker") == 3
    assert calls.count("slide_repair") == 1 + visual_attempts
    assert calls.count("prompt_critic") == 2 + visual_attempts
    assert calls.count("dark_theme_designer") == 0
    assert calls.count("infographic_designer") == 1
    assert report["infographic_plan"] == {"slide_2": "modules"}
    assert report["variants"][0]["metrics"]["infographic_slides"] == [2]
    assert all(item["metrics"]["relayout_slides"] == [3] for item in report["variants"])
    assert not any(issue["rule_id"] == "rendered_text_overlap" for variant in variants for issue in variant.issues)
    assert report["semantic_repair"]["status"] == "rechecked"
    assert report["variants"][0]["metrics"]["visual_repair_attempts"] == visual_attempts
    if visual_attempts:
        assert report["variants"][0]["visual_repair_history"][0]["issues"][0]["rule_id"] == "visual_text_clipping"
    if brief == "ИИ":
        assert any(issue["evidence"] == "Visual rewrite changed the meaning"
                   for variant in variants for issue in variant.issues)
        assert any(variant.metrics["quality_status"] == "failed_quality_gate" for variant in variants)
    assert report["slides"][0]["title"] == "Уточнённый проект"
    assert report["variant_design"]["b"]["theme"] == "source"
    assert report["variant_design"]["c"] == "alternate_layout_only"
    for variant, variant_report in zip(variants, report["variants"]):
        for expected, slide in zip(variant_report["slides"], Presentation(variant.pptx_path).slides):
            rendered = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
            assert expected["title"] in rendered
            assert all(bullet in rendered for bullet in expected["bullets"])
    assert all(call["host"] == "openrouter.ai" for call in report["model_calls"])
    assert "synthetic-key" not in json.dumps(report)


def test_shared_concurrency_limit_across_different_roles(monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY", "synthetic-key")
    monkeypatch.setenv("MODEL_MAX_PARALLEL", "2")
    active = peak = 0
    lock = Lock()
    def delayed(request, timeout):
        nonlocal active, peak
        with lock:
            active += 1
            peak = max(peak, active)
        time.sleep(0.04)
        with lock:
            active -= 1
        return io.BytesIO(b'{"choices":[{"message":{"content":"{}"}}]}')
    monkeypatch.setattr(provider, "_open", delayed)
    gateway = provider.ModelGateway()
    with ThreadPoolExecutor(max_workers=6) as executor:
        list(executor.map(lambda client: client.complete_json("JSON", "test"), gateway.clients.values()))
    assert peak == 2



def test_routine_template_normalization_is_a_diagnostic():
    assert pipeline._routine_template_note(
        "Template Analyst: слайд 1: роли 2 текстовых или групповых объектов исправлены для удаления старого содержания."
    )
    assert pipeline._routine_template_note(
        "Template Analyst: слайд 2: роли 1 объектов восстановлены безопасно."
    )
    assert not pipeline._routine_template_note(
        "В шаблоне нет слайда с двумя текстовыми областями; текстовые блоки будут созданы программно."
    )

def test_semantic_repair_respects_time_budget_and_preserves_blockers():
    slides = [SlideContent("slide_1", "Title", ["Fact"], ["f1"])]
    issues = [{"slide_id": "slide_1", "severity": "blocking"}]
    result, remaining, report = pipeline._repair_content(None, None, slides, [], issues, time.monotonic() + 20)
    assert result == slides and remaining == issues
    assert report["status"] == "insufficient_time"



def test_semantic_repair_rechecks_repetition_with_a_distinct_fact():
    facts = [
        Fact("f1", "Репозиторий помогает оценить навыки.", "brief.txt", "line 1", "Репозиторий помогает оценить навыки."),
        Fact("f2", "Учебный трек показывает следующие темы.", "brief.txt", "line 2", "Учебный трек показывает следующие темы."),
    ]
    slides = [
        SlideContent("slide_1", "Оценка навыков", ["Репозиторий помогает оценить навыки."], ["f1"]),
        SlideContent("slide_2", "Повтор оценки", ["Репозиторий помогает оценить навыки."], ["f1"]),
    ]
    issues = [{"slide_id": "slide_2", "rule_id": "repetition", "severity": "warning",
               "evidence": "Повторяет предыдущий слайд"}]
    assert pipeline._needs_semantic_repair(issues[0])
    assert not pipeline._needs_semantic_repair({"rule_id": "narrative_gap", "severity": "warning"})

    class Repairer:
        max_parallel = 1
        def complete_json(self, system, user, **kwargs):
            payload = json.loads(user)
            assert payload["allow_ref_change"] is True
            assert {item["fact_id"] for item in payload["facts"]} == {"f1", "f2"}
            assert payload["other_slides"][0]["title"] == "Оценка навыков"
            return {"title": "Учебный трек", "bullets": ["Учебный трек показывает следующие темы."],
                    "fact_ids": ["f2"]}

    class Critic:
        def complete_json(self, system, user, **kwargs):
            deck = json.loads(user)["slides"]
            assert deck[1]["fact_ids"] == ["f2"]
            return {"issues": []}

    revised, remaining, report = pipeline._repair_content(
        Repairer(), Critic(), slides, facts, issues, time.monotonic() + 300,
    )
    assert revised[0] == slides[0]
    assert revised[1].fact_ids == ["f2"]
    assert remaining == []
    assert report["rounds"] == 1
    assert report["status"] == "rechecked"

def test_semantic_repair_rejects_new_facts():
    fact = Fact("f1", "Revenue 12 percent.", "brief", "line 1", "Revenue 12 percent.")
    slides = [SlideContent("slide_1", "Revenue", [fact.text], ["f1"])]
    class Repairer:
        max_parallel = 1
        def complete_json(self, *args, **kwargs):
            return {"title": "Revenue", "bullets": ["Revenue 99 percent."], "fact_ids": ["f1"]}
    with pytest.raises(provider.ModelProviderError):
        pipeline._repair_content(Repairer(), None, slides, [fact],
                                 [{"slide_id": "slide_1", "severity": "blocking"}], time.monotonic() + 300)


def test_one_sentence_brief_gets_grounded_worker_text_after_retries():
    fact = Fact("f1", "Платформа помогает отслеживать пожары на карте.", "brief", "line 1",
                "Платформа помогает отслеживать пожары на карте.")
    brief = fact.text

    class Planner:
        def __init__(self):
            self.calls = 0

        def complete_json(self, system, user, **kwargs):
            self.calls += 1
            title = "Мониторинг в реальном времени" if self.calls == 1 else "Мониторинг пожаров"
            return {"slides": [{"title": title, "fact_ids": ["f1"]}]}

    class Worker:
        def __init__(self):
            self.calls = 0

        def complete_json(self, system, user, **kwargs):
            self.calls += 1
            payload = json.loads(user)
            assert payload["original_brief"] == brief
            assert payload["sparse_source"] is True
            return {"title": payload["slide"]["title"], "fact_ids": ["f1"],
                    "bullets": ["Пожары"] if self.calls == 1 else
                    ["Платформа помогает показывать пожары на карте.", "Карта служит способом представить сведения о пожарах."]}

    planner, worker = Planner(), Worker()
    planned = pipeline._model_plan(planner, [fact], 1, brief)
    result = pipeline._refine_slide(worker, planned[0], {"f1": fact}, [planned[0].title], brief=brief)
    assert planner.calls == worker.calls == 2
    assert result.bullets == ["Платформа помогает показывать пожары на карте.", "Карта служит способом представить сведения о пожарах."]
