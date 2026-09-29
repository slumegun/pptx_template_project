from pathlib import Path
import io
from concurrent.futures import ThreadPoolExecutor
from threading import Lock
import time
import pytest
from pptx import Presentation
from pptx.util import Inches, Pt
from engine.typography import text_height
from engine.renderer import _title_and_body, _clear_text_preserving_style
from engine.audit import audit_pptx
from engine.models import SlideContent, Fact


def test_font_metrics_distinguish_wide_and_narrow_glyphs():
    assert text_height(['Ш' * 30], 100, 20) > text_height(['i' * 30], 100, 20)


def test_eyebrow_does_not_steal_title_role():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shapes = []
    for y, size, text in [(.1, 9, 'РУБРИКА'), (.7, 31, 'Настоящий заголовок'), (2, 18, 'Содержание')]:
        shape = slide.shapes.add_textbox(Inches(.5), Inches(y), Inches(8), Inches(.7))
        shape.text = text
        shape.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
        shapes.append(shape)
    assert _title_and_body(shapes, deck.slide_width, deck.slide_height)[0] is shapes[1]


def test_overlapping_text_is_a_quality_blocker(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    for text in ('Заголовок', 'Наложенный текст'):
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        shape.text = text
    target = tmp_path / 'overlap.pptx'
    deck.save(target)
    issues = audit_pptx(target, [SlideContent('s1', 'Заголовок', ['Наложенный текст'], ['f1'])],
                        [Fact('f1', 'Текст', 'brief', '1', 'Текст')])
    assert any(i['rule_id'] == 'text_text_overlap' and i['severity'] == 'blocking' for i in issues)


def test_windows_exports_are_serialized_and_queue_wait_uses_budget(tmp_path, monkeypatch):
    from engine import export
    if export.os.name != 'nt':
        pytest.skip('Windows-specific LibreOffice startup regression')
    active = peak = 0
    lock = Lock()
    budgets = []
    def convert(source, output, *, timeout):
        nonlocal active, peak
        with lock:
            active += 1
            peak = max(peak, active)
            budgets.append(timeout)
        time.sleep(.03)
        with lock:
            active -= 1
        return output / 'result.pdf'
    monkeypatch.setattr(export, '_export_pdf', convert)
    with ThreadPoolExecutor(max_workers=3) as pool:
        list(pool.map(lambda _: export.export_pdf(Path('source'), tmp_path, timeout=1), range(3)))
    assert peak == 1
    assert min(budgets) < .98


def test_pdf_glyph_overflow_is_detected_independently_of_pptx_estimate(tmp_path, monkeypatch):
    from engine import rendered_audit
    from types import SimpleNamespace
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(.3))
    shape.text = 'Text'
    source = tmp_path / 'test.pptx'
    deck.save(source)
    page = SimpleNamespace(width=720, height=540,
        chars=[{'text':c,'x0':72+i*10,'x1':82+i*10,'top':72,'bottom':110} for i,c in enumerate('Text')], close=lambda:None)
    class PDF:
        pages = [page]
        def __enter__(self): return self
        def __exit__(self,*args): pass
    monkeypatch.setattr(rendered_audit.pdfplumber, 'open', lambda _: PDF())
    assert any(i['rule_id'] == 'rendered_text_overflow' for i in rendered_audit.audit_pdf(source, Path('test.pdf')))

def test_five_step_infographic_keeps_five_editable_nodes(tmp_path):
    from engine.ingest import inspect_template
    from engine.renderer import render_variant
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    heading = slide.shapes.add_textbox(Inches(.5), Inches(.4), Inches(12), Inches(.8))
    heading.text = 'Source title'
    heading.text_frame.paragraphs[0].runs[0].font.size = Pt(31)
    slide.shapes.add_textbox(Inches(.5), Inches(2), Inches(12), Inches(4)).text = 'Source body'
    source, target = tmp_path/'source.pptx', tmp_path/'result.pptx'
    deck.save(source)
    points = ['Нормализация сетки', 'Расчёт признаков', 'Голосование', 'Фильтрация', 'Экспорт']
    render_variant(source, inspect_template(source), [SlideContent('s1','Этапы',points)], 0, target,
                   infographic_layouts={'s1':'flow'})
    slide = Presentation(target).slides[0]
    nodes = [s for s in slide.shapes if s.name.startswith('Infographic node')]
    arrows = [s for s in slide.shapes if s.name.startswith('Infographic connector')]
    assert [s.text for s in nodes] == points
    assert len(arrows) == 4
    assert all(a.left + a.width < b.left for a,b in zip(nodes,nodes[1:]))
