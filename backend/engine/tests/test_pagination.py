from pathlib import Path
import pytest
from pptx import Presentation
from pptx.util import Inches, Pt
from engine.models import SlideContent
from engine.ingest import inspect_template
from engine.pagination import paginate_slides, dense_geometry, body_fits
from engine.renderer import render_variant, choose_compositions, _template_typography
from engine.audit import audit_pptx


@pytest.fixture
def source(tmp_path):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    for title in ("Обложка", "Основной материал", "Итоги"):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        heading = slide.shapes.add_textbox(Inches(.5), Inches(.4), Inches(12), Inches(.8))
        heading.text = title
        heading.text_frame.paragraphs[0].runs[0].font.name = "Arial"
        heading.text_frame.paragraphs[0].runs[0].font.size = Pt(32)
        # Deliberately restrictive body area reproduces the old fatal fallback.
        slide.shapes.add_textbox(Inches(7), Inches(4), Inches(2), Inches(.5)).text = "Старый текст"
    path = tmp_path / "source.pptx"
    deck.save(path)
    return path


@pytest.mark.parametrize("bullets", [
    ["Система сохраняет данные и передаёт результат пользователю. " * 100],
    ["Ш" * 5000],
    [f"Пункт {i}: содержательное объяснение работы системы." for i in range(80)],
    ["Абзац с переносом строки.\nСледующая строка.\n" * 100],
])
def test_long_copy_paginates_losslessly_and_renders_all_variants(source, tmp_path, bullets):
    template = inspect_template(source)
    original = [SlideContent("slide_1", "Подробное описание системы", bullets),
                SlideContent("slide_2", "Итоги", ["Заключительный короткий тезис."])]
    pages = paginate_slides(original, choose_compositions(template, 2), template.width, template.height, "Arial")
    assert len(pages) > 2
    assert ''.join(b for p in pages[:-1] for b in p.bullets) == ''.join(bullets)
    assert pages[-1].title == "Итоги"
    assert pages[-1].source_slide_index == 2
    assert all(p.source_slide_index == 0 for p in pages[:-1])
    assert [p.slide_id for p in pages] == [f"slide_{i+1}" for i in range(len(pages))]
    for variant in range(3):
        target = tmp_path / f"dense-{variant}.pptx"
        metrics = render_variant(source, template, pages, variant, target)
        assert metrics["slide_count"] == len(pages)
        assert len(metrics["source_slides"]) == len(pages)
        issues = audit_pptx(target, pages, [])
        assert not [i for i in issues if i["severity"] == "blocking"]
        deck = Presentation(target)
        for slide in list(deck.slides)[:-1]:
            body = next(s for s in slide.shapes if s.name == "Dense content body")
            assert all(r.font.size.pt >= 16 for p in body.text_frame.paragraphs for r in p.runs)


def test_short_copy_keeps_requested_count_and_source_layout(source):
    template = inspect_template(source)
    original = [SlideContent('slide_1', 'Короткий заголовок', ['Короткий тезис.'])]
    pages = paginate_slides(original, choose_compositions(template, 1), template.width, template.height, 'Arial')
    assert len(pages) == 1
    assert not pages[0].dense_layout
    assert pages[0].source_slide_index == 0


def test_single_dense_page_reclaims_space_instead_of_failing(source, tmp_path):
    template = inspect_template(source)
    content = SlideContent('slide_1', 'Полное описание', ['Подробное описание работы платформы. ' * 30])
    target = tmp_path / 'reflow.pptx'
    render_variant(source, template, [content], 0, target)
    assert not [i for i in audit_pptx(target, [content], []) if i['severity'] == 'blocking']
