"""Strict template mode: exact region data, limits from the template's own text, nothing moves."""
import json
from io import BytesIO
import time
from copy import deepcopy

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from engine import pipeline
from engine.audit import font_issues
from engine.ingest import inspect_template
from engine.models import Fact, SlideContent
from engine.renderer import fit_chars, render_variant, slot_budget, strict_regions


def textbox(slide, left, top, width, height, text, size, color, name=None):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    box.text_frame.word_wrap = True
    box.text_frame.text = text
    run = box.text_frame.paragraphs[0].runs[0]
    run.font.size, run.font.name = Pt(size), "Arial"
    run.font.color.rgb = RGBColor.from_string(color)
    if name:
        box.name = name
    return box


def test_region_data_is_exact_and_colour_is_resolved_through_inheritance(tmp_path):
    deck = Presentation()
    # The title colour lives only in the master's title style.
    level = deck.slide_master._element.find(qn("p:txStyles")).find(qn("p:titleStyle")).find(qn("a:lvl1pPr"))
    props = level.find(qn("a:defRPr"))
    fill = etree.SubElement(props, qn("a:solidFill"))
    etree.SubElement(fill, qn("a:srgbClr")).set("val", "C00000")
    props.insert(0, fill)
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Заголовок слайда"
    box = textbox(slide, 1, 5.5, 6, 0.8, "Короткий пример текста", 20, "123456")
    box.text_frame.paragraphs[0].runs[0].font.name = "Georgia"
    source = tmp_path / "template.pptx"
    deck.save(source)
    slots = {slot.shape_id: slot for slot in inspect_template(source).compositions[0].slots}
    title, sample = slots[slide.shapes.title.shape_id], slots[box.shape_id]
    assert title.color == "#C00000"
    # A bare label is a stub: the region's geometry decides its limit.
    assert title.chars_source == "geometry" and title.max_chars > len("Заголовок слайда")
    assert (sample.x, sample.y, sample.width, sample.height) == (box.left, box.top, box.width, box.height)
    assert (sample.font_family, sample.font_pt, sample.color) == ("Georgia", 20.0, "#123456")
    # A one-line sample allows one full line of the field at the template's size, never less than the sample.
    assert sample.chars_source == "sample_lines" and len("Короткий пример текста") <= sample.max_chars
    assert sample.max_chars <= int((box.width - Inches(0.2)) / 12700 / (20 * 0.55))


def three_point_template(tmp_path):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    textbox(slide, 0.8, 0.5, 11.5, 0.9, "Заголовок — ключевая мысль слайда", 32, "0077FF")
    for number in range(3):
        textbox(slide, 0.8 + number * 4, 2.3, 0.6, 0.6, str(number + 1), 20, "0077FF")
        textbox(slide, 0.8 + number * 4, 3.1, 3.4, 1.2, f"Описание пункта номер {number + 1} примера", 14, "222222")
    source = tmp_path / "template.pptx"
    deck.save(source)
    return source


def test_strict_render_fills_regions_in_place_within_the_template_text_volume(tmp_path):
    source = three_point_template(tmp_path)
    template = inspect_template(source)
    heading, regions = strict_regions(template.compositions[0], template.width, template.height)
    assert heading.text.startswith("Заголовок") and len(regions) == 3
    long_title = "Очень длинный заголовок, который никак не помещается в место заголовка шаблона и должен быть сокращён"
    content = [SlideContent("slide_1", long_title,
                            ["Короткий первый пункт.", "Второй пункт заметно длиннее своего места в макете шаблона."], ["f1"])]
    target = tmp_path / "strict.pptx"
    metrics = render_variant(source, template, content, 0, target, strict=True, plan=[0])
    original = {shape.shape_id: shape for shape in Presentation(source).slides[0].shapes}
    rendered = {shape.shape_id: shape for shape in Presentation(target).slides[0].shapes}
    # Nothing moves or resizes; two points of a row take its outer regions, the middle one
    # leaves with its number.
    for shape_id, shape in rendered.items():
        before = original[shape_id]
        assert (shape.left, shape.top, shape.width, shape.height) == (before.left, before.top, before.width, before.height)
    assert regions[1].shape_id not in rendered and "2" not in {shape.text for shape in rendered.values()}
    assert {"1", "3"} <= {shape.text for shape in rendered.values()}
    for slot in [heading, regions[0], regions[2]]:
        shape = rendered[slot.shape_id]
        assert len(shape.text) <= slot.max_chars
        run, before = shape.text_frame.paragraphs[0].runs[0], original[slot.shape_id].text_frame.paragraphs[0].runs[0]
        assert (run.font.size, run.font.name, run.font.color.rgb) == (before.font.size, before.font.name, before.font.color.rgb)
    written = metrics["written_slides"][0]
    assert written["bullets"][0] == "Короткий первый пункт." and len(written["title"]) <= heading.max_chars
    assert font_issues(target, metrics["font_expectations"]) == []


def test_budget_is_the_smallest_region_of_all_variant_layouts(tmp_path):
    source = three_point_template(tmp_path)
    template = inspect_template(source)
    wide = template.compositions[0]
    narrow = deepcopy(wide)
    narrow.slots[-1].max_chars = 10
    budget = slot_budget([wide, narrow], template.width, template.height)
    _, regions = strict_regions(wide, template.width, template.height)
    assert budget["title"] == regions[0].max_chars or budget["title"] > 0
    assert len(budget["texts"]) == 3 and min(budget["texts"]) == 10


def test_worker_gets_the_limits_retries_once_and_is_cut_to_the_template():
    fact = Fact("f1", "Сервис собирает презентации из брифа и шаблона.", "brief.txt", "1",
                "Сервис собирает презентации из брифа и шаблона.")
    facts = {"f1": fact, **{f"f{i}": Fact(f"f{i}", "x", "brief.txt", "1", "x") for i in range(2, 6)}}
    payloads, schemas = [], []

    class Worker:
        def complete_json(self, system, user, **kwargs):
            payloads.append(json.loads(user))
            schemas.append(kwargs["schema"])
            return {"title": "Сервис собирает презентации",
                    "bullets": ["Сервис собирает презентации из брифа и шаблона компании быстро."], "fact_ids": ["f1"]}

    slots = {"title": 20, "texts": [30], "paragraphs": [1]}
    slide = SlideContent("slide_1", "Сервис", [fact.excerpt], ["f1"])
    refined = pipeline._refine_slide(Worker(), slide, facts, ["Сервис"], time.monotonic() + 60, "", slots)
    assert payloads[0]["slots"] == slots
    assert schemas[0]["properties"]["title"]["maxLength"] == 20
    assert schemas[0]["properties"]["bullets"]["items"]["maxLength"] == 30
    assert "знаков" in payloads[1]["validation_feedback"]
    assert len(refined.title) <= 20 and len(refined.bullets[0]) <= 30


def test_cutting_keeps_whole_words_and_never_exceeds_the_limit():
    text = "Проверки ловят ошибки до показа. Файлы готовы к выступлению."
    assert fit_chars(text, 40) == "Проверки ловят ошибки до показа."
    cut = fit_chars("Сервис собирает презентации из брифа", 25)
    assert len(cut) <= 25 and not cut.endswith(" ") and "презентации из" not in cut


def test_generate_in_strict_mode_keeps_every_region_and_its_text_volume(tmp_path, monkeypatch):
    from PIL import Image
    from engine.ingest import save_prepared
    monkeypatch.setenv("AYA_STRICT_TEMPLATE", "1")
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    cover = deck.slides.add_slide(deck.slide_layouts[6])
    textbox(cover, 0.8, 2.5, 8, 1.2, "Тема презентации в две строки", 40, "0077FF")
    textbox(cover, 0.8, 4.0, 8, 0.6, "Имя и должность докладчика", 16, "333333")
    points = deck.slides.add_slide(deck.slide_layouts[6])
    textbox(points, 0.8, 0.5, 11.5, 0.9, "Заголовок — ключевая мысль слайда", 32, "0077FF")
    for number in range(3):
        textbox(points, 0.8 + number * 4, 3.1, 3.4, 1.2, f"Описание пункта номер {number + 1} примера", 14, "222222")
    closing = deck.slides.add_slide(deck.slide_layouts[6])
    textbox(closing, 0.8, 2.8, 8, 1.2, "Спасибо за внимание!", 48, "0077FF")
    textbox(closing, 0.8, 4.2, 8, 0.6, "Контакты и вопросы для обсуждения", 16, "333333")
    source = tmp_path / "template.pptx"
    deck.save(source)
    template = inspect_template(source)
    for composition, archetype in zip(template.compositions, ("title", "content", "closing")):
        composition.archetype, composition.analysis_complete = archetype, True
    template.analysis_mode = "vision_model"
    prepared = save_prepared(template, tmp_path / "prepared.json")

    long = "Сервис собирает презентации из брифа и шаблона компании и проверяет их перед выдачей пользователю"

    class Client:
        enabled, vision_enabled, text_model, vision_model, max_parallel = True, False, "fake", "fake", 2

        def complete_json(self, system, user, **kwargs):
            task = json.loads(user)
            if "slide" not in task:
                return {}  # the Infographic Designer answers off-contract: no diagrams
            return {"title": task["slide"]["title"] + " — развёрнутая версия заголовка слайда",
                    "bullets": [f"{word} пункт: {long}" for word in ("Первый", "Второй", "Третий")[:len(task["slots"]["texts"]) or 1]],
                    "fact_ids": task["slide"]["fact_ids"]}

    class Gateway:
        enabled, enable_visual_critic, repair_rounds = True, False, 0
        calls: list = []
        clients = {"slide_worker": Client()}

        def client(self, role):
            return Client()

        def configuration_summary(self):
            return {}

    brief = "Сервис собирает презентации из брифа и шаблона компании."
    monkeypatch.setattr(pipeline, "ModelGateway", Gateway)
    monkeypatch.setattr(pipeline, "_model_plan", lambda client, facts, count, brief, deadline: [
        SlideContent("slide_1", "Лукас", [facts[0].excerpt], [facts[0].fact_id]),
        SlideContent("slide_2", "Как работает", [facts[0].excerpt] * 3, [facts[0].fact_id]),
        SlideContent("slide_3", "Спасибо", [facts[0].excerpt], [facts[0].fact_id])])
    monkeypatch.setattr(pipeline, "critique_content", lambda *args, **kwargs: [])
    monkeypatch.setattr(pipeline, "critique_prompt", lambda *args, **kwargs: [])
    monkeypatch.setattr(pipeline, "audit_pdf", lambda *args: [])

    def export_pdf(pptx, folder, **kwargs):
        folder.mkdir(parents=True, exist_ok=True)
        (folder / "deck.pdf").write_bytes(b"pdf")
        (folder / "deck.pdf.count").write_text(str(len(Presentation(pptx).slides)))
        return folder / "deck.pdf"

    def export_previews(pdf, folder, **kwargs):
        folder.mkdir(parents=True, exist_ok=True)
        paths = []
        for number in range(int((pdf.parent / "deck.pdf.count").read_text())):
            Image.new("RGB", (80, 45), "white").save(folder / f"slide-{number + 1}.png")
            paths.append(folder / f"slide-{number + 1}.png")
        return paths

    monkeypatch.setattr(pipeline, "export_pdf", export_pdf)
    monkeypatch.setattr(pipeline, "export_previews", export_previews)
    variants = pipeline.generate(source, brief, 3, tmp_path / "out", prepared_path=prepared)
    report = json.loads((tmp_path / "out" / "run_report.json").read_text(encoding="utf-8"))
    assert report["template_mode"] == "strict" and report["infographic_plan"] == {}
    assert report["layouts"]["a"] == [1, 2, 3] and len(report["slot_budgets"][0][1]["texts"]) == 3
    assert report["variant_design"] == {"a": "base", "b": "more_text", "c": "more_visual"}
    original = [{shape.shape_id: shape for shape in slide.shapes} for slide in Presentation(source).slides]
    limits = {(composition.source_slide_index, slot.shape_id): slot.max_chars
              for composition in template.compositions for slot in composition.slots}
    for variant in variants:
        for index, slide in enumerate(Presentation(variant.pptx_path).slides):
            for shape in slide.shapes:
                before = original[index][shape.shape_id]
                assert (shape.left, shape.top, shape.width, shape.height) == (before.left, before.top, before.width, before.height)
                if shape.has_text_frame and shape.text.strip():
                    assert len(shape.text) <= limits[(index, shape.shape_id)]


def test_a_list_is_a_place_per_paragraph_and_unused_lines_leave_with_their_markers(tmp_path):
    from pptx.enum.shapes import MSO_SHAPE
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    textbox(slide, 0.8, 0.5, 11.5, 0.9, "Заголовок — ключевая мысль слайда", 32, "0077FF")
    points = slide.shapes.add_textbox(Inches(1.2), Inches(2), Inches(8), Inches(3.2))
    points.text_frame.word_wrap = True
    samples = ["Первый пункт списка шаблона", "Второй пункт списка шаблона",
               "Третий пункт списка шаблона", "Четвёртый пункт списка шаблона"]
    for number, sample in enumerate(samples):
        paragraph = points.text_frame.paragraphs[0] if number == 0 else points.text_frame.add_paragraph()
        paragraph.text = sample
        paragraph.space_after = Pt(12)
        paragraph.runs[0].font.size, paragraph.runs[0].font.name = Pt(20), "Arial"
    markers = []
    for number in range(4):
        marker = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.8), Inches(2.1 + number * 0.52), Inches(0.2), Inches(0.2))
        markers.append(marker.shape_id)
    source = tmp_path / "template.pptx"
    deck.save(source)
    template = inspect_template(source)
    heading, regions = strict_regions(template.compositions[0], template.width, template.height)
    assert [region.part for region in regions] == [0, 1, 2, 3]
    content = [SlideContent("slide_1", "Новый заголовок", ["Облачная или локальная модель", "Единая консоль"], ["f1"])]
    target = tmp_path / "list.pptx"
    render_variant(source, template, content, 0, target, strict=True, plan=[0])
    rendered = Presentation(target).slides[0]
    written = next(shape for shape in rendered.shapes if shape.shape_id == points.shape_id)
    assert [paragraph.text for paragraph in written.text_frame.paragraphs] == ["Облачная или локальная модель", "Единая консоль"]
    assert all(paragraph.space_after == Pt(12) for paragraph in written.text_frame.paragraphs)
    left = {shape.shape_id for shape in rendered.shapes}
    assert markers[0] in left and markers[1] in left and markers[2] not in left and markers[3] not in left


def test_themes_choose_layouts_for_more_text_and_for_more_visuals(tmp_path):
    from engine.renderer import choose_compositions
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    cover = deck.slides.add_slide(deck.slide_layouts[6])
    textbox(cover, 0.8, 2.5, 8, 1.2, "Тема презентации", 40, "0077FF")
    wide = deck.slides.add_slide(deck.slide_layouts[6])
    textbox(wide, 0.8, 0.5, 11.5, 0.9, "Заголовок — ключевая мысль слайда", 32, "0077FF")
    textbox(wide, 0.8, 1.8, 5.6, 4.8, "Большой текстовый блок шаблона на несколько строк. " * 6, 16, "222222")
    textbox(wide, 6.9, 1.8, 5.6, 4.8, "Второй большой текстовый блок шаблона на несколько строк. " * 6, 16, "222222")
    icons = deck.slides.add_slide(deck.slide_layouts[6])
    textbox(icons, 0.8, 0.5, 11.5, 0.9, "Заголовок — ключевая мысль слайда", 32, "0077FF")
    for number in range(2):
        icon = icons.shapes.add_picture(BytesIO(png()), Inches(0.8 + number * 6), Inches(2.2), Inches(0.6), Inches(0.6))
        textbox(icons, 0.8 + number * 6, 3.0, 3.4, 0.9, "Тезис под тематической иконкой", 16, "222222")
    closing = deck.slides.add_slide(deck.slide_layouts[6])
    textbox(closing, 0.8, 2.8, 8, 1.2, "Спасибо за внимание!", 48, "0077FF")
    source = tmp_path / "template.pptx"
    deck.save(source)
    template = inspect_template(source)
    for composition, archetype in zip(template.compositions, ("title", "content", "content", "closing")):
        composition.archetype, composition.analysis_complete = archetype, True
    for composition in template.compositions[2:3]:
        composition.object_roles.update({str(shape.shape_id): "reusable_asset" for shape in Presentation(source).slides[2].shapes
                                         if shape.shape_type == 13})
    slides = [SlideContent("slide_1", "Тема", ["Автор"], ["f1"]),
              SlideContent("slide_2", "Суть", ["Первый тезис", "Второй тезис"], ["f1"]),
              SlideContent("slide_3", "Спасибо", ["Вопросы"], ["f1"])]
    text = choose_compositions(template, 3, slides, strict=True, goal="more_text")
    visual = choose_compositions(template, 3, slides, strict=True, goal="more_visual")
    assert text[1].source_slide_index == 1 and visual[1].source_slide_index == 2


def png():
    from PIL import Image
    stream = BytesIO()
    Image.new("RGB", (20, 20), (0, 119, 255)).save(stream, format="PNG")
    return stream.getvalue()
