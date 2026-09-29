from base64 import b64decode
from decimal import Decimal
from io import BytesIO

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.util import Inches, Pt
from engine.ingest import inspect_template, load_prepared, save_prepared
from engine.models import SlideContent
from engine.native_data import NumericCsvData, NumericSeries
from engine.renderer import DARK_PALETTE, render_variant
import pytest


def test_variants_preserve_brand_and_change_composition(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for x, y, w, h, text in [(0.5, 0.3, 8, 0.7, "Old title"), (0.5, 1.4, 8, 4, "Old contents to replace"), (8.5, 0.1, 1, 0.3, "Brand")]:
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        shape.text = text
        shape.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
    logo_id = slide.shapes[-1].shape_id
    source = tmp_path / "unseen.pptx"
    presentation.save(source)
    template = inspect_template(source)
    template.compositions[0].object_roles[str(logo_id)] = "fixed_brand"
    content = [SlideContent("slide_1", "New title", ["First supported statement", "Second supported statement"], ["f1"])]
    geometries = []
    first_bullet_lefts = []
    for index in range(3):
        target = tmp_path / f"variant-{index}.pptx"
        render_variant(source, template, content, index, target)
        generated = Presentation(target)
        shapes = generated.slides[0].shapes
        texts = [shape.text for shape in shapes if shape.has_text_frame]
        assert "Brand" in texts
        assert "New title" in texts
        assert not any("Old" in text for text in texts)
        assert len({shape.shape_id for shape in shapes}) == len(shapes)
        geometries.append([(shape.left, shape.top, shape.width, shape.height) for shape in shapes if shape.has_text_frame and shape.text.strip()])
        first_bullet_lefts.append(next(shape.left for shape in shapes if shape.has_text_frame and "First supported statement" in shape.text))
    assert geometries[0] != geometries[1]
    assert geometries[2] != geometries[0]
    assert first_bullet_lefts[2] > first_bullet_lefts[0]
    ir = save_prepared(template, tmp_path / "template.json")
    source.write_bytes(b"another template")
    with pytest.raises(ValueError, match="different PPTX"):
        load_prepared(ir, source)


def test_resized_region_does_not_inherit_oversized_padding():
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from engine.renderer import _clear_text_preserving_style
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(0, 0, Inches(3), Inches(0.7))
    shape.text_frame.margin_top = Inches(0.5)
    shape.text_frame.margin_bottom = Inches(0.5)
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.space_before = Pt(36)
    paragraph.add_run().text = "Original"
    _clear_text_preserving_style(shape, ["Readable content"])
    assert shape.text_frame.margin_top + shape.text_frame.margin_bottom < shape.height / 4
    assert shape.text_frame.paragraphs[0].space_before == 0



def test_dark_variant_keeps_template_image_and_uses_contrasting_editable_text(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    panel = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, presentation.slide_width, presentation.slide_height
    )
    panel.name = "Neutral panel"
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor(255, 255, 255)
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(8), Inches(0.8))
    title.text = "Old heading"
    body = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(8), Inches(3.5))
    body.text = "Old body"
    # A tiny source image stands in for an unchanged logo or photograph.
    original_image = b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+j5s8AAAAASUVORK5CYII="
    )
    logo = slide.shapes.add_picture(
        BytesIO(original_image), Inches(8.5), Inches(0.25), width=Inches(0.6)
    )
    source = tmp_path / "source.pptx"
    presentation.save(source)
    template = inspect_template(source)
    template.compositions[0].object_roles[str(logo.shape_id)] = "fixed_brand"
    content = [SlideContent("s1", "New heading", ["First point", "Second point"])]
    palette = {
        "background": "#081725",
        "surface": "#203346",
        "text": "#F0F5F9",
        "accent": "#6FD7CA",
    }
    target = tmp_path / "dark.pptx"
    metrics = render_variant(source, template, content, 1, target, dark_palette=palette)

    generated = Presentation(target).slides[0]
    assert generated.background.fill.fore_color.rgb == RGBColor(8, 23, 37)
    assert next(shape for shape in generated.shapes if shape.name == "Neutral panel").fill.fore_color.rgb == RGBColor(8, 23, 37)
    assert next(shape for shape in generated.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE).image.blob == original_image
    heading = next(shape for shape in generated.shapes if shape.has_text_frame and shape.text == "New heading")
    body = next(shape for shape in generated.shapes if shape.has_text_frame and "First point" in shape.text)
    assert heading.text_frame.paragraphs[0].runs[0].font.color.rgb == RGBColor(111, 215, 202)
    assert body.text_frame.paragraphs[0].runs[0].font.color.rgb == RGBColor(240, 245, 249)
    assert metrics["dark_palette"] == {key: value.upper() for key, value in palette.items()}
    assert metrics["composition_strategy"] == "dark_theme"


def test_dark_native_data_remains_editable_and_readable(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(8), Inches(0.8)).text = "Old heading"
    slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(8), Inches(3.5)).text = "Old body"
    source = tmp_path / "source.pptx"
    presentation.save(source)
    template = inspect_template(source)
    content = [SlideContent("s1", "Revenue", ["Growth"])]
    data = NumericCsvData(
        "Quarter",
        ("Q1", "Q2"),
        (NumericSeries("Revenue", (Decimal("10"), Decimal("12"))),),
        (("Q1", "10"), ("Q2", "12")),
    )
    target = tmp_path / "dark-data.pptx"
    render_variant(source, template, content, 1, target, native_data=data, native_slide_index=0, dark_palette=DARK_PALETTE)
    generated = Presentation(target).slides[0]
    chart = next(shape.chart for shape in generated.shapes if shape.has_chart)
    table = next(shape.table for shape in generated.shapes if shape.has_table)
    text = RGBColor.from_string(DARK_PALETTE["text"][1:])
    surface = RGBColor.from_string(DARK_PALETTE["surface"][1:])
    assert generated.background.fill.fore_color.rgb == RGBColor.from_string(DARK_PALETTE["background"][1:])
    assert table.cell(0, 0).text == "Quarter"
    assert table.cell(1, 1).text == "10"
    assert table.cell(1, 0).fill.fore_color.rgb == surface
    assert table.cell(1, 0).text_frame.paragraphs[0].runs[0].font.color.rgb == text
    assert chart.category_axis.tick_labels.font.color.rgb == text
    assert chart.value_axis.tick_labels.font.color.rgb == text
    assert chart._chartSpace.find("{http://schemas.openxmlformats.org/drawingml/2006/chart}spPr") is not None
    assert chart._chartSpace.chart.plotArea.find("{http://schemas.openxmlformats.org/drawingml/2006/chart}spPr") is not None
    assert list(chart.series[0].values) == [10.0, 12.0]


def test_invalid_model_dark_palette_falls_back_to_readable_defaults(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(8), Inches(0.8)).text = "Old heading"
    slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(8), Inches(3.5)).text = "Old body"
    source = tmp_path / "source.pptx"
    presentation.save(source)
    template = inspect_template(source)
    target = tmp_path / "dark-fallback.pptx"
    metrics = render_variant(
        source, template, [SlideContent("s1", "Heading", ["Point"])], 1, target,
        dark_palette={"background": "#FFFFFF", "surface": "#FFFFFF", "text": "#FFFFFF", "accent": "#FFFFFF"},
    )
    assert metrics["dark_palette"] == DARK_PALETTE
    assert Presentation(target).slides[0].background.fill.fore_color.rgb == RGBColor.from_string(DARK_PALETTE["background"][1:])


def test_all_variants_keep_same_source_slide_sequence(tmp_path):
    presentation = Presentation()
    for index in range(3):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(0.5 + index * 0.2), Inches(0.3), Inches(8), Inches(0.7)).text = f"Source title {index}"
        slide.shapes.add_textbox(Inches(0.5 + index * 0.2), Inches(1.4), Inches(8), Inches(4)).text = f"Source body {index}"
    source = tmp_path / "varied-template.pptx"
    presentation.save(source)
    template = inspect_template(source)
    contents = [SlideContent(f"slide_{index}", f"Heading {index}", [f"Point {index}"]) for index in range(3)]
    source_sequences = []
    for variant in range(3):
        metrics = render_variant(source, template, contents, variant, tmp_path / f"variant-{variant}.pptx")
        source_sequences.append(metrics["source_slides"])
    assert source_sequences[0] == source_sequences[1] == source_sequences[2]

def test_source_layout_sampling_preserves_cover_closing_and_order(tmp_path):
    from engine.renderer import choose_compositions

    presentation = Presentation()
    for index in range(16):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(8), Inches(0.7)).text = f"Title {index}"
        slide.shapes.add_textbox(Inches(0.5), Inches(1.4), Inches(8), Inches(4)).text = f"Body {index}"
    source = tmp_path / "template-16.pptx"
    presentation.save(source)
    template = inspect_template(source)
    chosen = choose_compositions(template, 10)
    assert [item.source_slide_index + 1 for item in chosen] == [1, 2, 4, 6, 8, 9, 11, 13, 15, 16]

def test_infographic_cards_are_filled_and_invalid_source_xml_is_not_copied(tmp_path):
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.7)).text = "Old title"
    slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(9), Inches(0.5)).text = "Old subtitle"
    for index in range(3):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                      Inches(0.6 + index * 3), Inches(2.2), Inches(2.7), Inches(2.3))
        card.name = f"Card {index + 1}"
        label = slide.shapes.add_textbox(Inches(0.8 + index * 3), Inches(2.5), Inches(2), Inches(0.4))
        label.text = f"Old label {index + 1}"
        effect = OxmlElement("a:effectLst")
        shadow = OxmlElement("a:outerShdw")
        shadow.set("blurRad", "1.6e+24")
        shadow.set("dir", "3.5e+25")
        effect.append(shadow)
        card._element.spPr.append(effect)
    source = tmp_path / "source.pptx"
    deck.save(source)
    template = inspect_template(source)
    target = tmp_path / "generated.pptx"
    contents = [SlideContent("slide_1", "New title", ["First supported point", "Second supported point"])]
    render_variant(source, template, contents, 0, target)
    generated = Presentation(target)
    shapes = generated.slides[0].shapes
    assert len([shape for shape in shapes if shape.name.startswith("Card ")]) == 2
    text = "\n".join(shape.text for shape in shapes if shape.has_text_frame)
    assert all(point in text for point in contents[0].bullets)
    assert "Old label" not in text and "Old subtitle" not in text
    assert not list(generated.slides[0]._element.iter(qn("a:outerShdw")))


def test_generated_body_is_readable_on_dark_source_background(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(8, 8, 11)
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(8), Inches(0.8))
    title.text = "Old title"
    body = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(8), Inches(3))
    body.text = "Old body"
    source = tmp_path / "dark-source.pptx"
    presentation.save(source)
    template = inspect_template(source)
    for variant in (0, 2):
        target = tmp_path / f"dark-{variant}.pptx"
        render_variant(source, template, [SlideContent("slide_1", "Heading", ["New readable point"])], variant, target)
        generated = Presentation(target).slides[0]
        body = next(shape for shape in generated.shapes if shape.has_text_frame and "New readable point" in shape.text)
        assert {run.font.color.rgb for paragraph in body.text_frame.paragraphs for run in paragraph.runs} == {RGBColor(248, 250, 252)}


def test_single_rectangular_panel_is_used_without_mistaking_bleeding_oval_for_card():
    from engine.renderer import _card_regions
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(8), Inches(0.8))
    title.text = "Title"
    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4), Inches(2), Inches(5), Inches(3))
    slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10), Inches(5.7), Inches(4), Inches(4))
    assert [shape.shape_id for shape in _card_regions(slide, title, presentation.slide_width, presentation.slide_height)] == [panel.shape_id]



def test_light_slide_keeps_template_font_size_and_fixes_white_on_white(tmp_path):
    from engine.renderer import _contrast
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    title = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(8.0), Inches(1.2))
    body = slide.shapes.add_textbox(Inches(0.6), Inches(1.9), Inches(8.0), Inches(3.5))
    for shape, size in ((title, 40), (body, 18)):
        run = shape.text_frame.paragraphs[0].add_run()
        run.text = "Old text"
        run.font.name = "Play"
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor(255, 255, 255)
    source = tmp_path / "source.pptx"
    deck.save(source)
    template = inspect_template(source)
    target = tmp_path / "result.pptx"
    metrics = render_variant(
        source, template, [SlideContent("s1", "Readable heading", ["Readable first point", "Readable second point"])],
        0, target,
    )
    generated = Presentation(target).slides[0]
    heading = next(shape for shape in generated.shapes if shape.has_text_frame and shape.text == "Readable heading")
    body = next(shape for shape in generated.shapes if shape.has_text_frame and "Readable first point" in shape.text)
    assert heading.text_frame.paragraphs[0].runs[0].font.name == "Play"
    assert heading.text_frame.paragraphs[0].runs[0].font.size == Pt(40)
    assert body.text_frame.paragraphs[0].runs[0].font.name == "Play"
    assert body.text_frame.paragraphs[0].runs[0].font.size == Pt(18)
    assert metrics["minimum_text_contrast"] >= 4.5
    for shape in (heading, body):
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                assert _contrast(run.font.color.rgb, RGBColor(255, 255, 255)) >= 4.5


def test_infographic_is_native_editable_and_contrasting_in_light_and_dark(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(11), Inches(0.9))
    title.text = "Old heading"
    title.text_frame.paragraphs[0].runs[0].font.name = "Play"
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(38)
    body = slide.shapes.add_textbox(Inches(0.7), Inches(1.8), Inches(11), Inches(4.5))
    body.text = "Old material"
    body.text_frame.paragraphs[0].runs[0].font.name = "Play"
    body.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    source = tmp_path / "source.pptx"
    deck.save(source)
    template = inspect_template(source)
    content = [SlideContent("s1", "New heading", ["Inputs from the repository", "Individual learning track"])]
    for variant in (0, 1):
        target = tmp_path / f"infographic-{variant}.pptx"
        metrics = render_variant(
            source, template, content, variant, target,
            infographic_layouts={"s1": "flow"},
        )
        generated = Presentation(target).slides[0]
        nodes = [shape for shape in generated.shapes if shape.name.startswith("Infographic node")]
        connectors = [shape for shape in generated.shapes if shape.name.startswith("Infographic connector")]
        assert len(nodes) == 2 and len(connectors) == 1
        assert [node.text for node in nodes] == content[0].bullets
        assert all(node.text_frame.paragraphs[0].runs[0].font.name == "Play" for node in nodes)
        assert metrics["infographic_slides"] == [1]
        assert metrics["minimum_text_contrast"] >= 4.5
        assert "Old material" not in "\n".join(shape.text for shape in generated.shapes if shape.has_text_frame)


def test_analysed_template_uses_designed_layouts_and_a_region_per_point(tmp_path):
    from engine.renderer import choose_compositions, render_variant
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)

    def slide_with(heading, points=0):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        if heading:
            box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11), Inches(0.9))
            box.text = heading
            box.text_frame.paragraphs[0].runs[0].font.size = Pt(32)
        for number in range(points):
            marker = slide.shapes.add_textbox(Inches(0.8 + number * 4), Inches(2.4), Inches(0.8), Inches(0.6))
            marker.text = str(number + 1)
            point = slide.shapes.add_textbox(Inches(0.8 + number * 4), Inches(3.2), Inches(3.4), Inches(1.2))
            point.text = f"Описание пункта {number + 1}"
            point.text_frame.paragraphs[0].runs[0].font.size = Pt(16)

    slide_with("Тема презентации")                     # cover
    slide_with("")                                     # brand backdrop without text
    slide_with("Пример оформления таблиц")             # style-guide page
    slide_with("Заголовок — ключевая мысль слайда", 3)  # designed three-point layout
    slide_with("Спасибо за внимание")                  # closing
    source = tmp_path / "template.pptx"
    deck.save(source)
    template = inspect_template(source)
    for composition, archetype in zip(template.compositions, ["title", "content", "content", "content", "closing"]):
        composition.archetype, composition.analysis_complete = archetype, True
    slides = [SlideContent("slide_1", "Лукас", ["Команда проекта"], ["f1"]),
              SlideContent("slide_2", "Три шага", ["Загрузите шаблон.", "Опишите тему.", "Скачайте варианты."], ["f1"]),
              SlideContent("slide_3", "Спасибо", ["Вопросы"], ["f1"])]
    assert [item.source_slide_index for item in choose_compositions(template, 3, slides)] == [0, 3, 4]
    target = tmp_path / "deck.pptx"
    render_variant(source, template, slides, 0, target)
    texts = [shape.text for shape in Presentation(target).slides[1].shapes if shape.has_text_frame and shape.text.strip()]
    assert {"Загрузите шаблон.", "Опишите тему.", "Скачайте варианты."} <= set(texts)
    assert {"1", "2", "3"} <= set(texts)


def test_letter_collisions_ignore_empty_parts_of_frames():
    from engine.rendered_audit import text_collisions
    # A paragraph with a short last line and a block to the right of that line: the outlines
    # of the two objects cross, their drawn lines do not.
    first = [(0, 0, 300, 12), (0, 14, 80, 26)]
    beside = [(100, 16, 300, 28)]
    assert text_collisions(first, beside) == []
    over = [(50, 5, 200, 17)]
    assert text_collisions(first, over)


def test_pdf_letters_find_real_overlap_but_not_crossing_empty_frames(tmp_path):
    import shutil
    from engine.export import export_pdf
    from engine.rendered_audit import audit_pdf
    try:
        from engine.export import _find_soffice
        _find_soffice()
    except (ImportError, RuntimeError):
        if not (shutil.which("soffice") or shutil.which("soffice.com")):
            pytest.skip("LibreOffice is not installed")
    deck = Presentation()
    empty = deck.slides.add_slide(deck.slide_layouts[6])
    # Two tall frames crossing in their empty lower parts: the text sits apart.
    empty.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(5), Inches(4)).text = "Первый блок текста"
    empty.shapes.add_textbox(Inches(3), Inches(2.5), Inches(5), Inches(3)).text = "Второй блок текста"
    crowded = deck.slides.add_slide(deck.slide_layouts[6])
    crowded.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(6), Inches(1)).text = "Наложение строк на слайде"
    crowded.shapes.add_textbox(Inches(0.6), Inches(0.52), Inches(6), Inches(1)).text = "Второй текст поверх первого"
    source = tmp_path / "overlap.pptx"
    deck.save(source)
    issues = audit_pdf(source, export_pdf(source, tmp_path / "pdf"))
    overlaps = [issue["slide_id"] for issue in issues if issue["rule_id"] == "rendered_text_overlap"]
    assert overlaps == ["slide_2"]


def test_forced_relayout_moves_a_crowded_slide_into_a_clean_region(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.9)).text = "Heading"
    slide.shapes.add_textbox(Inches(0.5), Inches(1.4), Inches(4), Inches(1)).text = "Body"
    source = tmp_path / "template.pptx"
    deck.save(source)
    content = [SlideContent("slide_1", "Заголовок", ["Первый тезис", "Второй тезис"], ["f1"])]
    target = tmp_path / "relaid.pptx"
    metrics = render_variant(source, inspect_template(source), content, 0, target, relayout={0})
    assert metrics["relayout_slides"] == [1] and 1 in metrics["layout_repairs"]
    names = [shape.name for shape in Presentation(target).slides[0].shapes]
    assert "Reflowed content" in names
