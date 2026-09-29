import json
import struct
import time
from pathlib import Path
from types import SimpleNamespace

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from engine import fonts, pipeline, rendered_audit, typography
from engine.audit import font_issues
from engine.ingest import inspect_template
from engine.models import Fact, SlideContent
from engine.renderer import render_variant


def system_ttf() -> Path:
    for family in ("Arial", "DejaVu Sans", "Liberation Sans"):
        for path, _ in typography.font_files(family):
            if path.lower().endswith(".ttf"):
                return Path(path)
    pytest.skip("No TrueType font installed")


@pytest.fixture
def font_store(tmp_path, monkeypatch):
    """An empty application font store holding one 'Brand Sans' family."""
    store = tmp_path / "fonts"
    (store / "Brand_Sans").mkdir(parents=True)
    source = system_ttf().read_bytes()
    (store / "Brand_Sans" / "regular.ttf").write_bytes(source)
    monkeypatch.setenv("AYA_FONT_DIR", str(store))
    typography.refresh_fonts()
    yield store
    typography.refresh_fonts()


def test_eot_wraps_the_truetype_font_for_pptx_embedding():
    data = system_ttf().read_bytes()
    eot = fonts.eot_from_ttf(data)
    size, font_size, version, flags = struct.unpack_from("<IIII", eot, 0)
    assert size == len(eot) and font_size == len(data) and version == 0x00020001 and flags == 0
    assert struct.unpack_from("<H", eot, 34)[0] == 0x504C
    assert eot.endswith(data)


def test_store_font_is_found_under_the_template_name_and_embedded(font_store, tmp_path):
    assert typography.font_files("Brand Sans")
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box.text_frame.text = "Фирменный текст"
    box.text_frame.paragraphs[0].runs[0].font.name = "Brand Sans"
    assert "Brand Sans" in fonts.used_families(deck)
    report = fonts.embed_fonts(deck, fonts.used_families(deck))
    assert report["embedded"] == ["Brand Sans"] and "Calibri" in report["core"]
    target = tmp_path / "embedded.pptx"
    deck.save(target)
    reopened = Presentation(target)
    assert fonts.embedded_families(reopened) == {"Brand Sans"}
    assert reopened.part._element.get("embedTrueTypeFonts") == "1"
    assert font_issues(target, []) == []
    # Embedding twice keeps one entry.
    assert fonts.embed_fonts(reopened, {"Brand Sans"})["already_embedded"] == ["Brand Sans"]


def test_font_inspector_reports_family_size_and_missing_embedding(font_store, tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box.text_frame.text = "Текст"
    run = box.text_frame.paragraphs[0].runs[0]
    run.font.name, run.font.size = "Brand Sans", Pt(14)
    target = tmp_path / "deck.pptx"
    deck.save(target)
    issues = font_issues(target, [{"slide": 1, "shape_id": box.shape_id, "family": "Play", "size": 18.0}])
    rules = {issue["rule_id"]: issue for issue in issues}
    assert "«Brand Sans» вместо «Play»" in rules["font_mismatch"]["evidence"]
    assert "14 pt вместо 18 pt" in rules["font_size_mismatch"]["evidence"]
    assert rules["font_not_embedded"]["severity"] == "warning"


def test_rendered_audit_detects_a_substituted_typeface(tmp_path, monkeypatch):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box.text_frame.text = "Play"
    box.text_frame.paragraphs[0].runs[0].font.name = "Play"
    source = tmp_path / "deck.pptx"
    deck.save(source)

    def page(fontname):
        chars = [{"text": c, "x0": 80 + i * 10, "x1": 90 + i * 10, "top": 80, "bottom": 100, "fontname": fontname}
                 for i, c in enumerate("Play")]
        return SimpleNamespace(width=720, height=540, chars=chars, close=lambda: None)

    class PDF:
        def __init__(self, fontname):
            self.pages = [page(fontname)]
        def __enter__(self):
            return self
        def __exit__(self, *args):
            pass

    monkeypatch.setattr(rendered_audit.pdfplumber, "open", lambda _: PDF("BAAAAA+DejaVuSans"))
    issues = rendered_audit.audit_pdf(source, Path("deck.pdf"))
    assert [issue["rule_id"] for issue in issues] == ["font_substituted"]
    assert "«Play»" in issues[0]["evidence"] and "DejaVuSans" in issues[0]["evidence"]
    # Blocking only when the platform has the font; otherwise the author is told to add it.
    assert issues[0]["severity"] == ("blocking" if typography.font_files("Play") else "warning")
    monkeypatch.setattr(rendered_audit.pdfplumber, "open", lambda _: PDF("CAAAAA+Play-Regular"))
    assert rendered_audit.audit_pdf(source, Path("deck.pdf")) == []


def test_open_font_is_fetched_once_into_the_store(tmp_path, monkeypatch):
    monkeypatch.setenv("AYA_FONT_DIR", str(tmp_path / "fonts"))
    typography.refresh_fonts()
    data = system_ttf().read_bytes()
    requests = []

    def download(url, limit=0):
        requests.append(url)
        if "fonts.googleapis.com" in url:
            if "ital" in url:
                raise OSError("no italic")
            return b"src: url(https://fonts.gstatic.com/s/brand/v1/face.ttf) format('truetype');"
        return data

    monkeypatch.setattr(fonts, "_download", download)
    try:
        report = fonts.ensure_fonts({"Brand Sans SemiBold"}, download=True)
        assert report["downloaded"] == ["Brand Sans SemiBold"]
        assert "family=Brand+Sans:wght@600" in requests[0]
        assert sorted(path.name for path in (tmp_path / "fonts" / "Brand_Sans_SemiBold").iterdir()) == ["bold.ttf", "regular.ttf"]
        assert typography.font_files("Brand Sans SemiBold")
        requests.clear()
        assert fonts.ensure_fonts({"Brand Sans SemiBold"}, download=True)["available"] == ["Brand Sans SemiBold"]
        assert requests == []
        assert fonts.ensure_fonts({"Unknown Corporate"}, download=False)["missing"] == ["Unknown Corporate"]
    finally:
        typography.refresh_fonts()


def test_generated_text_keeps_the_template_family_and_size_exactly(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.9))
    title.text_frame.text = "Old heading"
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(34)
    title.text_frame.paragraphs[0].runs[0].font.name = "Georgia"
    body = slide.shapes.add_textbox(Inches(0.5), Inches(1.6), Inches(9), Inches(4.5))
    body.text_frame.text = "Old body text"
    body.text_frame.paragraphs[0].runs[0].font.size = Pt(13)
    body.text_frame.paragraphs[0].runs[0].font.name = "Verdana"
    source = tmp_path / "template.pptx"
    deck.save(source)
    content = [SlideContent("slide_1", "Новый заголовок", ["Первый тезис", "Второй тезис"], ["f1"])]
    for variant in range(3):
        target = tmp_path / f"{variant}.pptx"
        metrics = render_variant(source, inspect_template(source), content, variant, target)
        rendered = Presentation(target).slides[0]
        runs = {run.text: run for shape in rendered.shapes if shape.has_text_frame
                for paragraph in shape.text_frame.paragraphs for run in paragraph.runs}
        assert (runs["Новый заголовок"].font.name, runs["Новый заголовок"].font.size.pt) == ("Georgia", 34)
        for point in ("Первый тезис", "Второй тезис"):
            assert (runs[point].font.name, runs[point].font.size.pt) == ("Verdana", 13)
        assert {(item["family"], item["size"]) for item in metrics["font_expectations"]} == {("Georgia", 34.0), ("Verdana", 13.0)}
        assert font_issues(target, metrics["font_expectations"]) == []


def test_overflow_repair_grows_the_frame_and_keeps_the_template_size(tmp_path, monkeypatch):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(0.5))
    box.text_frame.word_wrap = True
    box.text_frame.text = "Длинный тезис, который не помещается в короткую рамку шаблона при его кегле"
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    base = tmp_path / "base.pptx"
    deck.save(base)

    def export_pdf(pptx, folder, **kwargs):
        target = Path(folder) / "out.pdf"
        target.write_bytes(b"%PDF")
        return target

    monkeypatch.setattr(pipeline, "export_pdf", export_pdf)
    monkeypatch.setattr(pipeline, "export_previews", lambda pdf, folder, **kwargs: [])
    monkeypatch.setattr(pipeline, "export_html", lambda pptx, previews, target: target)
    issue = {"issue_id": "i1", "slide_id": "slide_1", "object_id": str(box.shape_id),
             "rule_id": "possible_text_overflow", "repairability": "automatic"}
    result = pipeline.repair(base, base, "Бриф", 1, [issue], tmp_path / "out")[0]
    repaired = next(shape for shape in Presentation(result.pptx_path).slides[0].shapes if shape.shape_id == box.shape_id)
    assert repaired.text_frame.paragraphs[0].runs[0].font.size.pt == 24
    assert repaired.height > box.height


def test_planned_diagram_without_room_under_a_template_size_heading_stays_text(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(1))
    title.text_frame.text = "Heading"
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(60)
    slide.shapes.add_textbox(Inches(0.5), Inches(2), Inches(9), Inches(3)).text_frame.text = "Body"
    source = tmp_path / "template.pptx"
    deck.save(source)
    heading = "Очень длинный заголовок слайда, который в кегле шаблона занимает почти всю высоту и не оставляет места"
    content = [SlideContent("slide_1", heading, ["Первый этап работы", "Второй этап работы"], ["f1"])]
    metrics = render_variant(source, inspect_template(source), content, 0, tmp_path / "out.pptx",
                             infographic_layouts={"slide_1": "modules"})
    assert metrics["infographic_slides"] == [] and metrics["infographic_fallback_slides"] == [1]
    runs = [run for shape in Presentation(tmp_path / "out.pptx").slides[0].shapes if shape.has_text_frame
            for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if run.text == heading]
    assert runs and runs[0].font.size.pt == 60
