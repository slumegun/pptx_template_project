import json
from pathlib import Path
from io import BytesIO
from base64 import b64decode
from types import SimpleNamespace

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from engine import swarm
from engine.models import Composition, PreparedTemplate, SCHEMA_VERSION, SlideContent
from engine.renderer import render_variant
from engine.provider import ModelProviderError


def _template(tmp_path: Path, slide_count: int):
    path = tmp_path / "unseen.pptx"
    deck = Presentation()
    compositions = []
    for index in range(slide_count):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = f"Old slide {index}"
        compositions.append(Composition(
            source_slide_index=index, layout_index=6, fingerprint=f"slide-{index}",
            slots=[], shape_count=30, picture_area_ratio=0.0, chart_count=0,
            table_count=0, score=1.0,
        ))
    deck.save(path)
    template = PreparedTemplate(
        SCHEMA_VERSION, "test-sha", path.name, deck.slide_width, deck.slide_height,
        slide_count, len(deck.slide_layouts), compositions,
    )
    previews = [tmp_path / f"preview-{index}.png" for index in range(slide_count)]
    return path, template, previews


def _rows(inventory):
    return [{
        "source_slide_index": item["source_slide_index"],
        "archetype": "content",
        "object_roles": {shape["object_id"]: "replaceable" for shape in item["objects"]},
    } for item in inventory]


def test_incomplete_second_batch_retries_each_slide_and_checkpoints(tmp_path, monkeypatch):
    path, template, previews = _template(tmp_path, 4)
    clock = SimpleNamespace(now=100.0)
    clock.monotonic = lambda: clock.now
    def sleep(seconds):
        clock.now += seconds
    clock.sleep = sleep
    monkeypatch.setattr(swarm, "time", clock)
    monkeypatch.setattr(swarm, "contact_sheet", lambda paths: "data:image/jpeg;base64,dGVzdA==")
    monkeypatch.setenv("MODEL_ANALYSIS_RPM", "60")
    calls = []
    class Client:
        vision_model = "qwen/qwen3-vl-32b-instruct"
        def complete_json(self, system, user, **kwargs):
            assert "Template Analyst" in system
            assert kwargs["vision"] is True
            assert kwargs["image_data_url"].startswith("data:image/jpeg;base64,")
            inventory = json.loads(user)
            calls.append((clock.now, [item["source_slide_index"] for item in inventory]))
            if len(calls) == 2:
                return {"slides": _rows(inventory[:1])}
            return {"slides": _rows(inventory)}
    snapshots = []
    swarm.analyze_template(
        Client(), path, template, previews,
        checkpoint=lambda state: snapshots.append([item.analysis_complete for item in state.compositions]),
    )
    assert [indices for _, indices in calls] == [[0, 1], [2, 3], [2], [3]]
    assert [timestamp for timestamp, _ in calls] == pytest.approx([100, 101, 102, 103])
    assert snapshots == [[True, True, False, False], [True, True, True, False], [True, True, True, True]]
    assert all(item.object_roles and item.analysis_complete for item in template.compositions)
    assert template.analysis_mode == "vision_model"
    assert template.analysis_model == "qwen/qwen3-vl-32b-instruct"


def test_single_slide_failure_identifies_source_and_keeps_checkpoint(tmp_path, monkeypatch):
    path, template, previews = _template(tmp_path, 2)
    monkeypatch.setattr(swarm, "contact_sheet", lambda paths: "data:image/jpeg;base64,dGVzdA==")
    monkeypatch.setenv("MODEL_ANALYSIS_RPM", "6000")
    calls = []
    class Client:
        vision_model = "qwen/qwen3-vl-32b-instruct"
        def complete_json(self, system, user, **kwargs):
            inventory = json.loads(user)
            calls.append([item["source_slide_index"] for item in inventory])
            if len(calls) == 1:
                return {"slides": []}
            if len(calls) == 2:
                return {"slides": _rows(inventory)}
            return {"slides": []}
    snapshots = []
    with pytest.raises(ModelProviderError, match=r"слайд 2 \(source_slide_index=1\).*неполную классификацию"):
        swarm.analyze_template(
            Client(), path, template, previews,
            checkpoint=lambda state: snapshots.append([item.analysis_complete for item in state.compositions]),
        )
    assert calls == [[0, 1], [0], [1]]
    assert snapshots == [[True, False]]
    assert template.compositions[0].object_roles
    assert template.compositions[1].object_roles == {}
    assert template.analysis_mode == "structural"


def test_payment_error_does_not_trigger_per_slide_retries(tmp_path, monkeypatch):
    path, template, previews = _template(tmp_path, 2)
    monkeypatch.setattr(swarm, "contact_sheet", lambda paths: "data:image/jpeg;base64,dGVzdA==")
    calls = []
    class Client:
        vision_model = "qwen/qwen3-vl-32b-instruct"
        def complete_json(self, system, user, **kwargs):
            calls.append(json.loads(user))
            raise ModelProviderError("Model API HTTP 402: OpenRouter balance exhausted")
    with pytest.raises(ModelProviderError, match="HTTP 402"):
        swarm.analyze_template(Client(), path, template, previews)
    assert len(calls) == 1
    assert all(not item.analysis_complete for item in template.compositions)


def test_truncated_batch_retries_smaller_requests(tmp_path, monkeypatch):
    path, template, previews = _template(tmp_path, 2)
    monkeypatch.setattr(swarm, "contact_sheet", lambda paths: "data:image/jpeg;base64,dGVzdA==")
    monkeypatch.setenv("MODEL_ANALYSIS_RPM", "6000")
    calls = []
    class Client:
        vision_model = "qwen/qwen3-vl-32b-instruct"
        def complete_json(self, system, user, **kwargs):
            inventory = json.loads(user)
            calls.append([item["source_slide_index"] for item in inventory])
            if len(calls) == 1:
                raise ModelProviderError("Model response was truncated; reduce the task or increase max_tokens")
            return {"slides": _rows(inventory)}
    swarm.analyze_template(Client(), path, template, previews)
    assert calls == [[0, 1], [0], [1]]
    assert all(item.analysis_complete for item in template.compositions)


def test_partial_roles_are_safely_completed_and_old_text_is_removed(tmp_path, monkeypatch):
    path, template, previews = _template(tmp_path, 1)
    deck = Presentation(path)
    slide = deck.slides[0]
    old_text = slide.shapes[0]
    simple = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1), Inches(2), Inches(2), Inches(1),
    )
    image = b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+j5s8AAAAASUVORK5CYII="
    )
    logo = slide.shapes.add_picture(BytesIO(image), Inches(8), Inches(0.2), width=Inches(0.5))
    old_picture = slide.shapes.add_picture(BytesIO(image), Inches(5), Inches(2), width=Inches(2))
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(4), Inches(4), Inches(1)).text = "Old flood details"
    deck.save(path)
    template.compositions[0].shape_count = len(slide.shapes)
    monkeypatch.setattr(swarm, "contact_sheet", lambda paths: "data:image/jpeg;base64,dGVzdA==")

    class Client:
        vision_model = "qwen/qwen3-vl-32b-instruct"
        def complete_json(self, system, user, **kwargs):
            return {"slides": [{
                "source_slide_index": 0,
                "archetype": "content",
                "roles": {
                    "fixed_brand": [old_text.shape_id, logo.shape_id, group.shape_id, 99999],
                    "not_a_role": [simple.shape_id],
                },
            }]}

    swarm.analyze_template(Client(), path, template, previews)
    roles = template.compositions[0].object_roles
    assert set(roles) == {str(shape.shape_id) for shape in slide.shapes}
    assert roles[str(old_text.shape_id)] == "replaceable"
    assert roles[str(simple.shape_id)] == "replaceable"
    assert roles[str(logo.shape_id)] == "fixed_brand"
    assert roles[str(old_picture.shape_id)] == "unresolved"
    assert roles[str(group.shape_id)] == "unresolved"
    assert any("слайд 1" in warning and "2 объектов" in warning for warning in template.warnings)
    assert any("слайд 1" in warning and "2 текстовых" in warning for warning in template.warnings)

    output = tmp_path / "generated.pptx"
    render_variant(path, template, [SlideContent("s1", "Fire map", ["Satellite CV"])], 0, output)
    generated = Presentation(output).slides[0]
    texts = "\n".join(shape.text for shape in generated.shapes if shape.has_text_frame)
    assert "Old slide" not in texts
    assert "Old flood" not in texts
    assert "Fire map" in texts
    assert "Satellite CV" in texts


def test_partial_object_roles_map_ignores_unknown_ids_and_roles(tmp_path, monkeypatch):
    path, template, previews = _template(tmp_path, 1)
    shape_id = str(Presentation(path).slides[0].shapes[0].shape_id)
    monkeypatch.setattr(swarm, "contact_sheet", lambda paths: "data:image/jpeg;base64,dGVzdA==")
    class Client:
        vision_model = "qwen/qwen3-vl-32b-instruct"
        def complete_json(self, system, user, **kwargs):
            return {"slides": [{
                "source_slide_index": 0,
                "object_roles": {shape_id: "unknown", "99999": "fixed_brand"},
            }]}
    swarm.analyze_template(Client(), path, template, previews)
    assert template.compositions[0].object_roles == {shape_id: "replaceable"}
    assert any("слайд 1" in warning and "1 объектов" in warning for warning in template.warnings)
