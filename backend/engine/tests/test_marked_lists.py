"""A list drawn as one text box beside a column of checkmarks, and a heading in a narrow box that does not wrap."""
import json
import time
from io import BytesIO

from PIL import Image
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt

from engine import pipeline
from engine.ingest import inspect_template
from engine.models import Fact, SlideContent
from engine.renderer import render_variant, slot_budget, strict_regions


def png(size=(40, 40)):
    stream = BytesIO()
    Image.new("RGB", size, (255, 90, 40)).save(stream, format="PNG")
    return stream


def textbox(slide, left, top, width, height, text, size, wrap=True):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    box.text_frame.word_wrap = wrap
    box.text_frame.text = text
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
    box.text_frame.paragraphs[0].runs[0].font.name = "Arial"
    return box


def checkmark_template(tmp_path):
    """The layout of a real template: a short heading in an auto-size box, one text box, four checkmarks."""
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    heading = textbox(slide, 0.8, 0.55, 1.6, 0.8, "Кейс", 36, wrap=False)
    heading.text_frame.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
    body = textbox(slide, 1.25, 2.3, 5.2, 2.7, "Текст", 20)
    checks = [slide.shapes.add_picture(png(), Inches(0.82), Inches(2.4 + number * 0.62), Inches(0.26), Inches(0.26))
              for number in range(4)]
    caption = textbox(slide, 0.8, 5.75, 6, 0.4, "Подпись к слайду", 12)
    source = tmp_path / "template.pptx"
    deck.save(source)
    template = inspect_template(source)
    template.compositions[0].analysis_complete = True
    return source, template, heading, body, checks, caption


def test_a_column_of_checkmarks_beside_one_text_box_is_a_place_per_checkmark(tmp_path):
    source, template, heading, body, checks, caption = checkmark_template(tmp_path)
    composition = template.compositions[0]
    found, regions = strict_regions(composition, template.width, template.height)
    # The narrow "Кейс" is the heading and holds a real title, not four letters.
    assert found.shape_id == heading.shape_id and found.max_chars >= 25
    assert [region.kind for region in regions] == ["marked_item"] * 4 + ["text"]
    budget = slot_budget([composition], template.width, template.height)
    assert budget["min_texts"] == 4 and budget["title"] >= 25
    bullets = ["Облачная или своя модель", "Выбор по задаче и риску", "Единая консоль и квоты"]
    content = [SlideContent("slide_1", "Управление моделями ИИ", bullets, ["f1"])]
    target = tmp_path / "deck.pptx"
    metrics = render_variant(source, template, content, 0, target, strict=True, plan=[0])
    slide = Presentation(target).slides[0]
    shapes = {shape.shape_id: shape for shape in slide.shapes}
    title = shapes[heading.shape_id]
    assert title.text == "Управление моделями ИИ" and title.text_frame.word_wrap is True
    assert title.left + title.width <= template.width
    # Every point is its own box, level with its checkmark; the unused fourth checkmark leaves.
    points = sorted((shape for shape in slide.shapes if shape.has_text_frame and shape.text in bullets),
                    key=lambda shape: shape.top)
    assert [shape.text for shape in points] == bullets
    for point, check in zip(points, checks):
        line_middle = point.top + point.text_frame.margin_top + Pt(20) * 1.2 / 2
        assert abs(line_middle - (check.top + check.height / 2)) < Inches(0.12)
        assert point.left == body.left and point.width == body.width
    assert {check.shape_id for check in checks[:3]} <= set(shapes) and checks[3].shape_id not in shapes
    # The caption is not a fifth point: it had no text left and leaves.
    assert caption.shape_id not in shapes
    assert metrics["written_slides"][0]["bullets"] == bullets


def test_a_list_whose_paragraphs_match_its_markers_keeps_its_paragraphs(tmp_path):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    textbox(slide, 0.8, 0.5, 11.5, 0.9, "Заголовок — ключевая мысль слайда", 32)
    points = textbox(slide, 1.2, 2, 8, 2, "Первый пункт", 20)
    points.text_frame.add_paragraph().text = "Второй пункт"
    for number in range(2):
        slide.shapes.add_picture(png(), Inches(0.8), Inches(2.1 + number * 0.5), Inches(0.2), Inches(0.2))
    source = tmp_path / "template.pptx"
    deck.save(source)
    _, regions = strict_regions(inspect_template(source).compositions[0], deck.slide_width, deck.slide_height)
    assert [region.kind for region in regions] == ["list_line", "list_line"]


def test_a_heading_that_does_not_wrap_grows_only_up_to_its_neighbour(tmp_path):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    textbox(slide, 0.8, 0.55, 1.6, 0.8, "Кейс", 36, wrap=False)
    slide.shapes.add_picture(png(), Inches(7.0), Inches(0.5), Inches(1.5), Inches(1))
    source = tmp_path / "template.pptx"
    deck.save(source)
    slot = inspect_template(source).compositions[0].slots[0]
    assert Inches(5.5) < slot.grow_width < Inches(6.2)


def test_worker_fills_every_checkmark_and_writes_words_in_full():
    fact = Fact("f1", "Компания выбирает модели по задаче и риску.", "brief.txt", "1",
                "Компания выбирает модели по задаче и риску.")
    answers = [
        {"title": "Упр.", "bullets": ["Выбор облачной, локальной или специализированной модели"], "fact_ids": ["f1"]},
        {"title": "Управление моделями", "bullets": ["Облачные модели", "Локальные модели", "Выбор по задаче",
                                                     "Учёт риска"], "fact_ids": ["f1"]},
    ]
    payloads, schemas = [], []

    class Worker:
        def complete_json(self, system, user, **kwargs):
            payloads.append(json.loads(user))
            schemas.append(kwargs["schema"])
            return answers[len(payloads) - 1]

    slots = {"title": 34, "texts": [27, 27, 27, 27, 53], "paragraphs": [1] * 5,
             "kinds": ["marked_item"] * 4 + ["text"], "min_texts": 4}
    slide = SlideContent("slide_1", "Модели", [fact.excerpt], ["f1"])
    refined = pipeline._refine_slide(Worker(), slide, {"f1": fact}, ["Модели"], time.monotonic() + 60, "", slots)
    assert schemas[0]["properties"]["bullets"]["minItems"] == 4
    feedback = payloads[1]["validation_feedback"]
    assert "Упр." in feedback and "не сокращай" in feedback.casefold()
    assert refined.title == "Управление моделями" and len(refined.bullets) == 4


def test_visual_theme_finds_charts_and_figures_only_in_the_approved_text():
    from engine.visuals import chart_series, stat_parts, visual_options
    shares = ["Облачные модели — 45%", "Локальные модели — 30%", "Специализированные — 25%"]
    assert chart_series(shares).values == (45.0, 30.0, 25.0) and chart_series(shares).unit == "%"
    years = chart_series(["Рост по годам: 2023 — 12%, 2024 — 18%, 2025 — 25%"])
    assert years.categories == ("2023", "2024", "2025") and years.values == (12.0, 18.0, 25.0)
    mixed = ["Выручка выросла на 40%", "Клиентов стало 120 тыс.", "Затраты снизились на 15 млн руб."]
    assert chart_series(mixed) is None and [figure for figure, _ in stat_parts(mixed)] == ["40%", "120 тыс.", "15 млн руб."]
    # Years and list numbers are not figures: such points stay tiles.
    assert visual_options(["В 2024 году запущен сервис", "1. Команда выросла"]) == ["modules", "flow"]


def test_cutting_never_leaves_a_figure_without_its_unit():
    from engine.renderer import fit_chars
    assert fit_chars("Затраты снизились на 15 млн руб. за год", 22) == "Затраты снизились"


def test_chart_and_figure_tiles_are_editable_and_keep_the_points(tmp_path):
    source, template, heading, body, checks, caption = checkmark_template(tmp_path)
    shares = ["Облачные модели — 45%", "Локальные модели — 30%", "Специализированные — 25%"]
    results = ["Выручка выросла на 40%", "Клиентов стало 120 тыс.", "Затраты снизились на 15 млн"]
    content = [SlideContent("slide_1", "Структура использования", shares, ["f1"]),
               SlideContent("slide_2", "Результаты за год", results, ["f1"])]
    target = tmp_path / "visual.pptx"
    metrics = render_variant(source, template, content, 2, target, strict=True, plan=[0, 0],
                             infographic_layouts={"slide_1": "chart", "slide_2": "stats"})
    assert metrics["infographic_slides"] == [1, 2]
    chart_slide, stats_slide = Presentation(target).slides
    charts = [shape for shape in chart_slide.shapes if shape.has_chart]
    assert len(charts) == 1 and list(charts[0].chart.series[0].values) == [45.0, 30.0, 25.0]
    assert charts[0].chart.value_axis.axis_title.text_frame.text == "%"
    # The diagram replaced the marked list: its checkmarks leave, the heading is the template's.
    assert not [shape for shape in chart_slide.shapes if shape.shape_type == 13]
    assert any(shape.text == "Структура использования" for shape in chart_slide.shapes if shape.has_text_frame)
    texts = [shape.text for shape in stats_slide.shapes if shape.has_text_frame]
    assert {"40%", "120 тыс.", "15 млн"} <= set(texts) and all(point in texts for point in results)
