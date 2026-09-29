from __future__ import annotations

from decimal import Decimal

import pytest
from pptx import Presentation
from pptx.util import Inches

from engine.native_data import add_native_data_visuals, parse_numeric_csv


def test_csv_becomes_editable_chart_and_table_within_slide(tmp_path):
    source = tmp_path / "sales.csv"
    source.write_text(
        "\ufeffQuarter;Revenue;Cost\nQ1;12,5;7\nQ2;15,0;-3\n",
        encoding="utf-8",
    )
    data = parse_numeric_csv(source)
    assert data.categories == ("Q1", "Q2")
    assert data.series[0].values == (Decimal("12.5"), Decimal("15.0"))

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shapes = add_native_data_visuals(
        slide, data,
        slide_width=presentation.slide_width,
        slide_height=presentation.slide_height,
    )
    assert shapes.table.has_table
    assert shapes.chart.has_chart
    for shape in (shapes.table, shapes.chart):
        assert shape.left >= 0 and shape.top >= 0
        assert shape.left + shape.width <= presentation.slide_width
        assert shape.top + shape.height <= presentation.slide_height

    output = tmp_path / "native-data.pptx"
    presentation.save(output)
    reopened = Presentation(output)
    table = next(shape.table for shape in reopened.slides[0].shapes if shape.has_table)
    chart = next(shape.chart for shape in reopened.slides[0].shapes if shape.has_chart)
    assert table.cell(0, 1).text == "Revenue"
    assert table.cell(1, 1).text == "12,5"
    assert tuple(chart.series[0].values) == (12.5, 15.0)
    assert tuple(chart.series[1].values) == (7.0, -3.0)


@pytest.mark.parametrize(
    "body, error",
    [
        ("Category,Amount\\nQ1,NaN\\n", "plain number"),
        ("Category,Amount\\n=SUM(1:2),12\\n", "formula character"),
        ("Category,Amount\\nQ1,12\\nQ2\\n", "columns"),
        ("Category,Amount\\n" + "".join(f"Q{i},{i}\\n" for i in range(13)), "too many rows"),
    ],
)
def test_bad_csv_is_rejected(tmp_path, body, error):
    source = tmp_path / "bad.csv"
    source.write_text(body.replace("\\n", "\n"), encoding="utf-8")
    with pytest.raises(ValueError, match=error):
        parse_numeric_csv(source)


def test_out_of_bounds_area_fails_before_modifying_slide(tmp_path):
    source = tmp_path / "sales.csv"
    source.write_text("Category,Amount\nQ1,12\n", encoding="utf-8")
    data = parse_numeric_csv(source)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    existing = len(slide.shapes)
    with pytest.raises(ValueError, match="exceed"):
        add_native_data_visuals(
            slide, data,
            slide_width=presentation.slide_width,
            slide_height=presentation.slide_height,
            bounds=(Inches(2), Inches(2), presentation.slide_width, Inches(4)),
        )
    assert len(slide.shapes) == existing
