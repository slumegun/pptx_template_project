from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from engine import pipeline


def test_edit_changes_only_selected_slide(monkeypatch, tmp_path):
    source = Presentation()
    for number in range(1, 4):
        slide = source.slides.add_slide(source.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        box.text = f"Слайд {number}"
    base_path = tmp_path / "base.pptx"
    source.save(base_path)
    target_shape_id = source.slides[1].shapes[0].shape_id

    class FakeGateway:
        def client(self, role):
            assert role == "slide_repair"
            return self

        def complete_json(self, system, user, **kwargs):
            assert "Сократи заголовок" in user
            return {"edits": [{"shape_id": target_shape_id, "text": "Новый заголовок"}]}

    def fake_pdf(pptx_path, output_dir):
        result = output_dir / "edited.pdf"
        result.write_bytes(b"%PDF-1.4\n%%EOF")
        return result

    def fake_previews(pdf_path, output_dir):
        output_dir.mkdir()
        result = output_dir / "slide-1.png"
        result.write_bytes(b"\x89PNG\r\n\x1a\n")
        return [result]

    def fake_html(pptx_path, previews, html_path):
        html_path.write_text("<html></html>", encoding="utf-8")
        return html_path

    monkeypatch.setattr(pipeline, "ModelGateway", FakeGateway)
    monkeypatch.setattr(pipeline, "export_pdf", fake_pdf)
    monkeypatch.setattr(pipeline, "export_previews", fake_previews)
    monkeypatch.setattr(pipeline, "export_html", fake_html)
    monkeypatch.setattr(pipeline, "extract_facts", lambda brief, paths: [])
    monkeypatch.setattr(pipeline, "audit_pptx", lambda *args: [])
    monkeypatch.setattr(pipeline, "source_texts_from_template", lambda path: [])
    result = pipeline.edit_slide(base_path, base_path, "Тема", 3, 2, "Сократи заголовок", tmp_path / "out")
    edited = Presentation(result[0].pptx_path)
    assert [slide.shapes[0].text for slide in edited.slides] == [
        "Слайд 1", "Новый заголовок", "Слайд 3",
    ]
    assert result[0].pdf_path.is_file()
    assert result[0].html_path.is_file()
