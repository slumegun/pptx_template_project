"""Extract traceable statements from user supplied content."""

from __future__ import annotations

import csv
import hashlib
import json
import re
from pathlib import Path

from .models import Fact


MAX_SOURCE_BYTES = 20 * 1024 * 1024


def _fact(text: str, source: str, location: str) -> Fact:
    text = " ".join(text.split()).strip()
    digest = hashlib.sha256((source + ":" + location + ":" + text).encode("utf-8")).hexdigest()[:12]
    return Fact("fact_" + digest, text, source, location, text)


def _split_statements(text: str) -> list[str]:
    statements: list[str] = []
    for line in text.splitlines():
        if line.lstrip().startswith("#"):
            continue
        # Strip list markers only. A leading year, percentage or negative value
        # belongs to the source and must survive extraction unchanged.
        cleaned = re.sub(r"^\s*(?:[-*•]\s+|\d{1,2}[.)]\s+)", "", line).strip()
        if not cleaned:
            continue
        for part in re.split(r"(?<=[.!?])\s+(?=[А-ЯЁA-Z0-9])", cleaned):
            part = part.strip()
            if len(part) >= 12:
                statements.append(part)
    return statements


def _read_text_file(path: Path) -> list[tuple[str, str]]:
    suffix = path.suffix.lower()
    if path.stat().st_size > MAX_SOURCE_BYTES:
        raise ValueError(f"Content source is too large: {path.name}")
    if suffix in {".txt", ".md"}:
        text = path.read_text(encoding="utf-8-sig")
        return [(f"line {number}", line) for number, line in enumerate(text.splitlines(), 1)]
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            sample = stream.read(8192)
            stream.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            except csv.Error as exc:
                raise ValueError(f"Не удалось определить разделитель CSV: {path.name}") from exc
            reader = csv.DictReader(stream, dialect=dialect)
            records = []
            for number, row in enumerate(reader, 2):
                if None in row or any(value is None for value in row.values()):
                    raise ValueError(f"Неодинаковое число колонок CSV: {path.name}, строка {number}")
                records.append((f"row {number}", "; ".join(f"{key}: {value}" for key, value in row.items() if value)))
            return records
    if suffix == ".json":
        data = json.loads(path.read_text(encoding="utf-8-sig"))
        if isinstance(data, dict):
            return [(str(key), json.dumps(value, ensure_ascii=False)) for key, value in data.items()]
        if isinstance(data, list):
            return [(f"item {number}", json.dumps(value, ensure_ascii=False)) for number, value in enumerate(data, 1)]
        return [("root", str(data))]
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        records = [(f"page {number}", page.extract_text() or "") for number, page in enumerate(reader.pages, 1)]
        if not any(text.strip() for _, text in records):
            raise ValueError(f"PDF {path.name} не содержит текста. Для скана нужен текстовый файл или OCR.")
        return records
    if suffix == ".docx":
        from docx import Document

        document = Document(str(path))
        records = [(f"paragraph {number}", paragraph.text) for number, paragraph in enumerate(document.paragraphs, 1)]
        for table_index, table in enumerate(document.tables, 1):
            for row_index, row in enumerate(table.rows, 1):
                records.append((f"table {table_index}, row {row_index}", " | ".join(cell.text for cell in row.cells)))
        return records
    raise ValueError(f"Unsupported content source: {path.name}")


def extract_facts(brief: str, content_paths: list[Path] | None = None) -> list[Fact]:
    records: list[tuple[str, str, str]] = []
    for number, line in enumerate(brief.splitlines(), 1):
        records.append(("brief", f"line {number}", line))
    for path in content_paths or []:
        path = Path(path)
        for location, text in _read_text_file(path):
            records.append((path.name, location, text))
    facts: list[Fact] = []
    seen: set[str] = set()
    for source, location, raw in records:
        for statement in _split_statements(raw):
            # User instructions guide the planner but are not product evidence.
            if source == "brief" and (statement.endswith(":") or re.match(
                r"(?i)^(?:подготовь|создай|сделай|покажи|используй|сохрани|размести|переформулируй|не\s+(?:добавляй|переноси)|нужно\s+\d+)\b", statement)):
                continue
            normalized = statement.casefold()
            if normalized in seen:
                continue
            seen.add(normalized)
            facts.append(_fact(statement, source, location))
    return facts
