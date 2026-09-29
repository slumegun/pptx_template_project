"""Analyze a newly uploaded PPTX without using any design from another file."""

from __future__ import annotations

import hashlib
import json
import math
import re
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .models import Composition, PreparedTemplate, SCHEMA_VERSION, SlideSlot


MAX_ARCHIVE_BYTES = 100 * 1024 * 1024
MAX_UNPACKED_BYTES = 1024 * 1024 * 1024


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def validate_pptx(path: Path) -> None:
    if not path.is_file():
        raise ValueError("PPTX file does not exist")
    if path.suffix.lower() != ".pptx":
        raise ValueError("Template must be a .pptx file")
    if path.stat().st_size > MAX_ARCHIVE_BYTES:
        raise ValueError("Template is too large (100 MiB limit)")
    if not zipfile.is_zipfile(path):
        raise ValueError("Template is not a valid PowerPoint ZIP package")
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        if "[Content_Types].xml" not in names or "ppt/presentation.xml" not in names:
            raise ValueError("PowerPoint package is incomplete")
        total = 0
        for info in archive.infolist():
            normalized = info.filename.replace("\\", "/")
            if normalized.startswith("/") or ".." in normalized.split("/"):
                raise ValueError("PowerPoint package contains an unsafe path")
            if info.flag_bits & 1:
                raise ValueError("Encrypted PowerPoint packages are unsupported")
            total += info.file_size
            if total > MAX_UNPACKED_BYTES:
                raise ValueError("PowerPoint package expands beyond 1 GiB")


def _font_size(shape) -> int | None:
    if not shape.has_text_frame:
        return None
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.font.size is not None:
                return int(run.font.size)
    return None


# A bare label ("Заголовок", "Текст") is a stub, not a measure of the text a region holds.
_STUB_LABELS = {"заголовок", "подзаголовок", "текст", "описание", "подпись", "пункт", "title", "subtitle", "text",
                "click to add title", "click to add text", "заголовок слайда", "текст слайда"}


# Ordinary Russian prose with its spaces: the average character a limit in characters counts.
_PROSE = "Съешь же ещё этих мягких французских булок, да выпей чаю. Сервис готовит отчёт для команды."


def char_width(size: float, family: str | None = "Arial", bold: bool = False) -> float:
    """Average width (pt) of a character of prose in the template's own typeface and size.

    A wide bold face takes a fifth more room than a narrow one; a fixed 0.55 em per character
    lets a heading in a wide face overflow its line.
    """
    from .typography import _face
    try:
        return max(size * 0.4, _face(family or "Arial", float(size), bool(bold), False).getlength(_PROSE) / 4 / len(_PROSE))
    except (OSError, RuntimeError, ValueError):
        return size * 0.55


def _geometric_capacity(shape, size: float, outer_width: int | None = None, family: str | None = None,
                        bold: bool = False) -> int:
    """Characters a text object holds at its size: the typeface's average character, 1.2 lines."""
    frame = shape.text_frame
    outer = outer_width or shape.width
    width = max(1.0, (outer - (frame.margin_left or 0) - (frame.margin_right or 0)) / 12700)
    height = max(1.0, (shape.height - (frame.margin_top or 0) - (frame.margin_bottom or 0)) / 12700)
    return max(1, int(int(width / char_width(size, family, bold)) * max(1, int(height / (size * 1.2))) * 0.9))


def grow_width(shape, slide_width: int | None, align: str | None) -> int:
    """Width a frame that does not wrap may take: its line grows up to the slide's margin.

    A heading like "Кейс" in a narrow auto-size box is a heading all the same; new text there
    runs on to the right (or both ways when centred), not into a word cut to four letters. The
    line stops before the next object at its height: a column never grows into its neighbour.
    """
    if not slide_width or shape.text_frame.word_wrap is not False:
        return 0
    margin = max(int(slide_width * 0.04), min(shape.left, slide_width - shape.left - shape.width))
    lower, upper = margin, slide_width - margin
    try:
        neighbours = list(shape.part.slide.shapes)
    except AttributeError:
        neighbours = []
    gap = int(0.15 * 914400)
    for other in neighbours:
        if other.shape_id == shape.shape_id or other.width <= 0 or other.height <= 0:
            continue
        # A backdrop or a panel the frame sits on is not a neighbour.
        if other.left <= shape.left and other.left + other.width >= shape.left + shape.width:
            continue
        if other.top >= shape.top + shape.height or other.top + other.height <= shape.top:
            continue
        if other.left >= shape.left + shape.width:
            upper = min(upper, other.left - gap)
        elif other.left + other.width <= shape.left:
            lower = max(lower, other.left + other.width + gap)
    if align == "center":
        center = shape.left + shape.width // 2
        room = 2 * min(center - lower, upper - center)
    elif align == "right":
        room = shape.left + shape.width - lower
    else:
        room = upper - shape.left
    return max(int(shape.width), int(room))


def slot_details(shape, slide_width: int | None = None) -> dict:
    """Exact template data of one text object: typeface, size, weight, colour, alignment, volume.

    Everything is resolved the way PowerPoint draws the text: the run, the paragraph, the
    shape's list styles, the layout and master placeholders, the master text styles and the
    theme (colour scheme, colour map, theme fonts). slide_width lets a frame that does not wrap
    its lines count the room up to the slide's margin.
    """
    from types import SimpleNamespace
    from .design_audit import _Theme, _hex, _text_color
    from .fonts import effective_font

    frame = shape.text_frame
    filled = [paragraph for paragraph in frame.paragraphs if paragraph.text.strip()]
    paragraph = filled[0] if filled else frame.paragraphs[0]
    run = next((item for item in paragraph.runs if item.text.strip()), None)
    family, size, bold = effective_font(shape, paragraph, run)
    color = None
    try:
        theme = _Theme(shape.part.slide.slide_layout.slide_master)
        rgb = _text_color(shape, paragraph, run if run is not None else SimpleNamespace(_r=SimpleNamespace(rPr=None)), theme)
        color = _hex(rgb) if rgb is not None else None
    except (AttributeError, KeyError, ValueError, IndexError):
        pass
    alignment = paragraph.alignment
    align = alignment.name.lower() if alignment is not None and hasattr(alignment, "name") else None
    grow = grow_width(shape, slide_width, align)
    outer = grow or shape.width
    text = " ".join(shape.text.split())
    stub = not text or text.casefold().strip(" .:") in _STUB_LABELS
    if stub:
        paragraph_chars = [_geometric_capacity(shape, float(size), outer, family, bool(bold))]
    else:
        # A paragraph may take as many lines as the template's own text took there, at the
        # template's size: a one-line sample allows one full line, never a second one.
        from .typography import text_height
        inner = max(1.0, (outer - (frame.margin_left or 0) - (frame.margin_right or 0)) / 12700)
        per_line = max(1, int(inner / char_width(float(size), family, bool(bold))))
        one_line = text_height(["Ш"], inner, float(size), family or "Arial", bool(bold))
        paragraph_chars = []
        for item in filled:
            sample = " ".join(item.text.split())
            lines = max(1, round(text_height([sample], inner, float(size), family or "Arial", bool(bold)) / one_line))
            paragraph_chars.append(max(len(sample), int(lines * per_line * 0.85)))
    def bulleted(item) -> bool:
        # A bullet of the paragraph itself or of its list style, unless switched off nearer.
        from .fonts import _style_chain
        chain = ([item._p.pPr] if item._p.pPr is not None else []) + _style_chain(shape, item.level + 1)
        for props in chain:
            names = {child.tag.rsplit("}", 1)[-1] for child in props}
            if "buNone" in names:
                return False
            if names & {"buChar", "buAutoNum", "buBlip"}:
                return True
        return False

    return {
        "bulleted": len(filled) > 1 and all(bulleted(item) for item in filled),
        "font_family": family,
        "font_pt": round(float(size), 2),
        "bold": bool(bold),
        "italic": bool(run is not None and run.font.italic),
        "color": color,
        "align": align,
        "paragraphs": max(1, len(filled)),
        "max_chars": sum(paragraph_chars) + len(paragraph_chars) - 1,
        "chars_source": "geometry" if stub else "sample_lines",
        "paragraph_chars": paragraph_chars,
        "grow_width": grow if grow > shape.width else 0,
    }


_MARKER_TEXT = re.compile(r"[\W\d_]{0,3}")


def _letters(shape) -> int:
    return sum(character.isalpha() for character in shape.text) if shape.has_text_frame else 0


def list_items(slide, shape, details: dict, slide_height: int) -> list[dict]:
    """The items a column of markers beside a text object defines, one per marker.

    A template often draws a list as one text box with checkmarks, dots, icons or numbers placed
    beside it as separate objects. Each marker is a place for one short point written level with
    it; the text a marker's line holds is limited by the distance to the next marker. A list whose
    own paragraphs already match the markers keeps its paragraphs (they are the places).
    """
    from .fonts import effective_font, paragraph_spacing
    from .typography import text_height

    frame = shape.text_frame
    size = float(details.get("font_pt") or 18.0)
    line = size * 1.2 * 12700
    left_window, right_window = shape.left - int(1.2 * 914400), shape.left + min(shape.width // 4, int(0.8 * 914400))
    top, bottom = shape.top - line / 2, shape.top + shape.height + line / 2
    others = [item for item in slide.shapes if item.shape_id != shape.shape_id and _letters(item) >= 3]
    candidates = []
    for item in slide.shapes:
        if item.shape_id == shape.shape_id or item.is_placeholder or item.width <= 0 or item.height <= 0:
            continue
        if item.has_text_frame and not _MARKER_TEXT.fullmatch(item.text.strip()):
            continue
        if item.shape_type == MSO_SHAPE_TYPE.GROUP and _letters_in_group(item):
            continue
        if max(item.width, item.height) > max(int(0.9 * 914400), 3 * line) or min(item.width, item.height) < 36000:
            continue
        center_x, center_y = item.left + item.width // 2, item.top + item.height // 2
        if not (left_window <= center_x <= right_window and top <= center_y <= bottom):
            continue
        if item.left + item.width > shape.left + shape.width * 0.3:
            continue
        # A marker right before another text object at its height belongs to that text.
        if any(other.top <= center_y <= other.top + other.height
               and item.left + item.width - 91440 <= other.left < shape.left for other in others):
            continue
        candidates.append(item)
    best: list = []
    for item in candidates:
        column = [other for other in candidates
                  if abs((other.left + other.width // 2) - (item.left + item.width // 2)) <= max(item.width // 2, 73152)
                  and 0.6 <= other.width / item.width <= 1.6 and 0.6 <= other.height / item.height <= 1.6]
        distinct: list = []
        for other in sorted(column, key=lambda value: value.top):
            if not distinct or other.top + other.height // 2 - (distinct[-1].top + distinct[-1].height // 2) >= other.height * 0.8:
                distinct.append(other)
        if len(distinct) > len(best):
            best = distinct
    filled = [paragraph for paragraph in frame.paragraphs if paragraph.text.strip()]
    if len(best) < 2 or len(filled) == len(best):
        return []
    paragraph = filled[0] if filled else frame.paragraphs[0]
    family, _, bold = effective_font(shape, paragraph, next((run for run in paragraph.runs if run.text.strip()), None))
    factor, _, indent = paragraph_spacing(shape, paragraph)
    inner = max(20.0, (shape.width - (frame.margin_left or 0) - (frame.margin_right or 0)) / 12700 - indent)
    one_line = text_height(["Ш"], inner, size, family or "Arial", bool(bold), False, factor) * 12700
    per_line = max(1, int(inner / char_width(size, family, bool(bold))))
    centers = [marker.top + marker.height // 2 for marker in best]
    steps = [after - before for before, after in zip(centers, centers[1:])]
    items = []
    for number, (marker, center) in enumerate(zip(best, centers)):
        step = steps[number] if number < len(steps) else steps[-1]
        # The first line of the item sits level with its marker; the item ends where the next begins.
        item_top = int(center - one_line / 2 - (frame.margin_top or 0))
        height = int(min(step, slide_height - int(0.3 * 914400) - item_top))
        lines = max(1, int((step - 12700 * 2) // one_line))
        items.append({"top": item_top, "height": max(height, int(one_line + (frame.margin_top or 0))),
                      "max_chars": max(8, int(lines * per_line * 0.85)), "marker_id": marker.shape_id})
    return items


def _letters_in_group(group) -> int:
    return sum(_letters_in_group(item) if item.shape_type == MSO_SHAPE_TYPE.GROUP else _letters(item)
               for item in group.shapes)


def enrich_slots(template: PreparedTemplate, template_path: Path) -> PreparedTemplate:
    """Recompute every slot's exact data from the PPTX; the analysis (roles, archetypes) stays.

    The data is deterministic and cheap, so a template prepared before it existed gets it
    at generation time without a new paid analysis.
    """
    presentation = Presentation(str(template_path))
    slides = list(presentation.slides)
    for composition in template.compositions:
        if not 0 <= composition.source_slide_index < len(slides):
            continue
        slide = slides[composition.source_slide_index]
        shapes = {shape.shape_id: shape for shape in slide.shapes}
        for slot in composition.slots:
            shape = shapes.get(slot.shape_id)
            if shape is not None and shape.has_text_frame:
                details = slot_details(shape, presentation.slide_width)
                details["list_items"] = list_items(slide, shape, details, presentation.slide_height)
                for key, value in details.items():
                    setattr(slot, key, value)
    return template


def _fingerprint(slots: list[SlideSlot], width: int, height: int) -> str:
    # Quantized geometry distinguishes real compositions, not cosmetic colour changes.
    bins = []
    for slot in slots[:8]:
        bins.append(
            (
                min(4, int(5 * slot.x / max(width, 1))),
                min(4, int(5 * slot.y / max(height, 1))),
                min(4, int(5 * slot.width / max(width, 1))),
            )
        )
    return json.dumps(sorted(bins), separators=(",", ":"))


def inspect_template(path: Path) -> PreparedTemplate:
    path = Path(path)
    validate_pptx(path)
    presentation = Presentation(path)
    if not presentation.slides:
        raise ValueError("Template contains no slides")
    compositions: list[Composition] = []
    warnings: list[str] = []
    slide_area = presentation.slide_width * presentation.slide_height
    layouts = list(presentation.slide_layouts)
    for index, slide in enumerate(presentation.slides):
        slots: list[SlideSlot] = []
        picture_area = 0.0
        charts = 0
        tables = 0
        for shape in slide.shapes:
            if shape.has_text_frame and shape.width > 0 and shape.height > 0:
                text = shape.text.strip()
                if text or shape.is_placeholder:
                    details = slot_details(shape, presentation.slide_width)
                    details["list_items"] = list_items(slide, shape, details, presentation.slide_height)
                    slots.append(
                        SlideSlot(
                            shape_id=shape.shape_id,
                            x=int(shape.left),
                            y=int(shape.top),
                            width=int(shape.width),
                            height=int(shape.height),
                            text=text[:240],
                            font_size=_font_size(shape),
                            **details,
                        )
                    )
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                picture_area += max(0, shape.width) * max(0, shape.height)
            if shape.has_chart:
                charts += 1
            if shape.has_table:
                tables += 1
        try:
            layout_index = layouts.index(slide.slide_layout)
        except ValueError:
            layout_index = 0
        ratio = min(1.0, picture_area / max(slide_area, 1))
        # Dense infographic and full-image slides need semantic handling before reuse.
        score = (
            (6.0 if 2 <= len(slots) <= 8 else 0.0)
            + min(len(slots), 8) * 0.25
            - max(0, len(slide.shapes) - 20) * 0.11
            - ratio * 8
            - charts * 1.5
            - tables * 0.5
        )
        compositions.append(
            Composition(
                source_slide_index=index,
                layout_index=layout_index,
                fingerprint=_fingerprint(slots, presentation.slide_width, presentation.slide_height),
                slots=slots,
                shape_count=len(slide.shapes),
                picture_area_ratio=round(ratio, 4),
                chart_count=charts,
                table_count=tables,
                score=round(score, 3),
            )
        )
    if not any(len(item.slots) >= 2 for item in compositions):
        warnings.append("В шаблоне нет слайда с двумя текстовыми областями; текстовые блоки будут созданы программно.")
    if not any(item.picture_area_ratio < 0.3 for item in compositions):
        warnings.append("Все композиции насыщены изображениями; потребуется визуальная проверка на старое содержание.")
    return PreparedTemplate(
        schema_version=SCHEMA_VERSION,
        template_sha256=sha256_file(path),
        source_name=path.name,
        width=int(presentation.slide_width),
        height=int(presentation.slide_height),
        slide_count=len(presentation.slides),
        layout_count=len(layouts),
        compositions=compositions,
        warnings=warnings,
    )


def save_prepared(template: PreparedTemplate, path: Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(template.to_dict(), ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def load_prepared(path: Path, template_path: Path) -> PreparedTemplate:
    payload = json.loads(Path(path).read_text(encoding="utf-8"))
    result = PreparedTemplate.from_dict(payload)
    if result.template_sha256 != sha256_file(Path(template_path)):
        raise ValueError("Prepared analysis belongs to a different PPTX")
    return result


OBJECT_ROLES = {"fixed_brand", "replaceable", "reusable_asset", "remove", "unresolved"}
ARCHETYPES = {"title", "content", "comparison", "data", "closing", "unknown"}


def verify_imported(prepared_path: Path, template_path: Path) -> PreparedTemplate:
    """Check an analysis from a portable package against its own PPTX.

    The package comes from another installation and is untrusted: every slide
    index, object id and role must match the file, or it is rejected.
    """
    try:
        prepared = load_prepared(prepared_path, template_path)
    except (KeyError, TypeError, json.JSONDecodeError) as error:
        raise ValueError("Template analysis has an invalid structure") from error
    if prepared.analysis_mode != "vision_model":
        raise ValueError("Template analysis was not completed")
    actual = inspect_template(template_path)
    if (prepared.width, prepared.height, prepared.slide_count) != (actual.width, actual.height, actual.slide_count):
        raise ValueError("Template analysis does not match the PPTX geometry")
    indices = [item.source_slide_index for item in prepared.compositions]
    if sorted(indices) != list(range(actual.slide_count)):
        raise ValueError("Template analysis does not describe every slide exactly once")
    deck = Presentation(template_path)
    analysed = {item.source_slide_index: item for item in prepared.compositions}
    for composition in actual.compositions:
        imported = analysed[composition.source_slide_index]
        shape_ids = {str(shape.shape_id) for shape in deck.slides[composition.source_slide_index].shapes}
        if (not isinstance(imported.object_roles, dict)
                or any(key not in shape_ids or role not in OBJECT_ROLES
                       for key, role in imported.object_roles.items())):
            raise ValueError("Template analysis refers to unknown objects or roles")
        if imported.archetype not in ARCHETYPES:
            raise ValueError("Template analysis contains an unknown slide type")
        notes = imported.analysis_notes if isinstance(imported.analysis_notes, list) else []
        # Geometry is recomputed from the file; only the checked model semantics are kept.
        composition.object_roles = dict(imported.object_roles)
        composition.archetype = imported.archetype
        composition.analysis_notes = [note[:600] for note in notes[:8] if isinstance(note, str)]
        composition.analysis_complete = imported.analysis_complete is True
    actual.analysis_mode = "vision_model"
    actual.analysis_model = prepared.analysis_model if isinstance(prepared.analysis_model, str) else None
    actual.warnings = [item[:500] for item in prepared.warnings if isinstance(item, str)][:200] \
        if isinstance(prepared.warnings, list) else []
    return actual
