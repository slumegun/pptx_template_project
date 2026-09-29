"""Preparation, model planning, native rendering, audit and export."""

from __future__ import annotations

import json
import hashlib
import math
import os
import re
import time
import tempfile
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import asdict
from pathlib import Path
from typing import Any, Callable

from pptx import Presentation
from pptx.util import Pt

from .audit import audit_pptx, duplicate_content_issues, font_issues, source_texts_from_template
from .design_audit import design_issues
from .fields import mark_fields
from .fonts import ensure_fonts, template_families, text_block_height, with_embedded_fonts
from .content import extract_facts
from .export import export_html, export_pdf, export_previews
from .ingest import enrich_slots, inspect_template, load_prepared, save_prepared
from .models import Fact, GeneratedVariant, SlideContent, generation_budget_seconds
from .native_data import parse_numeric_csv
from .provider import ModelClient, ModelGateway, ModelProviderError
from .prompt_text import prompt_text
from .renderer import (GOALS, _point_slots, _template_typography, choose_compositions, fit_slide, render_variant,
                       slot_budget)
from .pagination import paginate_slides
from .rendered_audit import audit_pdf
from .visuals import visual_options
from .swarm import analyze_template, contact_sheet, critique_content, critique_prompt, critique_visual_slide, critique_visual_variant, remaining_seconds, validate_slide

Progress = Callable[[str, int | None], None]


class InsufficientContent(ValueError):
    pass


def _progress(callback: Progress | None, stage: str, percent: int | None) -> None:
    if callback:
        callback(stage, percent)


def _visual_references(template_path: Path, template, output_dir: Path, deadline: float | None = None) -> list[Path]:
    """Render immutable source references during preparation and reuse them for audit."""
    cache_root = os.getenv("MODEL_TEMPLATE_CACHE_DIR")
    destination = (Path(cache_root) / "visual-references-v1" / template.template_sha256
                   if cache_root else output_dir / "visual_reference")
    def ready():
        paths = sorted((destination / "slides").glob("*.png"), key=lambda p: int(p.stem.rsplit("-", 1)[-1]))
        return paths if len(paths) == len(template.compositions) and all(p.stat().st_size for p in paths) else []
    existing = ready()
    if existing:
        return existing
    destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="references-", dir=destination.parent) as scratch:
        staging = Path(scratch) / "ready"
        staging.mkdir()
        source = with_embedded_fonts(template_path, Path(scratch) / Path(template_path).name)
        pdf = export_pdf(source, staging, timeout=remaining_seconds(deadline))
        export_previews(pdf, staging / "slides", dpi=90, timeout=remaining_seconds(deadline))
        try:
            staging.rename(destination)
        except FileExistsError:
            if not ready():
                raise ValueError("Source preview cache is incomplete")
    paths = ready()
    if not paths:
        raise ValueError("Source preview count does not match the template")
    return paths


def prepare(template_path: Path, output_dir: Path) -> Path:
    """Analyze one uploaded PPTX before the five-minute generation clock starts."""
    template_path = Path(template_path)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    template = inspect_template(template_path)
    gateway = ModelGateway()
    client = gateway.client("template_analyst")
    if not client.vision_enabled:
        raise ModelProviderError("Set OPENROUTER_API_KEY before preparing a template")
    if client.vision_enabled:
        cache_path = None
        cache_root = os.getenv("MODEL_TEMPLATE_CACHE_DIR")
        if cache_root:
            identity = json.dumps({"parser": 3, "role": "template_analyst", "sha256": template.template_sha256,
                "model": client.vision_model, "endpoint": client.vision_base_url,
                "prompt": prompt_text("template_analyst"), "temperature": client.temperature, "reasoning": client.reasoning_effort}, sort_keys=True).encode("utf-8")
            cache_path = Path(cache_root) / (hashlib.sha256(identity).hexdigest() + ".json")
            if cache_path.is_file():
                try:
                    template = load_prepared(cache_path, template_path)
                except (ValueError, KeyError, TypeError, OSError):
                    pass
                else:
                    if template.analysis_mode == "vision_model":
                        return save_prepared(template, output_dir / "template_ir.json")
        def checkpoint(value):
            if cache_path is not None:
                temporary = cache_path.with_suffix(".tmp")
                save_prepared(value, temporary)
                temporary.replace(cache_path)
        previews = _visual_references(template_path, template, output_dir)
        analyze_template(client, template_path, template, previews, checkpoint=checkpoint)
        checkpoint(template)
    _mark_template_fields(template_path, template, output_dir)
    return save_prepared(template, output_dir / "template_ir.json")


def _mark_template_fields(template_path: Path, template, output_dir: Path) -> None:
    """Field Marker: exact free areas of the text fields, measured on the template without text."""
    try:
        mark_fields(template_path, template, output_dir)
    except (RuntimeError, OSError, ValueError) as error:
        # Without the measurement the frames themselves remain the fields.
        template.warnings.append(f"Field Marker: разметка полей не выполнена ({error}); используются рамки шаблона.")


def template_previews(template_path: Path, prepared_path: Path, output_dir: Path) -> list[Path]:
    """Slide images of a prepared template, reused from the analysis cache when present."""
    template = load_prepared(Path(prepared_path), Path(template_path))
    return _visual_references(Path(template_path), template, Path(output_dir))


def _model_plan(client: ModelClient, facts: list[Fact], slide_count: int, brief: str, deadline: float | None = None) -> list[SlideContent]:
    fact_json = json.dumps([{"fact_id": fact.fact_id, "excerpt": fact.excerpt} for fact in facts], ensure_ascii=False, separators=(",", ":"))
    if len(fact_json) + len(brief) > 120000:
        raise InsufficientContent("Промпт и материалы превышают контекст планировщика (120 000 знаков). Сократите ввод.")
    fact_lookup = {fact.fact_id: fact for fact in facts}
    feedback = ""
    for attempt in range(2):
        response = client.complete_json(
            prompt_text("deck_planner"),
            f"Число слайдов: {slide_count}. Полный пользовательский промпт: {brief}"
            + (f"\nПредыдущий план отклонён: {feedback}. Исправь его, не добавляя неподтверждённых возможностей." if feedback else "")
            + f"\nFact registry:\n{fact_json}",
            max_tokens=max(900, min(16000, 300 * slide_count)),
            deadline=deadline,
        )
        try:
            raw = response.get("slides")
            if not isinstance(raw, list) or len(raw) != slide_count:
                raise ModelProviderError("Deck Planner returned the wrong slide count")
            slides: list[SlideContent] = []
            for index, item in enumerate(raw):
                if not isinstance(item, dict):
                    raise ModelProviderError("Deck Planner returned an invalid slide")
                if "bullets" not in item:
                    refs = item.get("fact_ids", [])
                    if not isinstance(refs, list) or not refs or any(not isinstance(ref, str) or ref not in fact_lookup for ref in refs):
                        raise ModelProviderError("Deck Planner referenced an unknown fact")
                    item = {**item, "bullets": [fact_lookup[ref].excerpt for ref in refs]}
                slides.append(validate_slide(item, f"slide_{index+1}", fact_lookup))
            return slides
        except ModelProviderError as error:
            if attempt:
                raise
            feedback = str(error)
    raise ModelProviderError("Deck Planner could not produce a valid plan")


def _budget_violations(slide: SlideContent, budget: dict) -> list[str]:
    """What a written slide exceeds in its template regions (characters, whitespace collapsed)."""
    def length(text: str) -> int:
        return len("\n".join(" ".join(line.split()) for line in text.split("\n") if line.strip()))
    problems = []
    if length(slide.title) > budget["title"]:
        problems.append(f"заголовок {length(slide.title)} знаков при пределе {budget['title']}")
    texts = budget["texts"]
    if texts and len(slide.bullets) > len(texts):
        problems.append(f"{len(slide.bullets)} пунктов при {len(texts)} текстовых местах макета")
    required = min(budget.get("min_texts") or 0, len(texts), 6)
    if len(slide.bullets) < required:
        problems.append(f"{len(slide.bullets)} пунктов, а в списке макета {required} мест с галочками или иконками: "
                        "раздели мысль на отдельные короткие пункты, по одному на каждое место")
    for number, (text, limit) in enumerate(zip(slide.bullets, texts), 1):
        if length(text) > limit:
            problems.append(f"пункт {number}: {length(text)} знаков при пределе {limit}")
    return problems


# Short words a text may end with a dot: common abbreviations, not a word cut to fit.
_KNOWN_ABBREVIATIONS = {"т", "д", "п", "е", "др", "пр", "г", "гг", "в", "вв", "млн", "млрд", "трлн", "тыс", "руб",
                        "коп", "см", "ср", "им", "ул", "стр", "рис", "табл", "шт", "мин", "сек", "ч", "кв", "etc",
                        "vs", "inc", "ltd", "co", "no", "e.g", "i.e"}
_ABBREVIATED = re.compile(r"(?<![\w.])([A-Za-zА-Яа-яЁё]{1,6})\.(?=\s+[a-zа-яё]|\s*$)")


def _cut_words(slide: SlideContent) -> list[str]:
    """Words shortened with a dot to fit a limit ("Упр.", "разраб. модели") instead of rewritten."""
    found = []
    for number, text in enumerate([slide.title, *slide.bullets]):
        for match in _ABBREVIATED.finditer(text.strip()):
            word = match.group(1)
            at_end = match.end() >= len(text.strip())
            # A bullet may end a sentence with a short word; a title never ends with a dot.
            if word.casefold() in _KNOWN_ABBREVIATIONS or (at_end and number and len(text.split()) > 1):
                continue
            found.append(match.group(0).strip())
    return found


VARIANT_GOALS = {
    "base": "Тема «Базовая»: обычный слайд по шаблону. Пиши ясно и по делу, заполняй места макета без лишних подробностей.",
    "more_text": ("Тема «Больше текста»: слайд для чтения. Заполни все места макета и используй их пределы почти "
                  "полностью (80–100 %): пиши полными предложениями, раскрывай каждый факт подробнее — причины, "
                  "следствия, пояснения из тех же источников, — но без новых сведений и чисел."),
    "more_visual": ("Тема «Больше визуала»: слайд станет схемой, графиком или плитками с крупными цифрами. Пункты — "
                    "короткие ёмкие подписи до 8 слов, одна мысль в пункте. Если в выдержках есть числа (проценты, "
                    "суммы, количества, значения по годам или регионам), вынеси каждое число в свой пункт вместе с "
                    "тем, к чему оно относится, в одной единице измерения: из таких пунктов код построит график. "
                    "Если место одно, запиши ряд в одном пункте: «2023 — 12 %, 2024 — 18 %, 2025 — 25 %». "
                    "Числа бери только из выдержек, никогда не придумывай и не округляй их."),
}


def _refine_slide(client: ModelClient, slide: SlideContent, facts: dict[str, Fact], outline: list[str],
                  deadline: float | None = None, brief: str = "", slots: dict | None = None,
                  goal: str | None = None) -> SlideContent:
    selected = [{"fact_id": item, "excerpt": facts[item].excerpt} for item in slide.fact_ids if item in facts]
    sparse_source = len(facts) < max(4, len(outline) // 2) and all(
        facts[item].source == "brief" for item in slide.fact_ids
    )
    payload = {"slide": slide.to_dict(), "facts": selected, "frozen_outline": outline,
               "original_brief": brief[:12000], "sparse_source": sparse_source}
    title_limit, bullet_limit, items = 110, 450, 6
    if goal in VARIANT_GOALS:
        payload["variant_goal"] = VARIANT_GOALS[goal]
    if slots is not None:
        # Strict template mode: each text goes into one place of the template slide and may not
        # take more lines than the template's own text there.
        payload["slots"] = slots
        title_limit = max(1, min(110, slots["title"]))
        if slots["texts"]:
            bullet_limit, items = max(1, max(slots["texts"])), min(6, len(slots["texts"]))
        sparse_source = sparse_source and items >= 2
    minimum = min(items, 2 if sparse_source else 1)
    if slots is not None and slots.get("min_texts"):
        # A list drawn with checkmarks, icons or numbers gets a point for each of them.
        minimum = max(minimum, min(items, slots["min_texts"]))
    if goal == "more_text" and slots is not None:
        # The text theme fills the layout: every list line and point gets its own text.
        minimum = max(minimum, min(items, 3))
    for attempt in range(2):
        response = client.complete_json(
            prompt_text("slide_worker"),
            json.dumps(payload, ensure_ascii=False),
            max_tokens=1600,
            deadline=deadline,
            schema={"type": "object", "additionalProperties": False,
                    "required": ["title", "bullets", "fact_ids"], "properties": {
                        "title": {"type": "string", "minLength": 1, "maxLength": title_limit},
                        "bullets": {"type": "array", "minItems": minimum,
                                    "maxItems": items, "items": {"type": "string", "minLength": 1, "maxLength": bullet_limit}},
                        "fact_ids": {"type": "array", "enum": [slide.fact_ids]}}},
        )
        try:
            refined = validate_slide(response, slide.slide_id, facts, required_refs=set(slide.fact_ids))
            cut = _cut_words(refined)
            problems = _budget_violations(refined, slots) if slots is not None else []
            if (cut or problems) and not attempt:
                raise ModelProviderError(" ".join(
                    ([f"Текст не подходит к местам шаблона: {'; '.join(problems)}. Сократи, сохранив смысл: "
                      "больше писать нельзя."] if problems else [])
                    + ([f"Слова сокращены точкой: {', '.join(cut[:5])}. Не сокращай слова: напиши их полностью "
                        "или выбери другое, более короткое слово."] if cut else [])))
            if slots is not None:
                # The second answer is cut to the template at a sentence or word boundary.
                refined = fit_slide(refined, slots)
            source_lines = {" ".join(item["excerpt"].casefold().split()).strip(" .,!?:;") for item in selected}
            bullet_lines = [" ".join(item.casefold().split()).strip(" .,!?:;") for item in refined.bullets]
            if sparse_source and (
                len(bullet_lines) < 2 or len(set(bullet_lines)) < 2
                or all(len(item.split()) < 4 or item in source_lines for item in bullet_lines)
            ):
                raise ModelProviderError("Slide Worker must provide at least two distinct, grounded explanatory points for a sparse brief")
            return refined
        except ModelProviderError as error:
            if attempt:
                raise
            payload["validation_feedback"] = str(error)
    raise ModelProviderError("Slide Worker could not produce grounded text")


def _run_workers(client: ModelClient, slides: list[SlideContent], facts: list[Fact],
                 deadline: float | None = None, brief: str = "",
                 budgets: list[dict] | None = None, goal: str | None = None) -> list[SlideContent]:
    lookup = {fact.fact_id: fact for fact in facts}
    refined: list[SlideContent | None] = [None] * len(slides)
    workers = min(max(1, int(os.getenv("MODEL_MAX_PARALLEL", "4"))), 8)
    with ThreadPoolExecutor(max_workers=workers) as executor:
        outline = [slide.title for slide in slides]
        futures = {executor.submit(_refine_slide, client, slide, lookup, outline, deadline, brief,
                                   **({"slots": budgets[index]} if budgets else {}),
                                   **({"goal": goal} if goal else {})): index
                   for index, slide in enumerate(slides)}
        try:
            for future in as_completed(futures):
                refined[futures[future]] = future.result()
        except BaseException:
            for future in futures:
                future.cancel()
            raise
    return [item for item in refined if item is not None]


def _routine_template_note(message: str) -> bool:
    return message.startswith("Template Analyst: слайд ") and (
        message.endswith("восстановлены безопасно.")
        or message.endswith("исправлены для удаления старого содержания.")
    )


def _needs_semantic_repair(issue: dict[str, Any]) -> bool:
    return issue.get("severity") == "blocking" or issue.get("rule_id") in {"repetition", "duplicate_content"}


def _repair_content(repairer: ModelClient, critic: ModelClient, slides: list[SlideContent],
                    facts: list[Fact], issues: list[dict[str, Any]], deadline: float,
                    excluded_indices: set[int] | None = None,
                    prompt_critic: ModelClient | None = None, brief: str = ""):
    """One bounded repair round; grounded facts and an independent recheck."""
    excluded_indices = excluded_indices or set()
    targets = {item["slide_id"] for item in issues if _needs_semantic_repair(item)}
    indices = [i for i, slide in enumerate(slides) if slide.slide_id in targets and i not in excluded_indices]
    if not indices:
        return slides, issues, {"rounds": 0, "status": "manual_required", "initial_issues": issues}
    # Leave time to render and audit three decks; never hide unresolved blockers.
    if deadline - time.monotonic() < 150:
        return slides, issues, {"rounds": 0, "status": "insufficient_time", "initial_issues": issues}
    lookup = {fact.fact_id: fact for fact in facts}
    repaired = list(slides)
    def fix(index):
        slide = slides[index]
        sparse_source = len(facts) < max(4, len(slides) // 2) and all(
            lookup[ref].source == "brief" for ref in slide.fact_ids
        )
        slide_issues = [item for item in issues if item["slide_id"] == slide.slide_id]
        change_refs = any(item.get("rule_id") in {"repetition", "duplicate_content"} for item in slide_issues)
        candidate_facts = facts if change_refs else [lookup[ref] for ref in slide.fact_ids]
        payload = {"slide": slide.to_dict(),
                   "facts": [{"fact_id": fact.fact_id, "excerpt": fact.excerpt} for fact in candidate_facts],
                   "issues": slide_issues,
                   "frozen_outline": [item.title for item in slides],
                   "other_slides": [item.to_dict() for item in slides if item.slide_id != slide.slide_id]
                                   if change_refs else [],
                   "allow_ref_change": change_refs,
                   "sparse_source": sparse_source}
        for attempt in range(2):
            response = repairer.complete_json(
                prompt_text("slide_repair"), json.dumps(payload, ensure_ascii=False),
                max_tokens=1000, deadline=min(deadline - 95, time.monotonic() + 45))
            try:
                repaired = validate_slide(
                    response, slide.slide_id, lookup,
                    required_refs=None if change_refs else set(slide.fact_ids),
                )
                if sparse_source and len({" ".join(item.casefold().split()) for item in repaired.bullets}) < 2:
                    raise ModelProviderError("Slide Repair must keep two distinct grounded points for a sparse brief")
                return repaired
            except ModelProviderError as error:
                if attempt:
                    raise
                payload["validation_feedback"] = str(error)
        raise ModelProviderError("Slide Repair could not produce grounded text")
    with ThreadPoolExecutor(max_workers=repairer.max_parallel) as executor:
        futures = {executor.submit(fix, index): index for index in indices}
        try:
            for future in as_completed(futures):
                repaired[futures[future]] = future.result()
        except BaseException:
            for future in futures:
                future.cancel()
            raise
    final_issues = duplicate_content_issues(repaired) + critique_content(critic, repaired, facts, deadline)
    if prompt_critic is not None:
        final_issues += critique_prompt(prompt_critic, brief, repaired, facts, deadline)
    return repaired, final_issues, {"rounds": 1, "status": "blocked" if any(
        item["severity"] == "blocking" for item in final_issues) else "rechecked",
        "repaired_slide_ids": [slides[i].slide_id for i in indices], "initial_issues": issues}


DARK_PALETTE_DEFAULT = {
    "background": "#111827", "surface": "#1F2937", "text": "#F8FAFC", "accent": "#60A5FA"
}
DARK_PALETTE_CHOICES = {
    "background": ["#0B1220", "#101418", "#111827"],
    "surface": ["#172033", "#1F2937", "#252B36"],
    "text": ["#F3F4F6", "#F8FAFC"],
    "accent": ["#34D399", "#60A5FA", "#A78BFA", "#F59E0B"],
}


def _design_dark_theme(client: ModelClient, brief: str, deadline: float) -> dict[str, str]:
    """Qwen chooses a constrained dark palette; renderer applies it to variant B."""
    schema = {"type": "object", "additionalProperties": False,
              "required": list(DARK_PALETTE_CHOICES),
              "properties": {name: {"type": "string", "enum": options}
                             for name, options in DARK_PALETTE_CHOICES.items()}}
    answer = client.complete_json(
        prompt_text("dark_theme_designer"),
        json.dumps({"variant": "B", "instruction": "Тёмная тема; сохранить текст, порядок слайдов и брендовые изображения",
                    "brief": brief[:4000], "allowed_palette": DARK_PALETTE_CHOICES}, ensure_ascii=False),
        max_tokens=180, deadline=deadline, schema=schema)
    if set(answer) != set(DARK_PALETTE_CHOICES) or any(
        answer[name] not in options for name, options in DARK_PALETTE_CHOICES.items()
    ):
        raise ModelProviderError("Dark Theme Designer returned an invalid palette")
    return answer



def _plan_infographics(client: ModelClient, template, slides: list[SlideContent],
                       deadline: float, compositions: list | None = None, visual: bool = False) -> dict[str, str]:
    """Pick safe, content-bearing slides; the renderer uses only approved bullets.

    visual is the "more visuals" theme: diagrams on most content slides, and a chart or big
    figures wherever the points carry figures (found by code in the approved text, never invented).
    """
    if compositions is None:
        sampled = choose_compositions(template, len(slides))
        by_source = {c.source_slide_index: c for c in template.compositions}
        compositions = [by_source[s.source_slide_index] if s.source_slide_index is not None else sampled[i]
                        for i, s in enumerate(slides)]
    options = {slide.slide_id: visual_options(slide.bullets) if visual else ["modules", "flow"] for slide in slides}
    candidates = []
    for index in range(1, len(slides) - 1):
        slide, composition = slides[index], compositions[index]
        figures_ready = visual and options[slide.slide_id][0] in {"chart", "stats"}
        if (slide.dense_layout or not (1 if figures_ready else 2) <= len(slide.bullets) <= (8 if figures_ready else 6)
                or composition.picture_area_ratio >= 0.25):
            continue
        # A layout already built of icons, numbers or checkmarks is a visual of its own; only
        # figures that make a chart replace it.
        if (_point_slots(composition, template.width, template.height)
                or any(slot.list_items for slot in composition.slots)) and not figures_ready:
            continue
        has_heading = any(
            slot.y < template.height * 0.28
            and slot.width > template.width * 0.30
            and slot.height < template.height * 0.28
            for slot in composition.slots
        )
        if has_heading or visual:
            candidates.append(slide)
    if not candidates:
        candidates = [
            slides[index] for index in range(1, len(slides) - 1)
            if not slides[index].dense_layout and 2 <= len(slides[index].bullets) <= 6
            and compositions[index].picture_area_ratio < 0.25
        ]
    if not candidates:
        return {}
    minimum = 2 if len(slides) >= 7 and len(candidates) >= 2 else 1
    maximum = min(3, len(candidates))
    if visual:
        # The visual theme draws most of its content slides, not two or three of them.
        maximum = min(len(candidates), max(3, math.ceil((len(slides) - 2) * 0.6)))
        minimum = min(maximum, max(minimum, math.ceil(maximum / 2)))
    layouts = ["modules", "flow", "stats", "chart"] if visual else ["modules", "flow"]
    candidate_ids = [slide.slide_id for slide in candidates]
    schema = {
        "type": "object", "additionalProperties": False, "required": ["slides"],
        "properties": {"slides": {
            "type": "array", "minItems": minimum, "maxItems": maximum,
            "items": {"type": "object", "additionalProperties": False,
                      "required": ["slide_id", "layout"],
                      "properties": {
                          "slide_id": {"type": "string", "enum": candidate_ids},
                          "layout": {"type": "string", "enum": layouts},
                      }},
        }},
    }
    rows_in = [{**slide.to_dict(), "allowed_layouts": options[slide.slide_id]} if visual else slide.to_dict()
               for slide in candidates]
    answer = client.complete_json(
        prompt_text("infographic_designer"),
        json.dumps({"candidates": rows_in, "minimum": minimum, "maximum": maximum,
                    **({"theme": "more_visual"} if visual else {})}, ensure_ascii=False),
        max_tokens=450 + 40 * len(candidates), deadline=deadline, schema=schema,
    )
    rows = answer.get("slides") if isinstance(answer, dict) else None
    if not isinstance(rows, list) or not minimum <= len(rows) <= maximum:
        raise ModelProviderError("Infographic Designer returned the wrong number of slides")
    plan = {}
    for row in rows:
        if (not isinstance(row, dict) or set(row) != {"slide_id", "layout"}
                or row["slide_id"] not in candidate_ids or row["layout"] not in layouts
                or row["slide_id"] in plan):
            raise ModelProviderError("Infographic Designer returned an invalid diagram plan")
        allowed = options[row["slide_id"]]
        # A chart or big figures need figures in the points; otherwise the slide gets tiles.
        plan[row["slide_id"]] = row["layout"] if row["layout"] in allowed else allowed[0]
    if visual:
        # Figures that form a series are always drawn: a chart where the designer left the slide
        # out, and a chart instead of plain tiles where it chose them.
        for slide in candidates:
            if options[slide.slide_id][0] != "chart":
                continue
            if slide.slide_id not in plan and len(plan) < maximum or plan.get(slide.slide_id) in {"modules", "flow"}:
                plan[slide.slide_id] = "chart"
    return plan

def _variant_id(index: int) -> str:
    return ("a", "b", "c")[index]


def generate(
    template_path: Path,
    brief: str,
    slide_count: int,
    output_dir: Path,
    progress: Progress | None = None,
    content_paths: list[Path] | None = None,
    prepared_path: Path | None = None,
) -> list[GeneratedVariant]:
    """Generate three variants; generation clock starts after TemplateIR is ready."""
    template_path = Path(template_path)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    if not 1 <= slide_count <= 60:
        raise ValueError("slide_count must be between 1 and 60")
    _progress(progress, "Проверка подготовки шаблона", 1)
    if prepared_path is None:
        prepared_path = prepare(template_path, output_dir / "preparation")
    template = load_prepared(Path(prepared_path), template_path)
    preparation_started = time.monotonic()
    facts = extract_facts(brief, content_paths)
    topic_only = False
    if not facts and brief.strip():
        # A short topic is a valid prompt, though it does not support numbers.
        excerpt = brief.strip()
        facts = [Fact("fact_prompt_" + hashlib.sha256(excerpt.encode("utf-8")).hexdigest()[:12],
                      excerpt, "brief", "prompt", excerpt)]
        topic_only = True
    if not facts:
        raise InsufficientContent("Промпт и материалы не содержат текста для генерации.")
    # Template fonts are measured and embedded; open ones the store lacks are fetched once.
    font_report = ensure_fonts(template_families(Presentation(str(template_path))))
    typography = _template_typography(Presentation(str(template_path)))
    # Strict template mode (default): the text fills the template's own regions, exactly as they are.
    strict = os.getenv("AYA_STRICT_TEMPLATE", "1") != "0"
    if strict:
        enrich_slots(template, template_path)
        # Cached per template; a template prepared before the Field Marker is measured once here.
        _mark_template_fields(template_path, template, output_dir / "preparation")
    content_preparation_seconds = time.monotonic() - preparation_started
    generation_started = time.monotonic()
    budget = generation_budget_seconds(slide_count)
    deadline = generation_started + budget  # Reserve 10 seconds for publication.
    _progress(progress, "Планирование презентации", 10)
    gateway = ModelGateway()
    client = gateway.client("deck_planner")
    worker = gateway.client("slide_worker")
    content_critic = gateway.client("deck_critic")
    prompt_critic = gateway.client("prompt_critic")
    infographic_designer = gateway.client("infographic_designer")
    critic = gateway.client("visual_critic") if gateway.enabled and gateway.enable_visual_critic else None
    if gateway.enabled and template.analysis_mode != "vision_model":
        raise ModelProviderError("Template was prepared without vision analysis; prepare it again with OpenRouter")
    if not client.enabled:
        raise ModelProviderError("Модель не настроена. Укажите OPENROUTER_API_KEY.")
    slides = _model_plan(client, facts, slide_count, brief, deadline)
    layout_plans: list[list] = []
    budgets: list[list[dict]] | None = None
    variant_texts: list[list[SlideContent]] | None = None
    if strict:
        # Three themes, each with its own layouts and text: "base" takes the closest template
        # layouts, "more_text" those holding the most text, "more_visual" those with icons,
        # numbers and parallel points. A slide's text is written for its own layout's places.
        for goal in GOALS:
            layout_plans.append(choose_compositions(template, len(slides), slides, avoid=layout_plans,
                                                    strict=True, goal=goal))
        budgets = [[slot_budget([plan[index]], template.width, template.height) for index in range(len(slides))]
                   for plan in layout_plans]
        with ThreadPoolExecutor(max_workers=3) as executor:
            variant_texts = list(executor.map(
                lambda number: _run_workers(worker, slides, facts, deadline, brief, budgets[number], goal=GOALS[number]),
                range(3)))
        slides = variant_texts[0]
    else:
        slides = _run_workers(worker, slides, facts, deadline, brief)
    model_mode = "configured_api"
    native_data = None
    native_source_name = None
    native_warning = None
    for content_path in content_paths or []:
        candidate = Path(content_path)
        if candidate.suffix.lower() != ".csv":
            continue
        try:
            native_data = parse_numeric_csv(candidate)
            native_source_name = candidate.name
            break
        except ValueError as exc:
            native_warning = f"{candidate.name}: {exc}"
    native_slide_index = None
    if native_data is not None:
        native_slide_index = slide_count // 2
        data_refs = [fact.fact_id for fact in facts if fact.source == native_source_name]
        slides[native_slide_index] = SlideContent(
            slides[native_slide_index].slide_id,
            "\u0414\u0430\u043d\u043d\u044b\u0435 \u0438\u0437 \u0442\u0430\u0431\u043b\u0438\u0446\u044b",
            [],
            data_refs or slides[native_slide_index].fact_ids,
        )
        for texts in variant_texts or []:
            texts[native_slide_index] = slides[native_slide_index]
    def review(variant_slides: list[SlideContent]):
        issues = (duplicate_content_issues(variant_slides) + critique_content(content_critic, variant_slides, facts, deadline)
                  + critique_prompt(prompt_critic, brief, variant_slides, facts, deadline))
        outcome = {"rounds": 0, "status": "not_needed"}
        if gateway.repair_rounds and any(_needs_semantic_repair(item) for item in issues):
            _progress(progress, "Исправление содержания по замечаниям критика", 16)
            variant_slides, issues, outcome = _repair_content(
                gateway.client("slide_repair"), content_critic, variant_slides, facts, issues, deadline,
                excluded_indices={native_slide_index} if native_slide_index is not None else set(),
                prompt_critic=prompt_critic, brief=brief)
        return variant_slides, issues, outcome

    variant_issues: list[list[dict[str, Any]]] | None = None
    if strict:
        # Every theme's text is checked by the critics; the three reviews run together.
        with ThreadPoolExecutor(max_workers=3) as executor:
            reviewed = list(executor.map(review, variant_texts))
        # A repaired slide is held to its places too; no page is added.
        variant_texts = [[slide if index == native_slide_index else fit_slide(slide, budgets[number][index])
                          for index, slide in enumerate(texts)] for number, (texts, _, _) in enumerate(reviewed)]
        variant_issues = [issues for _, issues, _ in reviewed]
        slides, content_issues, repair_report = variant_texts[0], variant_issues[0], reviewed[0][2]
        # The visual theme gets the Infographic Designer's diagrams on slides whose layout has one
        # plain text place (a list); the other themes keep the template's own objects only.
        try:
            infographic_plan = _plan_infographics(infographic_designer, template, variant_texts[2], deadline,
                                                  compositions=layout_plans[2], visual=True)
        except ModelProviderError:
            infographic_plan = {}
    else:
        content_issues = (duplicate_content_issues(slides) + critique_content(content_critic, slides, facts, deadline)
                          + critique_prompt(prompt_critic, brief, slides, facts, deadline))
        repair_report = {"rounds": 0, "status": "not_needed"}
        if gateway.repair_rounds and any(_needs_semantic_repair(item) for item in content_issues):
            _progress(progress, "Исправление содержания по замечаниям критика", 16)
            slides, content_issues, repair_report = _repair_content(
                gateway.client("slide_repair"), content_critic, slides, facts, content_issues, deadline,
                excluded_indices={native_slide_index} if native_slide_index is not None else set(),
                prompt_critic=prompt_critic, brief=brief)
        native_original_id = slides[native_slide_index].slide_id if native_slide_index is not None else None
        # Long copy continues on further pages in the template's own typeface and sizes.
        slides = paginate_slides(slides, choose_compositions(template, len(slides), slides),
                                 template.width, template.height, typography["font"],
                                 typography["body_pt"], typography["title_pt"], typography["title_font"])
        # Audit identifiers and native data positions refer to physical pages after pagination.
        content_issues = [{**issue, "slide_id": page.slide_id,
                          "issue_id": issue["issue_id"] + "_" + page.slide_id}
                         for issue in content_issues for page in slides
                         if page.continuation_of == issue["slide_id"]]
        native_slide_index = next((i for i, page in enumerate(slides)
                                   if page.continuation_of == native_original_id), None)
        infographic_plan = _plan_infographics(infographic_designer, template, slides, deadline)
    planning_seconds = round(time.monotonic() - generation_started, 3)
    variants: list[GeneratedVariant] = []
    source_texts = source_texts_from_template(template_path)
    report: dict[str, Any] = {
        "schema_version": "2.0",
        "pagination": {"requested_slide_count": slide_count, "actual_slide_count": len(slides),
                       "added_slides": len(slides) - slide_count, "body_pt": typography["body_pt"]},
        "template_sha256": template.template_sha256,
        "model_mode": model_mode,
        "model_id": client.text_model,
        "model_calls": gateway.calls,
        "provider_configuration": gateway.configuration_summary(),
        "semantic_repair": repair_report,
        "variant_design": ({"a": "base", "b": "more_text", "c": "more_visual"} if strict else
                           {"a": "source_template", "b": {"theme": "source", "layout": "stacked"},
                            "c": "alternate_layout_only"}),
        "infographic_plan": infographic_plan,
        "template_mode": "strict" if strict else "adaptive",
        "layouts": {variant: [item.source_slide_index + 1 for item in plan]
                    for variant, plan in zip("abc", layout_plans)} if strict else None,
        "slot_budgets": budgets,
        "facts": [fact.to_dict() for fact in facts],
        "slides": [slide.to_dict() for slide in slides],
        "variants": [],
        "native_csv": native_source_name,
        "fonts": font_report,
        # Deterministic inspectors that run after every build, next to the model agents.
        "inspectors": {"field_marker": "free area of every text field, measured on the template rendered without text",
                       "field_checker": "letters inside their field's free area, not overlapped, not covered by a picture",
                       "font_inspector": "template typeface and size of every generated text region",
                       "design_inspector": "palette, contrast, guides, margins, pictures, brand, fill, chart axes"},
        "warnings": [message for message in template.warnings if not _routine_template_note(message)]
                   + ([native_warning] if native_warning else [])
                   + ([f"Для сохранения текста и читаемого шрифта добавлено слайдов-продолжений: {len(slides) - slide_count}."] if len(slides) > slide_count else [])
                   + (["Короткий промпт задаёт тему, но не содержит проверяемых фактов; числовые утверждения запрещены."] if topic_only else []),
        "agents": {role: agent.text_model for role, agent in gateway.clients.items()},
        "template_analysis": {"mode": template.analysis_mode, "model": template.analysis_model},
        "template_diagnostics": [message for message in template.warnings if _routine_template_note(message)],
        "timings": {"budget_seconds": budget + 10, "planning_seconds": planning_seconds, "content_preparation_seconds": round(content_preparation_seconds, 3)},
    }
    source_previews = []
    if critic:
        _progress(progress, "Подготовка визуального сравнения", 18)
        source_previews = _visual_references(template_path, template, output_dir, deadline)
        report["agents"]["visual_critic"] = critic.vision_model
    _progress(progress, "Сборка и экспорт трёх вариантов", 20)

    def build_variant(index: int, variant_slides: list[SlideContent] | None = None, revision: int = 0,
                      semantic_issues: list[dict[str, Any]] | None = None):
        remaining_seconds(deadline)
        variant_slides = variant_slides or slides
        variant_id = _variant_id(index)
        root = output_dir if revision == 0 else output_dir / f"visual_revision_{revision}"
        folder = root / f"variant_{variant_id}"
        folder.mkdir(parents=True, exist_ok=True)
        started = time.monotonic()
        pptx_path = folder / f"aiya_{variant_id}.pptx"
        relayout: set[int] = set()
        # In strict mode only the visual theme draws diagrams.
        variant_plan = infographic_plan if not strict or index == 2 else {}
        for attempt in range(2):
            render_metrics = render_variant(
                template_path, template, variant_slides, index, pptx_path,
                native_data=native_data, native_slide_index=native_slide_index,
                infographic_layouts=variant_plan, relayout=relayout, strict=strict,
                plan=[item.source_slide_index for item in layout_plans[index]] if strict else None,
            )
            expected_infographics = {
                slide_index + 1 for slide_index, planned in enumerate(variant_slides)
                if planned.slide_id in variant_plan
            }
            if (set(render_metrics["infographic_slides"])
                    | set(render_metrics.get("infographic_fallback_slides", []))) != expected_infographics:
                raise ValueError(f"Инфографика по плану не построена в варианте {variant_id.upper()}")
            expected_fonts = render_metrics["font_expectations"]
            issues = (audit_pptx(pptx_path, variant_slides, facts, source_texts)
                      + font_issues(pptx_path, expected_fonts)
                      + design_issues(pptx_path, template_path, variant_slides, template,
                                      render_metrics["source_slides"], render_metrics.get("dark_palette"))
                      + (content_issues if semantic_issues is None else semantic_issues))
            # Small text is fine where it is the template's own size.
            templated = {(f"slide_{item['slide']}", str(item["shape_id"])) for item in expected_fonts}
            issues = [issue for issue in issues if not (issue["rule_id"] == "small_body_text"
                                                       and (issue["slide_id"], issue["object_id"]) in templated)]
            issues = list({issue["issue_id"]: issue for issue in issues}.values())
            pdf_path = export_pdf(pptx_path, folder, timeout=remaining_seconds(deadline))
            # The letters in the PDF decide about overlap: frames may cross in their empty parts
            # (an oversized frame over an image, two frames over each other) without any collision.
            issues = [issue for issue in issues if issue["rule_id"] not in {"text_image_overlap", "text_text_overlap"}]
            # A slide drawn as a diagram has new objects instead of the template's fields.
            fields = ({number: {slot.shape_id: slot.clear_box for slot in composition.slots if slot.clear_box}
                       for number, composition in enumerate(layout_plans[index])
                       if variant_slides[number].slide_id not in variant_plan} if strict else None)
            issues += audit_pdf(pptx_path, pdf_path, variant_slides, fields)
            # A slide whose letters of different objects overlap is laid out again, once, in a
            # clean region under its heading; text leaving its frame and what remains after the
            # new layout go to the repair loop, which shortens the text.
            crowded = {int(issue["slide_id"].split("_")[1]) - 1 for issue in issues
                       if issue["rule_id"] == "rendered_text_overlap"} - relayout
            # Strict mode never moves a template region: its findings go to the repair loop.
            if strict or attempt or not crowded or deadline - time.monotonic() < 90:
                break
            relayout |= crowded
        previews = export_previews(pdf_path, folder / "previews", timeout=remaining_seconds(deadline))
        html_path = export_html(pptx_path, previews, folder / f"aiya_{variant_id}.html")
        return pptx_path, pdf_path, html_path, previews, render_metrics, issues, started

    # Each variant has its own Presentation and LibreOffice profile. Build them
    # together so a slow PDF export cannot consume the entire final slot.
    built = [None] * 3
    with ThreadPoolExecutor(max_workers=3) as executor:
        futures = {executor.submit(build_variant, index, variant_texts[index] if strict else None, 0,
                                   variant_issues[index] if strict else None): index for index in range(3)}
        try:
            for future in as_completed(futures):
                built[futures[future]] = future.result()
        except BaseException:
            for future in futures:
                future.cancel()
            raise

    def visual_feedback(index: int, result, variant_slides: list[SlideContent],
                        changed: list[int] | None = None) -> list[dict[str, Any]]:
        if critic is None:
            return []
        previews, render_metrics = result[3], result[4]
        own = strict or index == 0
        baseline = source_previews if own else built[0][3]
        source_indices = (render_metrics["source_slides"] if own
                          else list(range(1, len(variant_slides) + 1)))
        theme = ("source", "reflow", "reflow")[index]
        if changed is None:
            return critique_visual_variant(critic, baseline, previews, source_indices,
                                           variant_slides, deadline, variant_theme=theme)
        with ThreadPoolExecutor(max_workers=critic.max_parallel) as executor:
            futures = [executor.submit(
                critique_visual_slide, critic, baseline[source_indices[i] - 1],
                previews[i], variant_slides[i], i, deadline, theme,
            ) for i in changed]
            return [issue for future in futures for issue in future.result()]

    def repair_visual_slide(index: int, slide_index: int, current, variant_slides,
                            observations: list[dict[str, Any]]) -> SlideContent:
        original = variant_slides[slide_index]
        lookup = {fact.fact_id: fact for fact in facts}
        sparse = len(facts) < max(4, len(variant_slides) // 2) and all(
            lookup[ref].source == "brief" for ref in original.fact_ids
        )
        own = strict or index == 0
        baseline = source_previews if own else built[0][3]
        source_indices = current[4]["source_slides"] if own else list(range(1, len(variant_slides) + 1))
        image = contact_sheet([baseline[source_indices[slide_index] - 1], current[3][slide_index]])
        payload = {
            "slide": original.to_dict(),
            "facts": [{"fact_id": ref, "excerpt": lookup[ref].excerpt} for ref in original.fact_ids],
            "visual_issues": observations,
            "variant_theme": ("source", "reflow", "reflow")[index],
            "frozen_outline": [item.title for item in variant_slides],
            "sparse_source": sparse,
        }
        if strict:
            payload["slots"] = budgets[index][slide_index]
        repairer = gateway.client("slide_repair")
        for attempt in range(2):
            response = repairer.complete_json(
                prompt_text("slide_repair"), json.dumps(payload, ensure_ascii=False),
                vision=True, image_data_url=image, max_tokens=1000,
                deadline=min(deadline - 35, time.monotonic() + 30),
            )
            try:
                revised = validate_slide(response, original.slide_id, lookup,
                                         required_refs=set(original.fact_ids))
                if sparse and len({" ".join(item.casefold().split()) for item in revised.bullets}) < 2:
                    raise ModelProviderError("Visual repair must keep two distinct grounded points")
                revised.source_slide_index = original.source_slide_index
                revised.dense_layout = original.dense_layout
                revised.continuation_of = original.continuation_of
                return fit_slide(revised, budgets[index][slide_index]) if strict else revised
            except ModelProviderError as error:
                if attempt:
                    raise
                payload["validation_feedback"] = str(error)
        raise ModelProviderError("Visual repair did not produce a valid slide")

    def audit_and_repair(index: int):
        current = built[index]
        variant_slides = list(variant_texts[index] if strict else slides)
        try:
            observations = visual_feedback(index, current, variant_slides)
        except (ModelProviderError, TimeoutError) as error:
            from .audit import _issue
            observations = [_issue(0, None, "visual_audit_incomplete", "blocking", str(error), "manual")]
            return current, variant_slides, observations, 0, [], [{"status": "audit_failed", "error": str(error)}]
        observations += [issue for issue in current[5] if issue["rule_id"].startswith(("rendered_", "field_"))]
        attempts = 0
        repaired_ids: list[str] = []
        history = []
        while observations and attempts < 10 and deadline - time.monotonic() > 65:
            affected = sorted({int(item["slide_id"].split("_")[1]) - 1 for item in observations})
            previous = {i: [item for item in observations if item["slide_id"] == variant_slides[i].slide_id]
                        for i in affected}
            revised = list(variant_slides)
            attempts += 1
            try:
                with ThreadPoolExecutor(max_workers=min(len(affected), gateway.client("slide_repair").max_parallel)) as executor:
                    futures = {executor.submit(repair_visual_slide, index, i, current, variant_slides,
                                               previous[i]): i for i in affected}
                    for future in as_completed(futures):
                        revised[futures[future]] = future.result()
                if revised == variant_slides:
                    history.append({"attempt": attempts, "status": "unchanged", "issues": observations})
                    break
                semantic_issues = (duplicate_content_issues(revised) + critique_content(content_critic, revised, facts, deadline)
                                   + critique_prompt(prompt_critic, brief, revised, facts, deadline))
                rebuilt = build_variant(index, revised, revision=attempts, semantic_issues=semantic_issues)
                refreshed = visual_feedback(index, rebuilt, revised, affected)
                refreshed += [issue for issue in rebuilt[5] if issue["rule_id"].startswith(("rendered_", "field_"))]
            except (ModelProviderError, TimeoutError, RuntimeError, ValueError) as error:
                history.append({"attempt": attempts, "slide_ids": [variant_slides[i].slide_id for i in affected],
                                "issues": observations, "error": str(error)})
                break
            repaired_ids.extend(revised[i].slide_id for i in affected)
            history.append({"attempt": attempts, "slide_ids": [revised[i].slide_id for i in affected],
                            "issues": observations})
            variant_slides = revised
            current = rebuilt
            observations = [item for item in observations if item["slide_id"] not in
                            {variant_slides[i].slide_id for i in affected}]
            observations += refreshed
        return current, variant_slides, observations, attempts, sorted(set(repaired_ids)), history

    if critic:
        _progress(progress, "Визуальная проверка и исправление вариантов", 75)
        with ThreadPoolExecutor(max_workers=3) as executor:
            futures = {executor.submit(audit_and_repair, index): index for index in range(3)}
            audited = [None] * 3
            try:
                for future in as_completed(futures):
                    audited[futures[future]] = future.result()
            except BaseException:
                for future in futures:
                    future.cancel()
                raise
    else:
        audited = [(built[index], variant_texts[index] if strict else slides, [], 0, [], []) for index in range(3)]

    for index, (result, variant_slides, visual_issues, attempts, repaired_ids, history) in enumerate(audited):
        pptx_path, pdf_path, html_path, previews, render_metrics, issues, started = result
        issues = list({issue["issue_id"]: issue for issue in [*issues, *visual_issues]}.values())
        quality = "failed_quality_gate" if any(issue["severity"] == "blocking" for issue in issues) else (
            "completed_with_warnings" if issues else "completed"
        )
        metrics = {**render_metrics, "seconds": round(time.monotonic() - started, 3),
                   "quality_status": quality, "model_mode": model_mode,
                   "visual_repair_attempts": attempts, "visual_repaired_slide_ids": repaired_ids}
        variant_id = _variant_id(index)
        variants.append(GeneratedVariant(variant_id, pptx_path, pdf_path, html_path, previews, issues, metrics))
        report["variants"].append({"variant_id": variant_id, "metrics": metrics, "issues": issues,
                                   "visual_repair_history": history,
                                   "slides": [slide.to_dict() for slide in variant_slides]})
    generation_seconds = round(time.monotonic() - generation_started, 3)
    report["model_calls"] = gateway.calls
    report["timings"]["three_variants_seconds"] = generation_seconds
    report["timings"]["within_300_seconds"] = generation_seconds <= 300
    (output_dir / "run_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    _progress(progress, "Готово", 100)
    return variants


def repair(
    template_path: Path,
    base_pptx_path: Path,
    brief: str,
    slide_count: int,
    selected_issues: list[dict[str, Any]],
    output_dir: Path,
    progress: Progress | None = None,
    content_paths: list[Path] | None = None,
) -> list[GeneratedVariant]:
    """Apply only selected, safe geometric/text repairs to one immutable version."""
    if not selected_issues:
        raise ValueError("No audit issues selected")
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    presentation = Presentation(str(base_pptx_path))
    # The immutable PPTX is authoritative after automatic continuation pages.
    slide_count = len(presentation.slides)
    allowed_rules = {"possible_text_overflow", "out_of_bounds", "old_template_content"}
    for issue in selected_issues:
        rule = issue.get("rule_id")
        if rule not in allowed_rules or issue.get("repairability") != "automatic":
            raise ValueError(f"Selected issue requires manual correction: {rule}")
        match = re.fullmatch(r"slide_(\d+)", str(issue.get("slide_id", "")))
        if not match:
            raise ValueError("Invalid issue slide_id")
        index = int(match.group(1)) - 1
        if not 0 <= index < slide_count:
            raise ValueError("Issue slide_id is outside the deck")
        shape_id = int(issue["object_id"])
        shape = next((item for item in presentation.slides[index].shapes if item.shape_id == shape_id), None)
        if shape is None:
            raise ValueError("Issue object_id does not exist in base version")
        if rule == "old_template_content":
            if shape.has_text_frame:
                shape.text_frame.clear()
        elif rule == "possible_text_overflow":
            if not shape.has_text_frame:
                raise ValueError("Text overflow issue points to a non-text object")
            # The template's type size is kept: the frame grows into free space below it.
            frame = shape.text_frame
            needed = int(text_block_height(shape) * 12700) + frame.margin_top + frame.margin_bottom + int(Pt(4))
            room = presentation.slide_height - shape.top - int(Pt(14))
            if needed > room:
                raise ValueError("Текст не помещается при кегле шаблона: сократите его правкой слайда")
            left, top, width = shape.left, shape.top, shape.width
            shape.left, shape.top, shape.width, shape.height = left, top, width, max(shape.height, needed)
        elif rule == "out_of_bounds":
            shape.left = max(0, min(shape.left, presentation.slide_width - shape.width))
            shape.top = max(0, min(shape.top, presentation.slide_height - shape.height))
    pptx_path = output_dir / "aiya_repair.pptx"
    presentation.save(str(pptx_path))
    _progress(progress, "Проверка исправлений", 60)
    pdf_path = export_pdf(pptx_path, output_dir)
    previews = export_previews(pdf_path, output_dir / "previews")
    html_path = export_html(pptx_path, previews, output_dir / "aiya_repair.html")
    facts = extract_facts(brief, content_paths)
    slides = []
    for index, slide in enumerate(presentation.slides):
        texts = [shape.text.strip() for shape in slide.shapes if shape.has_text_frame and shape.text.strip()]
        slides.append(SlideContent(f"slide_{index+1}", texts[0] if texts else "", texts[1:], []))
    issues = audit_pptx(pptx_path, slides, facts, source_texts_from_template(Path(template_path)))
    quality = "failed_quality_gate" if any(item["severity"] == "blocking" for item in issues) else "completed_with_warnings" if issues else "completed"
    metrics = {"slide_count": slide_count, "quality_status": quality, "repaired_issue_ids": [item["issue_id"] for item in selected_issues]}
    _progress(progress, "Готово", 100)
    return [GeneratedVariant("repair", pptx_path, pdf_path, html_path, previews, issues, metrics)]


def edit_slide(
    template_path: Path,
    base_pptx_path: Path,
    brief: str,
    slide_count: int,
    slide_index: int,
    instruction: str,
    output_dir: Path,
    progress: Progress | None = None,
    content_paths: list[Path] | None = None,
) -> list[GeneratedVariant]:
    """Apply a requested text or layout change to one slide of one version."""
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    presentation = Presentation(str(base_pptx_path))
    slide_count = len(presentation.slides)
    if not 1 <= slide_index <= slide_count:
        raise ValueError("Selected slide is outside the presentation")
    slide = presentation.slides[slide_index - 1]
    width, height = int(presentation.slide_width), int(presentation.slide_height)
    shapes = [{
        "shape_id": shape.shape_id,
        "type": str(shape.shape_type),
        "text": shape.text[:1500] if shape.has_text_frame else "",
        "left": round(shape.left / width, 4),
        "top": round(shape.top / height, 4),
        "width": round(shape.width / width, 4),
        "height": round(shape.height / height, 4),
    } for shape in slide.shapes]
    if not shapes:
        raise ValueError("Selected slide has no editable objects")
    _progress(progress, "Применяем правки к слайду", 20)
    system = prompt_text("slide_edit")
    payload = {
        "brief": brief[:4000],
        "slide_number": slide_index,
        "instruction": instruction,
        "shapes": shapes,
    }
    response = ModelGateway().client("slide_repair").complete_json(
        system, json.dumps(payload, ensure_ascii=False), max_tokens=1800,
    )
    edits = response.get("edits")
    if not isinstance(edits, list) or not 1 <= len(edits) <= len(shapes):
        raise ValueError("The editing model returned no valid slide changes")
    by_id = {shape.shape_id: shape for shape in slide.shapes}
    seen: set[int] = set()
    for edit in edits:
        if not isinstance(edit, dict) or type(edit.get("shape_id")) is not int:
            raise ValueError("The editing model returned an invalid shape")
        shape_id = edit["shape_id"]
        if shape_id in seen or shape_id not in by_id:
            raise ValueError("The editing model referred to an unknown shape")
        seen.add(shape_id)
        shape = by_id[shape_id]
        if "text" in edit:
            value = edit["text"]
            if not shape.has_text_frame or not isinstance(value, str) or len(value) > 2000:
                raise ValueError("The editing model returned invalid slide text")
            paragraphs = shape.text_frame.paragraphs
            if paragraphs and paragraphs[0].runs:
                paragraphs[0].runs[0].text = value
                for run in paragraphs[0].runs[1:]:
                    run.text = ""
                for paragraph in paragraphs[1:]:
                    paragraph.text = ""
            else:
                shape.text_frame.text = value
        geometry = {key: edit[key] for key in ("left", "top", "width", "height") if key in edit}
        if geometry:
            current = {
                "left": shape.left / width,
                "top": shape.top / height,
                "width": shape.width / width,
                "height": shape.height / height,
            }
            current.update(geometry)
            if any(type(value) not in (int, float) or not 0 <= value <= 1 for value in current.values()):
                raise ValueError("The editing model returned invalid slide geometry")
            if current["width"] <= 0 or current["height"] <= 0 or current["left"] + current["width"] > 1.01 or current["top"] + current["height"] > 1.01:
                raise ValueError("The editing model placed an object outside the slide")
            shape.left = int(current["left"] * width)
            shape.top = int(current["top"] * height)
            shape.width = int(current["width"] * width)
            shape.height = int(current["height"] * height)
    pptx_path = output_dir / "lukas_edit.pptx"
    presentation.save(str(pptx_path))
    _progress(progress, "Экспорт исправленной презентации", 60)
    pdf_path = export_pdf(pptx_path, output_dir)
    previews = export_previews(pdf_path, output_dir / "previews")
    html_path = export_html(pptx_path, previews, output_dir / "lukas_edit.html")
    facts = extract_facts(brief, content_paths)
    slides = []
    for index, current_slide in enumerate(presentation.slides):
        texts = [shape.text.strip() for shape in current_slide.shapes if shape.has_text_frame and shape.text.strip()]
        slides.append(SlideContent(f"slide_{index + 1}", texts[0] if texts else "", texts[1:], []))
    issues = audit_pptx(pptx_path, slides, facts, source_texts_from_template(Path(template_path)))
    quality = "failed_quality_gate" if any(item["severity"] == "blocking" for item in issues) else "completed_with_warnings" if issues else "completed"
    _progress(progress, "Готово", 100)
    return [GeneratedVariant("edit", pptx_path, pdf_path, html_path, previews, issues, {
        "slide_count": slide_count, "quality_status": quality, "edited_slide": slide_index, "instruction": instruction,
    })]
