import io
import sys
import types
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select

from app import db as db_module
from app import jobs
from app import main as api
from app.config import get_settings
from app.models import Run, Source, Version
from app.storage import get_storage


@pytest.fixture
def environment(tmp_path, monkeypatch):
    monkeypatch.setenv("AYA_DATABASE_URL", f"sqlite:///{(tmp_path / 'test.db').as_posix()}")
    monkeypatch.setenv("AYA_STORAGE_ROOT", str(tmp_path / "storage"))
    monkeypatch.setenv("AYA_INLINE_JOBS", "false")
    monkeypatch.setenv("AYA_LOCAL_JOBS", "false")
    get_settings.cache_clear()
    db_module.get_engine.cache_clear()
    db_module.get_session_factory.cache_clear()
    monkeypatch.setattr(api, "enqueue_preparation", lambda source_id: None)
    monkeypatch.setattr(api, "enqueue_run", lambda run_id: None)
    engine = types.ModuleType("engine")
    engine.__path__ = []
    pipeline = types.ModuleType("engine.pipeline")
    engine.pipeline = pipeline
    monkeypatch.setitem(sys.modules, "engine", engine)
    monkeypatch.setitem(sys.modules, "engine.pipeline", pipeline)
    with TestClient(api.app) as client:
        registered = client.post("/api/auth/register", json={"email": "first@example.com", "password": "correct-horse-42"})
        assert registered.status_code == 201, registered.text
        yield client, pipeline, tmp_path
    db_module.get_engine().dispose()
    db_module.get_engine.cache_clear()
    db_module.get_session_factory.cache_clear()
    get_settings.cache_clear()


def test_email_auth_and_account_isolation(environment):
    client, _, _ = environment
    first_workspace = client.get("/api/workspace")
    assert first_workspace.status_code == 200
    first_project_id = first_workspace.json()["id"]
    assert client.get("/api/projects").json()[0]["id"] == first_project_id

    assert client.post("/api/auth/logout").status_code == 204
    assert client.get("/api/workspace").status_code == 401
    assert client.post("/api/auth/login", json={"email": "first@example.com", "password": "wrong-password"}).status_code == 401
    second = client.post("/api/auth/register", json={"email": "second@example.com", "password": "another-password-42"})
    assert second.status_code == 201
    assert client.get(f"/api/projects/{first_project_id}").status_code == 404
    assert client.get(f"/api/projects/{first_project_id}/sources").status_code == 404
    assert all(project["id"] != first_project_id for project in client.get("/api/projects").json())

    assert client.post("/api/auth/logout").status_code == 204
    login = client.post("/api/auth/login", json={"email": "first@example.com", "password": "correct-horse-42"})
    assert login.status_code == 200
    assert client.get("/api/workspace").json()["id"] == first_project_id

def template_bytes():
    data = io.BytesIO()
    with zipfile.ZipFile(data, "w") as archive:
        archive.writestr("ppt/presentation.xml", "<presentation/>")
    return data.getvalue()


def create_prepared_template(client, pipeline, tmp_path):
    project = client.post("/api/projects", json={"name": "Демо"}).json()
    response = client.post(
        f"/api/projects/{project['id']}/sources",
        data={"kind": "template"},
        files={"file": ("new-template.pptx", template_bytes(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
    )
    assert response.status_code == 201, response.text
    source = response.json()
    assert source["preparation"]["status"] == "queued"

    def prepare(template_path, output_dir):
        assert template_path.is_file()
        result = output_dir / "template_ir.json"
        result.write_text('{"layouts":[]}', encoding="utf-8")
        return result

    pipeline.prepare = prepare
    jobs.prepare_source(source["id"])
    ready = client.get(f"/api/sources/{source['id']}").json()
    assert ready["preparation"]["status"] == "ready"
    return project, source


def create_run(client, project, source, *, key="create-one"):
    return client.post(
        f"/api/projects/{project['id']}/runs",
        headers={"Idempotency-Key": key},
        json={"brief": "Новые факты для презентации", "slide_count": 3, "template_source_id": source["id"], "content_source_ids": []},
    )


def test_one_template_and_text_prompt_need_no_extra_files_or_slide_setting(environment):
    client, pipeline, tmp_path = environment
    project, source = create_prepared_template(client, pipeline, tmp_path)
    response = client.post(f"/api/projects/{project['id']}/runs",
                           json={"template_source_id": source["id"], "brief": "ИИ"})
    assert response.status_code == 202, response.text
    with db_module.get_session_factory()() as db:
        stored = db.get(Run, response.json()["id"])
        assert stored.config["slide_count"] == 10
        assert stored.config["brief"] == "ИИ"
        assert stored.config["content_source_ids"] == []


def fake_variants(output_dir):
    variants = []
    for number in range(1, 4):
        pptx = output_dir / f"variant-{number}.pptx"
        pdf = output_dir / f"variant-{number}.pdf"
        html = output_dir / f"variant-{number}.html"
        preview = output_dir / f"variant-{number}.png"
        pptx.write_bytes(template_bytes())
        pdf.write_bytes(b"%PDF-1.4\n%%EOF")
        html.write_text("<html></html>", encoding="utf-8")
        preview.write_bytes(b"\x89PNG\r\n\x1a\n")
        variants.append(SimpleNamespace(
            variant_id=f"variant-{number}", pptx_path=pptx, pdf_path=pdf,
            html_path=html, preview_paths=[preview], issues=[], metrics={"slides": 3},
        ))
    return variants


def test_prepare_and_idempotent_generation(environment):
    client, pipeline, tmp_path = environment
    project, source = create_prepared_template(client, pipeline, tmp_path)
    first = create_run(client, project, source)
    assert first.status_code == 202, first.text
    again = create_run(client, project, source)
    assert again.status_code == 202
    assert first.json()["id"] == again.json()["id"]
    conflict = client.post(
        f"/api/projects/{project['id']}/runs",
        headers={"Idempotency-Key": "create-one"},
        json={"brief": "Другой бриф", "slide_count": 3, "template_source_id": source["id"]},
    )
    assert conflict.status_code == 409

    def generate(template_path, brief, slide_count, output_dir, progress=None, content_paths=None, prepared_path=None):
        assert prepared_path.is_file()
        assert slide_count == 3
        return fake_variants(output_dir)

    pipeline.generate = generate
    jobs.execute_run(first.json()["id"])
    result = client.get(f"/api/runs/{first.json()['id']}").json()
    assert result["status"] == "completed"
    assert len(result["versions"]) == 3
    assert result["durations"]["generation_seconds"] >= 0
    assert result["artifacts"][0]["kind"] == "audit"
    versions = client.get(f"/api/projects/{project['id']}/versions").json()
    assert len(versions) == 3
    assert all(v["quality_status"] == "passed" for v in versions)
    assert all(len(v["preview_urls"]) == 1 for v in versions)
    pptx_id = next(a["id"] for a in versions[0]["artifacts"] if a["kind"] == "pptx")
    downloaded = client.get(f"/api/artifacts/{pptx_id}").content
    with zipfile.ZipFile(io.BytesIO(downloaded)) as archive:
        assert archive.read("ppt/presentation.xml") == b"<presentation/>"


def test_cancelled_execution_cannot_publish(environment):
    client, pipeline, tmp_path = environment
    project, source = create_prepared_template(client, pipeline, tmp_path)
    created = create_run(client, project, source)
    run_id = created.json()["id"]

    def generate(template_path, brief, slide_count, output_dir, progress=None, content_paths=None, prepared_path=None):
        with db_module.get_session_factory()() as db:
            run = db.get(Run, run_id)
            run.cancel_requested = True
            db.commit()
        return fake_variants(output_dir)

    pipeline.generate = generate
    jobs.execute_run(run_id)
    result = client.get(f"/api/runs/{run_id}").json()
    assert result["status"] == "cancelled"
    assert result["versions"] == []
    with db_module.get_session_factory()() as db:
        assert not list(db.scalars(select(Version).where(Version.run_id == run_id)))



def test_fast_worker_status_is_not_overwritten(environment, monkeypatch):
    client, pipeline, tmp_path = environment
    project = client.post("/api/projects", json={"name": "Быстрая очередь"}).json()

    def prepare(template_path, output_dir):
        result = output_dir / "template_ir.json"
        result.write_text("{}", encoding="utf-8")
        return result

    def generate(template_path, brief, slide_count, output_dir, progress=None, content_paths=None, prepared_path=None):
        return fake_variants(output_dir)

    pipeline.prepare = prepare
    pipeline.generate = generate
    monkeypatch.setattr(api, "enqueue_preparation", jobs.prepare_source)
    source_response = client.post(
        f"/api/projects/{project['id']}/sources",
        data={"kind": "template"},
        files={"file": ("template.pptx", template_bytes(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
    )
    assert source_response.status_code == 201, source_response.text
    source = source_response.json()
    assert source["preparation"]["status"] == "ready"

    monkeypatch.setattr(api, "enqueue_run", jobs.execute_run)
    run_response = create_run(client, project, source, key="fast")
    assert run_response.status_code == 202, run_response.text
    assert run_response.json()["status"] == "completed"
    assert len(run_response.json()["versions"]) == 3



def test_selected_repair_creates_child_version(environment):
    client, pipeline, tmp_path = environment
    project, source = create_prepared_template(client, pipeline, tmp_path)
    raw_issue = {
        "issue_id": "issue-overflow-1",
        "slide_id": "slide_1",
        "object_id": "4",
        "rule_id": "possible_text_overflow",
        "severity": "warning",
        "evidence": "Possible overflow",
        "repairability": "automatic",
    }

    def generate(template_path, brief, slide_count, output_dir, progress=None, content_paths=None, prepared_path=None):
        variants = fake_variants(output_dir)
        variants[0].issues = [raw_issue]
        return variants

    def repair(template_path, base_pptx_path, brief, slide_count, selected_issues, output_dir, progress=None, content_paths=None):
        assert base_pptx_path.is_file()
        assert len(selected_issues) == 1
        assert selected_issues[0]["issue_id"] == raw_issue["issue_id"]
        assert selected_issues[0]["repairability"] == "automatic"
        result = fake_variants(output_dir)[0]
        result.variant_id = "repair"
        return [result]

    pipeline.generate = generate
    pipeline.repair = repair
    run_id = create_run(client, project, source).json()["id"]
    jobs.execute_run(run_id)
    run = client.get(f"/api/runs/{run_id}").json()
    assert run["status"] == "completed_with_warnings"
    base_version_id = run["versions"][0]
    issues = client.get(f"/api/versions/{base_version_id}/issues").json()
    assert len(issues) == 1
    assert issues[0]["repairability"] == "automatic"
    repair_response = client.post(
        f"/api/versions/{base_version_id}/repairs",
        json={"issue_ids": [issues[0]["id"]]},
    )
    assert repair_response.status_code == 202, repair_response.text
    repair_run_id = repair_response.json()["id"]
    duplicate = client.post(
        f"/api/versions/{base_version_id}/repairs",
        json={"issue_ids": [issues[0]["id"]]},
    )
    assert duplicate.status_code == 409
    jobs.execute_run(repair_run_id)
    repaired = client.get(f"/api/runs/{repair_run_id}").json()
    assert repaired["status"] == "completed"
    assert len(repaired["versions"]) == 1
    child = client.get(f"/api/versions/{repaired['versions'][0]}").json()
    base = client.get(f"/api/versions/{base_version_id}").json()
    assert child["parent_version_id"] == base_version_id
    assert child["variant_id"] == base["variant_id"]
    assert child["ordinal"] == 2
    assert len(client.get(f"/api/projects/{project['id']}/versions").json()) == 4




def test_content_files_keep_distinct_source_names(environment):
    client, pipeline, tmp_path = environment
    project, template = create_prepared_template(client, pipeline, tmp_path)
    content_ids = []
    for filename in ("facts-a.txt", "facts-b.txt"):
        response = client.post(
            f"/api/projects/{project['id']}/sources",
            data={"kind": "content"},
            files={"file": (filename, b"A distinct statement about this presentation.", "text/plain")},
        )
        assert response.status_code == 201, response.text
        content_ids.append(response.json()["id"])

    def generate(template_path, brief, slide_count, output_dir, progress=None, content_paths=None, prepared_path=None):
        names = [path.name for path in content_paths]
        assert len(names) == 2
        assert len(set(names)) == 2
        assert any("facts-a.txt" in name for name in names)
        assert any("facts-b.txt" in name for name in names)
        return fake_variants(output_dir)

    pipeline.generate = generate
    response = client.post(
        f"/api/projects/{project['id']}/runs",
        json={
            "brief": "A long enough fact for a presentation.",
            "slide_count": 3,
            "template_source_id": template["id"],
            "content_source_ids": content_ids,
        },
    )
    assert response.status_code == 202, response.text
    jobs.execute_run(response.json()["id"])
    assert client.get(f"/api/runs/{response.json()['id']}").json()["status"] == "completed"



def test_system_mode_comes_from_worker_report(environment, monkeypatch):
    client, pipeline, tmp_path = environment
    unknown = client.get("/api/system").json()
    assert unknown["model_mode"] == "unknown"
    assert unknown["worker_status"] == "unknown"

    monkeypatch.setenv("OPENROUTER_API_KEY", "never-expose-this-key")
    jobs.publish_capability()
    reported = client.get("/api/system")
    assert reported.status_code == 200
    assert reported.json()["model_mode"] == "configured_api"
    assert reported.json()["worker_status"] == "reported"
    assert reported.json()["text_model"] == "qwen/qwen3.8-27b"
    assert len(reported.json()["agents"]) == 9
    assert reported.json()["provider"] == "openrouter"
    assert "never-expose-this-key" not in reported.text




def test_stale_execution_is_interrupted_and_retryable(environment, monkeypatch):
    from datetime import datetime, timedelta, timezone

    from app.models import Execution

    client, pipeline, tmp_path = environment
    project, source = create_prepared_template(client, pipeline, tmp_path)
    run_id = create_run(client, project, source).json()["id"]
    lease = "stale-lease"
    with db_module.get_session_factory()() as db:
        run = db.get(Run, run_id)
        run.status = "running"
        run.lease_token = lease
        db.add(Execution(
            run_id=run_id,
            attempt_no=1,
            lease_token=lease,
            status="running",
            heartbeat_at=datetime.now(timezone.utc) - timedelta(minutes=11),
        ))
        db.commit()

    jobs.reconcile_pending()
    interrupted = client.get(f"/api/runs/{run_id}").json()
    assert interrupted["status"] == "interrupted"

    monkeypatch.setattr(api, "enqueue_run", lambda run_id, force=False: None)
    retried = client.post(f"/api/runs/{run_id}/retry")
    assert retried.status_code == 202, retried.text
    assert retried.json()["status"] == "queued"




def test_unsafe_archive_is_rejected(environment):
    client, _, _ = environment
    project = client.post("/api/projects", json={"name": "Archive validation"}).json()
    data = io.BytesIO()
    with zipfile.ZipFile(data, "w") as archive:
        archive.writestr("ppt/presentation.xml", "<presentation/>")
        archive.writestr("../outside.xml", "bad")
    response = client.post(f"/api/projects/{project['id']}/sources", data={"kind": "template"},
        files={"file": ("unsafe.pptx", data.getvalue())})
    assert response.status_code == 422
    assert client.get(f"/api/projects/{project['id']}/sources").json() == []


def test_archive_expansion_limit(environment, monkeypatch):
    client, _, _ = environment
    monkeypatch.setenv("AYA_MAX_ARCHIVE_BYTES", "100")
    get_settings.cache_clear()
    project = client.post("/api/projects", json={"name": "Archive limit"}).json()
    data = io.BytesIO()
    with zipfile.ZipFile(data, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("ppt/presentation.xml", "x" * 101)
    response = client.post(f"/api/projects/{project['id']}/sources", data={"kind": "template"},
        files={"file": ("large.pptx", data.getvalue())})
    assert response.status_code == 422


def test_preparation_queue_does_not_reference_undefined_force(monkeypatch):
    from app import queueing
    calls = []
    class ExistingJob:
        def get_status(self, refresh): return "queued"
        def delete(self): calls.append("delete")
    queue = SimpleNamespace(fetch_job=lambda key: ExistingJob(), enqueue=lambda *args, **kwargs: calls.append("enqueue"))
    monkeypatch.setattr(queueing, "get_settings", lambda: SimpleNamespace(inline_jobs=False, local_jobs=False))
    monkeypatch.setattr(queueing, "get_queue", lambda: queue)
    queueing.enqueue_preparation("source")
    assert calls == []
    queueing.enqueue_preparation("source", force=True)
    assert calls == ["delete", "enqueue"]


def test_local_preparation_returns_before_job_finishes(environment, monkeypatch):
    from threading import Event
    from app import queueing
    client, pipeline, _ = environment
    started, release = Event(), Event()
    def prepare(path, output_dir):
        started.set()
        if not release.wait(5): raise RuntimeError("test worker was not released")
        target = output_dir / "template_ir.json"
        target.write_text("{}")
        return target
    pipeline.prepare = prepare
    monkeypatch.setenv("AYA_LOCAL_JOBS", "true")
    get_settings.cache_clear()
    monkeypatch.setattr(api, "enqueue_preparation", queueing.enqueue_preparation)
    project = client.post("/api/projects", json={"name": "Background job"}).json()
    try:
        response = client.post(f"/api/projects/{project['id']}/sources", data={"kind": "template"},
            files={"file": ("template.pptx", template_bytes())})
        assert response.status_code == 201
        assert started.wait(2)
        assert response.json()["preparation"]["status"] in {"queued", "running"}
        source_id = response.json()["id"]
        with queueing._lock:
            future = queueing._pending[f"prepare-{source_id}"]
        release.set()
        future.result(timeout=5)
        assert client.get(f"/api/sources/{source_id}").json()["preparation"]["status"] == "ready"
    finally:
        release.set()
        queueing.stop_local_jobs()


def test_storage_refuses_path_escape_and_overwrite(tmp_path):
    from app.storage import LocalStorage
    import pytest
    storage = LocalStorage(tmp_path / "store")
    source = tmp_path / "source.txt"
    source.write_text("first")
    with pytest.raises(ValueError): storage.path("../escape.txt")
    storage.put_file(source, "result.txt")
    source.write_text("second")
    with pytest.raises(FileExistsError): storage.put_file(source, "result.txt")
    assert storage.path("result.txt").read_text() == "first"


def test_requested_slide_edit_creates_only_one_new_variant(environment):
    client, pipeline, tmp_path = environment
    project, source = create_prepared_template(client, pipeline, tmp_path)
    pipeline.generate = lambda template_path, brief, slide_count, output_dir, progress=None, content_paths=None, prepared_path=None: fake_variants(output_dir)
    run_id = create_run(client, project, source).json()["id"]
    jobs.execute_run(run_id)
    original_ids = client.get(f"/api/runs/{run_id}").json()["versions"]
    assert len(original_ids) == 3

    def edit_slide(template_path, base_pptx_path, brief, slide_count, slide_index, instruction, output_dir, progress=None, content_paths=None):
        assert base_pptx_path.is_file()
        assert slide_count == 3
        assert slide_index == 2
        assert instruction == "Сократи заголовок"
        result = fake_variants(output_dir)[0]
        result.variant_id = "edit"
        return [result]

    pipeline.edit_slide = edit_slide
    invalid = client.post(f"/api/versions/{original_ids[0]}/edits", json={"slide_index": 4, "prompt": "Сократи заголовок"})
    assert invalid.status_code == 422
    response = client.post(f"/api/versions/{original_ids[0]}/edits", json={"slide_index": 2, "prompt": "Сократи заголовок"})
    assert response.status_code == 202, response.text
    jobs.execute_run(response.json()["id"])
    edited = client.get(f"/api/runs/{response.json()['id']}").json()
    assert edited["status"] == "completed"
    assert len(edited["versions"]) == 1
    child = client.get(f"/api/versions/{edited['versions'][0]}").json()
    assert child["parent_version_id"] == original_ids[0]
    assert len(client.get(f"/api/projects/{project['id']}/versions").json()) == 4
    assert all(client.get(f"/api/versions/{version_id}").status_code == 200 for version_id in original_ids[1:])


def real_template_bytes():
    from pptx import Presentation as Deck
    deck = Deck()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Old title"
    slide.placeholders[1].text = "Old body"
    data = io.BytesIO()
    deck.save(data)
    return data.getvalue()


def test_template_library_is_shared_portable_and_used_by_runs(environment, monkeypatch):
    import json as _json
    from app import library
    from app.models import Template
    from engine.ingest import inspect_template, save_prepared

    client, pipeline, tmp_path = environment
    monkeypatch.setattr(library, "enqueue_template", lambda template_id: None)
    uploaded = client.post("/api/templates", files={"file": ("Корпоративный.pptx", real_template_bytes(),
                                                              "application/vnd.openxmlformats-officedocument.presentationml.presentation")})
    assert uploaded.status_code == 201, uploaded.text
    template = uploaded.json()
    assert template["name"] == "Корпоративный" and template["preparation"]["status"] == "queued"

    def prepare(template_path, output_dir):
        prepared = inspect_template(template_path)
        prepared.analysis_mode = "vision_model"
        for composition in prepared.compositions:
            composition.object_roles = {str(slot.shape_id): "replaceable" for slot in composition.slots}
            composition.archetype = "content"
        return save_prepared(prepared, output_dir / "template_ir.json")

    def template_previews(template_path, prepared_path, output_dir):
        image = output_dir / "slide-1.png"
        image.write_bytes(b"\x89PNG\r\n\x1a\n" + b"0" * 16)
        return [image]

    pipeline.prepare, pipeline.template_previews = prepare, template_previews
    jobs.prepare_template(template["id"])
    ready = client.get(f"/api/templates/{template['id']}").json()
    assert ready["preparation"]["status"] == "ready" and ready["slide_count"] == 1
    assert client.get(ready["preview_urls"][0]).content.startswith(b"\x89PNG")

    # The library is not tied to an account: another user sees the same template.
    client.post("/api/auth/logout")
    client.post("/api/auth/register", json={"email": "other@example.com", "password": "another-password-42"})
    assert [item["id"] for item in client.get("/api/templates").json()] == [template["id"]]

    package = client.get(ready["export_url"])
    assert package.status_code == 200
    with zipfile.ZipFile(io.BytesIO(package.content)) as archive:
        assert _json.loads(archive.read("manifest.json"))["format"] == "lukas-template"
        members = {name: archive.read(name) for name in archive.namelist()}

    # Importing on a server that has never seen the template needs no new analysis.
    with db_module.get_session_factory()() as db:
        db.delete(db.get(Template, template["id"]))
        db.commit()
    imported = client.post("/api/templates/import", files={"file": ("Корпоративный.template.zip", package.content, "application/zip")})
    assert imported.status_code == 201, imported.text
    imported = imported.json()
    assert imported["origin"] == "import" and imported["preparation"]["status"] == "ready"
    assert len(imported["preview_urls"]) == 1

    tampered_ir = _json.loads(members["template_ir.json"])
    tampered_ir["compositions"][0]["object_roles"] = {"999": "replaceable"}
    tampered = io.BytesIO()
    with zipfile.ZipFile(tampered, "w") as archive:
        for name, data in members.items():
            archive.writestr(name, _json.dumps(tampered_ir) if name == "template_ir.json" else data)
    rejected = client.post("/api/templates/import", files={"file": ("bad.zip", tampered.getvalue(), "application/zip")})
    assert rejected.status_code == 422

    workspace = client.get("/api/workspace").json()
    created = client.post(f"/api/projects/{workspace['id']}/runs",
                          json={"brief": "Новые факты", "slide_count": 3, "template_id": imported["id"]})
    assert created.status_code == 202, created.text

    def generate(template_path, brief, slide_count, output_dir, progress=None, content_paths=None, prepared_path=None):
        assert template_path.is_file() and prepared_path.is_file()
        return fake_variants(output_dir)

    pipeline.generate = generate
    jobs.execute_run(created.json()["id"])
    runs = client.get(f"/api/projects/{workspace['id']}/runs").json()
    assert runs[0]["status"] == "completed" and runs[0]["template_name"] == "Корпоративный"
    assert runs[0]["brief_preview"] == "Новые факты" and runs[0]["slide_count"] == 3
    both = client.post(f"/api/projects/{workspace['id']}/runs",
                       json={"brief": "x", "template_id": imported["id"], "template_source_id": imported["id"]})
    assert both.status_code == 422


def test_project_template_package_moves_with_its_fonts_and_needs_no_new_analysis(environment, monkeypatch):
    import shutil
    from pptx import Presentation as Deck
    from engine import typography
    from engine.ingest import inspect_template, save_prepared

    client, pipeline, tmp_path = environment
    system_font = next((path for family in ("Arial", "DejaVu Sans", "Liberation Sans")
                        for path, _ in typography.font_files(family) if path.lower().endswith(".ttf")), None)
    if system_font is None:
        pytest.skip("No TrueType font installed")
    store = tmp_path / "fonts"
    (store / "Brand_Sans").mkdir(parents=True)
    shutil.copyfile(system_font, store / "Brand_Sans" / "regular.ttf")
    monkeypatch.setenv("AYA_FONT_DIR", str(store))
    typography.refresh_fonts()
    deck = Deck()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Фирменный заголовок"
    slide.shapes.title.text_frame.paragraphs[0].runs[0].font.name = "Brand Sans"
    data = io.BytesIO()
    deck.save(data)
    project = client.post("/api/projects", json={"name": "Демо"}).json()
    source = client.post(f"/api/projects/{project['id']}/sources", data={"kind": "template"},
                         files={"file": ("Бренд.pptx", data.getvalue(), "application/octet-stream")}).json()

    def prepare(template_path, output_dir):
        prepared = inspect_template(template_path)
        prepared.analysis_mode = "vision_model"
        return save_prepared(prepared, output_dir / "template_ir.json")

    pipeline.prepare = prepare
    jobs.prepare_source(source["id"])
    url = f"/api/projects/{project['id']}/sources/{source['id']}"
    package = client.get(f"{url}/export")
    assert package.status_code == 200, package.text
    with zipfile.ZipFile(io.BytesIO(package.content)) as archive:
        names = set(archive.namelist())
    assert {"manifest.json", "template.pptx", "template_ir.json", "fonts/Brand_Sans/regular.ttf"} <= names

    # Another machine: its font store is empty and the template is not in the project.
    other_store = tmp_path / "other-fonts"
    monkeypatch.setenv("AYA_FONT_DIR", str(other_store))
    typography.refresh_fonts()
    assert client.delete(url).status_code == 204
    imported = client.post(f"/api/projects/{project['id']}/sources/import",
                           files={"file": ("Бренд.template.zip", package.content, "application/zip")})
    assert imported.status_code == 201, imported.text
    imported = imported.json()
    assert imported["id"] != source["id"] and imported["filename"] == "Бренд.pptx"
    assert imported["preparation"]["status"] == "ready"
    assert (other_store / "Brand_Sans" / "regular.ttf").is_file() and typography.font_files("Brand Sans")
    again = client.post(f"/api/projects/{project['id']}/sources/import",
                        files={"file": ("Бренд.template.zip", package.content, "application/zip")})
    assert again.json()["id"] == imported["id"]
    assert [item["id"] for item in client.get(f"/api/projects/{project['id']}/sources").json()] == [imported["id"]]
    created = client.post(f"/api/projects/{project['id']}/runs",
                          json={"brief": "Новые факты", "slide_count": 3, "template_source_id": imported["id"]})
    assert created.status_code == 202, created.text
    for name, body in (("template.pptx", data.getvalue()), ("broken.zip", b"not a zip")):
        rejected = client.post(f"/api/projects/{project['id']}/sources/import", files={"file": (name, body, "application/zip")})
        assert rejected.status_code == 422
    typography.refresh_fonts()


def test_stale_template_preparation_becomes_retryable(environment, monkeypatch):
    from datetime import datetime, timedelta, timezone
    client, pipeline, tmp_path = environment
    project, source = create_prepared_template(client, pipeline, tmp_path)
    old = (datetime.now(timezone.utc) - timedelta(minutes=15)).isoformat()
    with db_module.get_session_factory()() as db:
        record = db.get(Source, source["id"])
        record.preparation_status = "running"
        record.metadata_json = {"preparation_lease": "old-worker", "preparation_started_at": old}
        db.commit()
    monkeypatch.setattr(jobs, "enqueue_preparation", lambda source_id: None)
    jobs.reconcile_pending()
    result = client.get(f"/api/sources/{source['id']}").json()
    assert result["preparation"]["status"] == "failed"
    assert "Повторить" in result["preparation"]["error"]
    with db_module.get_session_factory()() as db:
        assert db.get(Source, source["id"]).metadata_json["preparation_lease"] is None
    retry = client.post(f"/api/projects/{project['id']}/sources/{source['id']}/prepare")
    assert retry.status_code == 202
    assert retry.json()["preparation"]["status"] == "queued"


def test_live_preparation_heartbeat_prevents_false_interruption(environment, monkeypatch):
    from datetime import datetime, timedelta, timezone
    client, pipeline, tmp_path = environment
    _, source = create_prepared_template(client, pipeline, tmp_path)
    with db_module.get_session_factory()() as db:
        record = db.get(Source, source["id"])
        record.preparation_status = "running"
        record.metadata_json = {
            "preparation_lease": "active-worker",
            "preparation_started_at": (datetime.now(timezone.utc) - timedelta(minutes=15)).isoformat(),
            "preparation_heartbeat_at": datetime.now(timezone.utc).isoformat(),
        }
        db.commit()
    jobs.reconcile_pending()
    assert client.get(f"/api/sources/{source['id']}").json()["preparation"]["status"] == "running"


def test_project_runs_restore_progress_and_enforce_ownership(environment):
    client, pipeline, tmp_path = environment
    project, source = create_prepared_template(client, pipeline, tmp_path)
    run = create_run(client, project, source).json()
    response = client.get(f"/api/projects/{project['id']}/runs")
    assert response.status_code == 200
    assert [item["id"] for item in response.json()] == [run["id"]]
    client.post("/api/auth/logout")
    client.post("/api/auth/register", json={"email": "other-run@example.com", "password": "different-password-42"})
    assert client.get(f"/api/projects/{project['id']}/runs").status_code == 404


def test_slide_count_accepts_one_to_sixty(environment):
    client, pipeline, tmp_path = environment
    project, source = create_prepared_template(client, pipeline, tmp_path)
    url = f"/api/projects/{project['id']}/runs"
    for count in (1, 60):
        response = client.post(url, json={"brief": "Тема", "slide_count": count, "template_source_id": source["id"]})
        assert response.status_code == 202, response.text
        with db_module.get_session_factory()() as db:
            assert db.get(Run, response.json()["id"]).config["slide_count"] == count
    for count in (0, 61):
        response = client.post(url, json={"brief": "Тема", "slide_count": count, "template_source_id": source["id"]})
        assert response.status_code == 422


def test_delete_template_removes_it_from_library_and_keeps_existing_run(environment):
    client, pipeline, tmp_path = environment
    project, source = create_prepared_template(client, pipeline, tmp_path)
    run = create_run(client, project, source, key="before-delete")
    assert run.status_code == 202
    url = f"/api/projects/{project['id']}/sources/{source['id']}"
    assert client.delete(url).status_code == 204
    assert client.get(f"/api/projects/{project['id']}/sources").json() == []
    assert client.get(f"/api/sources/{source['id']}").status_code == 404
    assert client.get(f"/api/runs/{run.json()['id']}").status_code == 200
    later = client.post(f"/api/projects/{project['id']}/runs", json={"brief": "Новая тема", "slide_count": 10, "template_source_id": source["id"]})
    assert later.status_code == 422


def test_deleted_template_is_hidden_but_its_presentations_remain(environment):
    client, pipeline, tmp_path = environment
    project, source = create_prepared_template(client, pipeline, tmp_path)
    run = create_run(client, project, source).json()
    url = f"/api/projects/{project['id']}/sources/{source['id']}"
    assert client.delete(url).status_code == 204
    assert client.get(f"/api/projects/{project['id']}/sources").json() == []
    assert client.get(f"/api/sources/{source['id']}").status_code == 404
    assert client.post(f"{url}/prepare").status_code == 404
    assert client.delete(url).status_code == 404
    response = client.post(f"/api/projects/{project['id']}/runs",
                           json={"template_source_id": source["id"], "brief": "Новый бриф"})
    assert response.status_code == 422
    # The file stays for the presentation already made from it.
    assert client.get(f"/api/runs/{run['id']}").status_code == 200
    with db_module.get_session_factory()() as db:
        assert get_storage().exists(db.get(Source, source["id"]).storage_key)
    client.post("/api/auth/logout")
    client.post("/api/auth/register", json={"email": "other-delete@example.com", "password": "different-password-42"})
    assert client.delete(url).status_code == 404


def test_stale_library_template_preparation_becomes_retryable(environment, monkeypatch):
    from datetime import datetime, timedelta, timezone
    from app import library
    from app.models import Template

    client, pipeline, tmp_path = environment
    monkeypatch.setattr(library, "enqueue_template", lambda template_id: None)
    template = client.post("/api/templates", files={"file": ("Зависший.pptx", real_template_bytes(),
                                                             "application/vnd.openxmlformats-officedocument.presentationml.presentation")}).json()
    with db_module.get_session_factory()() as db:
        record = db.get(Template, template["id"])
        record.preparation_status = "running"
        record.metadata_json = {"preparation_lease": "old-worker",
                                "preparation_heartbeat_at": (datetime.now(timezone.utc) - timedelta(minutes=15)).isoformat()}
        db.commit()
    jobs.reconcile_pending()
    result = client.get(f"/api/templates/{template['id']}").json()
    assert result["preparation"]["status"] == "failed" and "Повторить" in result["preparation"]["error"]
    retry = client.post(f"/api/templates/{template['id']}/prepare")
    assert retry.status_code == 202 and retry.json()["preparation"]["status"] == "queued"
