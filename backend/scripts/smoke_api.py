"""Real local API smoke: one sample template -> three downloadable variants."""

from __future__ import annotations

import io
import argparse
import json
import os
import sys
import time
import zipfile
from datetime import datetime
from pathlib import Path
from uuid import uuid4

from fastapi.testclient import TestClient
from pptx import Presentation


def require(response, expected: int):
    if response.status_code != expected:
        raise RuntimeError(f"{response.request.method} {response.request.url.path}: HTTP {response.status_code}: {response.text[:1200]}")
    return response.json()


def run(template_path: Path, brief_path: Path, data_path: Path | None = None, *, live: bool = False, slide_count: int = 3) -> dict:
    workspace = Path(__file__).resolve().parents[1]
    output_dir = workspace / "data" / "api-smoke" / datetime.now().strftime("%Y%m%d-%H%M%S")
    output_dir = output_dir.with_name(output_dir.name + "-" + uuid4().hex[:6])
    output_dir.mkdir(parents=True)
    os.environ["AYA_DATABASE_URL"] = f"sqlite:///{(output_dir / 'smoke.db').as_posix()}"
    os.environ["AYA_STORAGE_ROOT"] = str(output_dir / "storage")
    os.environ["AYA_INLINE_JOBS"] = "true"
    os.environ["AYA_LOCAL_JOBS"] = "false"
    if not live:
        raise RuntimeError("API smoke requires --live: all generation uses paid OpenRouter models")
    from dotenv import load_dotenv
    load_dotenv(workspace / ".env", override=False)
    if not os.getenv("OPENROUTER_API_KEY"):
        raise RuntimeError("--live requires OPENROUTER_API_KEY in backend/.env")

    from app.config import get_settings
    from app.db import get_engine, get_session_factory

    get_settings.cache_clear()
    get_engine.cache_clear()
    get_session_factory.cache_clear()
    from app.main import app

    report: dict = {
        "template": str(template_path),
        "brief": str(brief_path),
        "output_dir": str(output_dir),
        "started_at": datetime.now().isoformat(timespec="seconds"),
    }
    try:
        with TestClient(app) as client:
            report["health"] = require(client.get("/api/health"), 200)
            report["system"] = require(client.get("/api/system"), 200)
            smoke_email = f"smoke-{uuid4().hex}@example.invalid"
            require(client.post("/api/auth/register", json={
                "email": smoke_email, "password": uuid4().hex + uuid4().hex,
            }), 201)
            project = require(client.post("/api/projects", json={"name": "API smoke"}), 201)
            report["project_id"] = project["id"]

            with template_path.open("rb") as file:
                source = require(client.post(
                    f"/api/projects/{project['id']}/sources",
                    data={"kind": "template"},
                    files={"file": (template_path.name, file, "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
                ), 201)
            report["source_id"] = source["id"]
            report["preparation"] = source["preparation"]
            if source["preparation"]["status"] != "ready":
                raise RuntimeError(f"Template preparation did not complete: {source['preparation']}")

            content_source_ids = []
            if data_path is not None:
                with data_path.open("rb") as file:
                    content = require(client.post(
                        f"/api/projects/{project['id']}/sources",
                        data={"kind": "content"},
                        files={"file": (data_path.name, file, "text/csv")},
                    ), 201)
                content_source_ids.append(content["id"])
                report["data_source_id"] = content["id"]

            brief = brief_path.read_text(encoding="utf-8")
            started = time.monotonic()
            created = require(client.post(
                f"/api/projects/{project['id']}/runs",
                headers={"Idempotency-Key": "smoke-generation"},
                json={
                    "brief": brief,
                    "slide_count": slide_count,
                    "template_source_id": source["id"],
                    "content_source_ids": content_source_ids,
                },
            ), 202)
            report["request_seconds"] = round(time.monotonic() - started, 3)
            report["run_id"] = created["id"]
            run_status = require(client.get(f"/api/runs/{created['id']}"), 200)
            report["run"] = run_status
            if run_status["status"] not in {"completed", "completed_with_warnings"}:
                raise RuntimeError(f"Generation failed: {run_status['status']}: {run_status['error']}")
            if len(run_status["versions"]) != 3:
                raise RuntimeError(f"Expected 3 versions, got {len(run_status['versions'])}")

            versions = require(client.get(f"/api/projects/{project['id']}/versions"), 200)
            report["variants"] = []
            for version_id in run_status["versions"]:
                version = require(client.get(f"/api/versions/{version_id}"), 200)
                issues = require(client.get(f"/api/versions/{version_id}/issues"), 200)
                entry = {
                    "id": version_id,
                    "variant_id": version["variant_id"],
                    "quality_status": version["quality_status"],
                    "issues": len(issues),
                    "automatic_issues": [issue["id"] for issue in issues if issue["repairability"] == "automatic"],
                    "artifacts": {},
                }
                by_kind: dict[str, list[dict]] = {}
                for artifact in version["artifacts"]:
                    by_kind.setdefault(artifact["kind"], []).append(artifact)
                for kind in ("pptx", "pdf", "html", "preview"):
                    if kind not in by_kind:
                        raise RuntimeError(f"Variant {version_id} missing {kind}")
                    response = client.get(by_kind[kind][0]["url"])
                    if response.status_code != 200:
                        raise RuntimeError(f"Could not download {kind} for {version_id}: HTTP {response.status_code}")
                    data = response.content
                    if not data:
                        raise RuntimeError(f"Empty {kind} for {version_id}")
                    if kind == "pptx":
                        with zipfile.ZipFile(io.BytesIO(data)) as archive:
                            if "ppt/presentation.xml" not in archive.namelist():
                                raise RuntimeError("Invalid PPTX package")
                        deck = Presentation(io.BytesIO(data))
                        if len(deck.slides) != slide_count:
                            raise RuntimeError(f"Expected {slide_count} slides in PPTX, got {len(deck.slides)}")
                        entry["editable_text_shapes"] = sum(
                            shape.has_text_frame for slide in deck.slides for shape in slide.shapes
                        )
                        if entry["editable_text_shapes"] < 3:
                            raise RuntimeError("PPTX has too few editable text objects")
                        if data_path is not None:
                            entry["native_chart_count"] = sum(shape.has_chart for slide in deck.slides for shape in slide.shapes)
                            entry["native_table_count"] = sum(shape.has_table for slide in deck.slides for shape in slide.shapes)
                            if entry["native_chart_count"] != 1 or entry["native_table_count"] != 1:
                                raise RuntimeError("PPTX lost its editable CSV chart or table")
                    elif kind == "pdf" and not data.startswith(b"%PDF"):
                        raise RuntimeError("Invalid PDF signature")
                    elif kind == "html" and b"<html" not in data[:500].lower():
                        raise RuntimeError("Invalid HTML")
                    elif kind == "preview" and not data.startswith(b"\x89PNG\r\n\x1a\n"):
                        raise RuntimeError("Invalid PNG preview")
                    entry["artifacts"][kind] = {"bytes": len(data), "count": len(by_kind[kind])}
                report["variants"].append(entry)

            audit_artifacts = [item for item in run_status["artifacts"] if item["kind"] == "audit"]
            if not audit_artifacts:
                raise RuntimeError("Run-level audit report is missing")
            audit_response = client.get(audit_artifacts[0]["url"])
            if audit_response.status_code != 200:
                raise RuntimeError("Run-level audit report is not downloadable")
            report["audit_report_bytes"] = len(audit_response.content)
            json.loads(audit_response.content)

            automatic = next(
                ((variant, issue_id) for variant in report["variants"] for issue_id in variant["automatic_issues"]),
                None,
            )
            if automatic:
                base, issue_id = automatic
                repair = require(client.post(
                    f"/api/versions/{base['id']}/repairs",
                    json={"issue_ids": [issue_id]},
                ), 202)
                repair_status = require(client.get(f"/api/runs/{repair['id']}"), 200)
                report["repair"] = {
                    "run_id": repair["id"],
                    "status": repair_status["status"],
                    "error": repair_status["error"],
                    "versions": repair_status["versions"],
                }
            else:
                report["repair"] = {"status": "skipped_no_automatic_issue"}
            if report["repair"]["status"] not in {"completed", "completed_with_warnings", "skipped_no_automatic_issue"}:
                raise RuntimeError("Selected repair did not complete")
            report["version_list_count"] = len(versions)
            report["result"] = "passed"
    except Exception as exc:
        report["result"] = "failed"
        report["error"] = str(exc)
        raise
    finally:
        report["finished_at"] = datetime.now().isoformat(timespec="seconds")
        (output_dir / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps(report, ensure_ascii=False, indent=2))
    return report


if __name__ == "__main__":
    project_root = Path(__file__).resolve().parents[2]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("template", nargs="?", type=Path, default=project_root / "sample_templates" / "VK WorkSpace.pptx")
    parser.add_argument("brief", nargs="?", type=Path, default=project_root / "examples" / "demo_brief.md")
    parser.add_argument("data", nargs="?", type=Path)
    parser.add_argument("--live", action="store_true", help="Use configured model API; sends the selected files to that provider")
    parser.add_argument("--slides", type=int, default=3, choices=range(3, 21))
    args = parser.parse_args()
    run(args.template, args.brief, args.data, live=args.live, slide_count=args.slides)
